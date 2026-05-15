import os
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from datetime import datetime, date, timedelta
from flask import Flask, render_template, request, jsonify, session, redirect, url_for, send_file, send_from_directory
from flask_cors import CORS
from pymongo import MongoClient
from bson import ObjectId
from io import BytesIO
import json
from dotenv import load_dotenv
from werkzeug.security import generate_password_hash, check_password_hash
import bcrypt
import secrets
import csv
import uuid

load_dotenv()

app = Flask(__name__)
app.secret_key = os.getenv('SECRET_KEY', 'pos_secret_key_2024')
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['PERMANENT_SESSION_LIFETIME'] = timedelta(hours=24)
CORS(app, supports_credentials=True)

# Email Configuration
SMTP_SERVER = os.getenv('SMTP_SERVER', '')
SMTP_PORT = int(os.getenv('SMTP_PORT', 587))
SMTP_USERNAME = os.getenv('SMTP_USERNAME', '')
SMTP_PASSWORD = os.getenv('SMTP_PASSWORD', '')
SMTP_SENDER = os.getenv('SMTP_SENDER', SMTP_USERNAME)
SMTP_USE_TLS = os.getenv('SMTP_USE_TLS', 'true').lower() == 'true'

# MongoDB Configuration
MONGO_URI = os.getenv('MONGO_URI', 'mongodb://localhost:27017/')

# In-memory data store (fallback if MongoDB not available)
USE_MEMORY_DB = False
memory_db = None
client = None
db = None

try:
    client = MongoClient(MONGO_URI, serverSelectionTimeoutMS=5000)
    client.server_info()
    db = client['pos_supermarket']
    print("Connected to MongoDB successfully")
except Exception as e:
    print(f"MongoDB connection failed: {e}")
    print("Using in-memory fallback database")
    USE_MEMORY_DB = True

    class MemoryDB:
        def __init__(self):
            self.products = []
            self.categories = []
            self.product_groups = []
            self.product_comments = []
            self.sales = []
            self.users = []
            self.vendors = []
            self.purchase_orders = []
            self.product_prices = []
            self.write_offs = []
            self._init_data()

        def _init_data(self):
            categories = ["Fruits", "Dairy", "Bakery", "Grains", "Meat", "Beverages", "Vegetables", "Snacks", "Clothing", "Shoes", "Grocery"]
            self.categories = [{"_id": str(i+1), "name": cat} for i, cat in enumerate(categories)]

            # Product Groups
            groups = [
                {"name": "Electronics", "description": "Electronic devices and accessories"},
                {"name": "Clothing", "description": "Apparel and fashion items"},
                {"name": "Food & Beverages", "description": "Edible items and drinks"},
                {"name": "Home & Kitchen", "description": "Household and kitchen products"},
                {"name": "Personal Care", "description": "Health and beauty products"},
            ]
            for i, g in enumerate(groups):
                g['_id'] = str(i + 1)
            self.product_groups = groups

            sample_products = [
                {"name": "Apple", "price": 2.50, "category": "Fruits", "barcode": "100001", "stock": 100, "unit": "kg", "supplier": "Fresh Farms", "cost_price": 1.50, "mrp": 3.00, "reorder_level": 20, "tax_rate": 5, "hsn_code": "0808", "expiry_date": "2026-08-15", "batch_no": "BT001", "manufacturer": "Fresh Farms Co"},
                {"name": "Banana", "price": 1.80, "category": "Fruits", "barcode": "100002", "stock": 150, "unit": "kg", "supplier": "Fresh Farms", "cost_price": 0.90, "mrp": 2.50, "reorder_level": 30, "tax_rate": 5, "hsn_code": "0808", "expiry_date": "2026-07-20", "batch_no": "BT002", "manufacturer": "Fresh Farms Co"},
                {"name": "Orange", "price": 3.20, "category": "Fruits", "barcode": "100003", "stock": 80, "unit": "kg", "supplier": "Citrus Valley", "cost_price": 2.00, "mrp": 4.00, "reorder_level": 20, "tax_rate": 5, "hsn_code": "0808", "expiry_date": "2026-09-10", "batch_no": "BT003", "manufacturer": "Citrus Valley Ltd"},
                {"name": "Milk", "price": 4.50, "category": "Dairy", "barcode": "200001", "stock": 50, "unit": "liter", "supplier": "DairyBest", "cost_price": 3.00, "mrp": 5.50, "reorder_level": 15, "tax_rate": 12, "hsn_code": "0401", "expiry_date": "2026-06-25", "batch_no": "BT004", "manufacturer": "DairyBest Inc"},
                {"name": "Cheese", "price": 8.99, "category": "Dairy", "barcode": "200002", "stock": 30, "unit": "piece", "supplier": "DairyBest", "cost_price": 5.50, "mrp": 10.00, "reorder_level": 10, "tax_rate": 12, "hsn_code": "0406", "expiry_date": "2026-10-01", "batch_no": "BT005", "manufacturer": "DairyBest Inc"},
                {"name": "Yogurt", "price": 2.20, "category": "Dairy", "barcode": "200003", "stock": 40, "unit": "cup", "supplier": "DairyBest", "cost_price": 1.20, "mrp": 3.00, "reorder_level": 15, "tax_rate": 12, "hsn_code": "0401", "expiry_date": "2026-07-30", "batch_no": "BT006", "manufacturer": "DairyBest Inc"},
                {"name": "Bread", "price": 3.50, "category": "Bakery", "barcode": "300001", "stock": 60, "unit": "loaf", "supplier": "Golden Bakery", "cost_price": 2.00, "mrp": 4.50, "reorder_level": 15, "tax_rate": 5, "hsn_code": "1905", "expiry_date": "2026-06-20", "batch_no": "BT007", "manufacturer": "Golden Bakery"},
                {"name": "Croissant", "price": 1.99, "category": "Bakery", "barcode": "300002", "stock": 45, "unit": "piece", "supplier": "Golden Bakery", "cost_price": 1.00, "mrp": 2.50, "reorder_level": 10, "tax_rate": 5, "hsn_code": "1905", "expiry_date": "2026-07-15", "batch_no": "BT008", "manufacturer": "Golden Bakery"},
                {"name": "Cookies", "price": 4.25, "category": "Bakery", "barcode": "300003", "stock": 55, "unit": "pack", "supplier": "Sweet Treats", "cost_price": 2.50, "mrp": 5.50, "reorder_level": 15, "tax_rate": 12, "hsn_code": "1905", "expiry_date": "2026-12-01", "batch_no": "BT009", "manufacturer": "Sweet Treats Co"},
                {"name": "Rice", "price": 12.99, "category": "Grains", "barcode": "400001", "stock": 70, "unit": "kg", "supplier": "Grain House", "cost_price": 8.00, "mrp": 15.00, "reorder_level": 20, "tax_rate": 5, "hsn_code": "1006", "expiry_date": "2027-01-15", "batch_no": "BT010", "manufacturer": "Grain House Ltd"},
                {"name": "Pasta", "price": 2.99, "category": "Grains", "barcode": "400002", "stock": 85, "unit": "pack", "supplier": "Grain House", "cost_price": 1.50, "mrp": 4.00, "reorder_level": 20, "tax_rate": 12, "hsn_code": "1902", "expiry_date": "2026-11-30", "batch_no": "BT011", "manufacturer": "Grain House Ltd"},
                {"name": "T-Shirt", "price": 15.99, "category": "Clothing", "barcode": "500001", "stock": 60, "unit": "piece", "supplier": "StyleWear", "cost_price": 8.00, "mrp": 20.00, "reorder_level": 15, "tax_rate": 5, "hsn_code": "6109", "expiry_date": "2027-06-01", "batch_no": "BT012", "manufacturer": "StyleWear Inc"},
                {"name": "Jeans", "price": 29.99, "category": "Clothing", "barcode": "500002", "stock": 40, "unit": "piece", "supplier": "StyleWear", "cost_price": 15.00, "mrp": 35.00, "reorder_level": 10, "tax_rate": 5, "hsn_code": "6204", "expiry_date": "2027-06-01", "batch_no": "BT013", "manufacturer": "StyleWear Inc"},
                {"name": "Running Shoes", "price": 45.99, "category": "Shoes", "barcode": "600001", "stock": 30, "unit": "pair", "supplier": "FootFit", "cost_price": 25.00, "mrp": 55.00, "reorder_level": 8, "tax_rate": 5, "hsn_code": "6403", "expiry_date": "2027-12-31", "batch_no": "BT014", "manufacturer": "FootFit Ltd"},
                {"name": "Sneakers", "price": 35.99, "category": "Shoes", "barcode": "600002", "stock": 25, "unit": "pair", "supplier": "FootFit", "cost_price": 20.00, "mrp": 42.00, "reorder_level": 8, "tax_rate": 5, "hsn_code": "6404", "expiry_date": "2027-12-31", "batch_no": "BT015", "manufacturer": "FootFit Ltd"},
                {"name": "Chicken", "price": 9.99, "category": "Meat", "barcode": "700001", "stock": 25, "unit": "kg", "supplier": "Farm Fresh", "cost_price": 6.00, "mrp": 12.00, "reorder_level": 8, "tax_rate": 5, "hsn_code": "0207", "expiry_date": "2026-06-15", "batch_no": "BT016", "manufacturer": "Farm Fresh Co"},
                {"name": "Beef", "price": 15.99, "category": "Meat", "barcode": "700002", "stock": 20, "unit": "kg", "supplier": "Farm Fresh", "cost_price": 10.00, "mrp": 18.00, "reorder_level": 8, "tax_rate": 5, "hsn_code": "0201", "expiry_date": "2026-06-10", "batch_no": "BT017", "manufacturer": "Farm Fresh Co"},
                {"name": "Fish", "price": 12.50, "category": "Meat", "barcode": "700003", "stock": 15, "unit": "kg", "supplier": "Sea Foods Ltd", "cost_price": 8.00, "mrp": 15.00, "reorder_level": 5, "tax_rate": 5, "hsn_code": "0306", "expiry_date": "2026-06-08", "batch_no": "BT018", "manufacturer": "Sea Foods Ltd"},
                {"name": "Coke", "price": 1.99, "category": "Beverages", "barcode": "800001", "stock": 200, "unit": "bottle", "supplier": "Beverage World", "cost_price": 0.80, "mrp": 2.50, "reorder_level": 50, "tax_rate": 12, "hsn_code": "2201", "expiry_date": "2026-10-15", "batch_no": "BT019", "manufacturer": "Beverage World"},
                {"name": "Juice", "price": 3.49, "category": "Beverages", "barcode": "800002", "stock": 120, "unit": "bottle", "supplier": "Beverage World", "cost_price": 1.50, "mrp": 4.00, "reorder_level": 30, "tax_rate": 12, "hsn_code": "2202", "expiry_date": "2026-09-30", "batch_no": "BT020", "manufacturer": "Beverage World"},
                {"name": "Water", "price": 0.99, "category": "Beverages", "barcode": "800003", "stock": 300, "unit": "bottle", "supplier": "Beverage World", "cost_price": 0.40, "mrp": 1.50, "reorder_level": 50, "tax_rate": 5, "hsn_code": "2201", "expiry_date": "2026-11-20", "batch_no": "BT021", "manufacturer": "Beverage World"},
                {"name": "Carrot", "price": 1.50, "category": "Vegetables", "barcode": "900001", "stock": 80, "unit": "kg", "supplier": "Green Valley", "cost_price": 0.70, "mrp": 2.00, "reorder_level": 20, "tax_rate": 5, "hsn_code": "0706", "expiry_date": "2026-06-20", "batch_no": "BT022", "manufacturer": "Green Valley Farms"},
                {"name": "Potato", "price": 1.20, "category": "Vegetables", "barcode": "900002", "stock": 100, "unit": "kg", "supplier": "Green Valley", "cost_price": 0.50, "mrp": 1.80, "reorder_level": 30, "tax_rate": 5, "hsn_code": "0701", "expiry_date": "2026-07-25", "batch_no": "BT023", "manufacturer": "Green Valley Farms"},
                {"name": "Snack Mix", "price": 3.99, "category": "Snacks", "barcode": "100001", "stock": 90, "unit": "pack", "supplier": "Crunchy Foods", "cost_price": 2.00, "mrp": 5.00, "reorder_level": 20, "tax_rate": 12, "hsn_code": "1904", "expiry_date": "2026-11-15", "batch_no": "BT024", "manufacturer": "Crunchy Foods Ltd"},
                {"name": "Ice Cream", "price": 2.50, "category": "Grocery", "barcode": "110001", "stock": 40, "unit": "liter", "supplier": "Cool Creams", "cost_price": 1.20, "mrp": 3.50, "reorder_level": 10, "tax_rate": 12, "hsn_code": "2105", "expiry_date": "2026-07-10", "batch_no": "BT025", "manufacturer": "Cool Creams Co"},
            ]
            for i, p in enumerate(sample_products):
                p['_id'] = str(i + 1)
            self.products = sample_products

            admin_pass_hash = bcrypt.hashpw("admin123".encode('utf-8'), bcrypt.gensalt())
            self.users = [
                {
                    "_id": "1", "username": "admin", "password": admin_pass_hash.decode('utf-8'),
                    "role": "admin", "full_name": "System Admin", "email": "admin@pos.com",
                    "created_at": datetime.utcnow().isoformat()
                }
            ]

            self.vendors = [
                {"_id": "1", "name": "Fresh Farms", "contact": "John Doe", "phone": "+1234567890",
                 "email": "fresh@farms.com", "address": "123 Farm Road", "status": "active",
                 "payment_terms": "Net 30", "created_at": datetime.utcnow().isoformat()},
                {"_id": "2", "name": "DairyBest", "contact": "Jane Smith", "phone": "+1234567891",
                 "email": "contact@dairybest.com", "address": "456 Dairy Lane", "status": "active",
                 "payment_terms": "Net 15", "created_at": datetime.utcnow().isoformat()},
                {"_id": "3", "name": "StyleWear", "contact": "Mike Brown", "phone": "+1234567892",
                 "email": "info@stylewear.com", "address": "789 Fashion Ave", "status": "active",
                 "payment_terms": "Net 45", "created_at": datetime.utcnow().isoformat()},
            ]

            # Colors
            colors = [
                {"_id": "1", "value": "red", "label": "Red"},
                {"_id": "2", "value": "blue", "label": "Blue"},
                {"_id": "3", "value": "green", "label": "Green"},
                {"_id": "4", "value": "yellow", "label": "Yellow"},
                {"_id": "5", "value": "orange", "label": "Orange"},
                {"_id": "6", "value": "purple", "label": "Purple"},
                {"_id": "7", "value": "pink", "label": "Pink"},
                {"_id": "8", "value": "black", "label": "Black"},
                {"_id": "9", "value": "white", "label": "White"},
                {"_id": "10", "value": "gray", "label": "Gray"},
                {"_id": "11", "value": "brown", "label": "Brown"},
                {"_id": "12", "value": "navy", "label": "Navy"},
                {"_id": "13", "value": "teal", "label": "Teal"},
                {"_id": "14", "value": "maroon", "label": "Maroon"},
                {"_id": "15", "value": "olive", "label": "Olive"},
                {"_id": "16", "value": "silver", "label": "Silver"},
                {"_id": "17", "value": "gold", "label": "Gold"},
                {"_id": "18", "value": "multi", "label": "Multi-Color"},
            ]
            self.colors = colors

            # Units of Measure
            units = [
                {"_id": "1", "name": "piece"},
                {"_id": "2", "name": "kg"},
                {"_id": "3", "name": "liter"},
                {"_id": "4", "name": "meter"},
                {"_id": "5", "name": "pack"},
                {"_id": "6", "name": "box"},
                {"_id": "7", "name": "pair"},
                {"_id": "8", "name": "set"},
            ]
            self.units = units

        def count_documents(self, collection_name, query=None):
            collection = getattr(self, collection_name, [])
            if query:
                if collection_name == 'users':
                    return len([u for u in collection if all(u.get(k) == v for k, v in query.items() if k != 'password')])
                return len([p for p in collection if all(p.get(k) == v for k, v in query.items())])
            return len(collection)

        def find(self, collection_name, query=None):
            collection = getattr(self, collection_name, [])
            if collection_name == 'users':
                return [u for u in collection]
            if query and '_id' in query:
                pid = str(query['_id'])
                return [p for p in collection if str(p.get('_id')) == pid]
            elif query and 'category' in query:
                return [p for p in collection if p.get('category') == query['category']]
            elif query and 'vendor_name' in query:
                return [p for p in collection if query['vendor_name'].lower() in (p.get('supplier') or '').lower()]
            return collection[:]

        def find_one(self, collection_name, query):
            collection = getattr(self, collection_name, [])
            if collection_name == 'users':
                if 'username' in query:
                    for u in collection:
                        if u.get('username') == query['username']:
                            return u
                if '_id' in query:
                    uid = str(query['_id'])
                    for u in collection:
                        if str(u.get('_id')) == uid:
                            return u
                return None
            if '_id' in query:
                pid = str(query['_id'])
                for p in collection:
                    if str(p.get('_id')) == pid:
                        return p
            return None

        def insert_one(self, collection_name, doc):
            collection = getattr(self, collection_name, [])
            next_id = str(len(collection) + 1)
            doc['_id'] = next_id
            collection.append(doc)
            setattr(self, collection_name, collection)
            class Result:
                def __init__(self, inserted_id):
                    self.inserted_id = inserted_id
            return Result(next_id)

        def update_one(self, collection_name, query, update):
            collection = getattr(self, collection_name, [])
            pid = str(query['_id'])
            for i, item in enumerate(collection):
                if str(item.get('_id')) == pid:
                    for k, v in update.get('$set', {}).items():
                        collection[i][k] = v
                    class Result:
                        def __init__(self, modified):
                            self.modified_count = modified
                    return Result(1)
            class Result:
                def __init__(self, modified):
                    self.modified_count = modified
            return Result(0)

        def delete_one(self, collection_name, query):
            collection = getattr(self, collection_name, [])
            pid = str(query['_id'])
            for i, item in enumerate(collection):
                if str(item.get('_id')) == pid:
                    del collection[i]
                    setattr(self, collection_name, collection)
                    class Result:
                        def __init__(self, deleted):
                            self.deleted_count = deleted
                    return Result(1)
            class Result:
                def __init__(self, deleted):
                    self.deleted_count = deleted
            return Result(0)

        def find_many(self, collection_name, query=None):
            collection = getattr(self, collection_name, [])
            if not query:
                return collection[:]
            results = []
            for item in collection:
                match = True
                for k, v in query.items():
                    if str(item.get(k, '')) != str(v):
                        match = False
                        break
                if match:
                    results.append(item)
            return results

        def aggregate(self, pipeline):
            results = []
            for sale in self.sales:
                results.append(sale)
            return results

    memory_db = MemoryDB()
    db = memory_db
    print("Using in-memory database (data will not persist after restart)")

# Initialize database collections
def init_db():
    """Initialize database with sample data if empty"""
    if not USE_MEMORY_DB:
        if db.products.count_documents({}) == 0:
            sample_products = [
                {"name": "Apple", "price": 2.50, "category": "Fruits", "barcode": "100001", "stock": 100, "unit": "kg",
                 "supplier": "Fresh Farms", "cost_price": 1.50, "mrp": 3.00, "reorder_level": 20, "tax_rate": 5,
                 "hsn_code": "0808", "expiry_date": "2026-08-15", "batch_no": "BT001", "manufacturer": "Fresh Farms Co",
                 "created_at": datetime.utcnow().isoformat()},
                {"name": "Banana", "price": 1.80, "category": "Fruits", "barcode": "100002", "stock": 150, "unit": "kg",
                 "supplier": "Fresh Farms", "cost_price": 0.90, "mrp": 2.50, "reorder_level": 30, "tax_rate": 5,
                 "hsn_code": "0808", "expiry_date": "2026-07-20", "batch_no": "BT002", "manufacturer": "Fresh Farms Co",
                 "created_at": datetime.utcnow().isoformat()},
                {"name": "Rice", "price": 12.99, "category": "Grains", "barcode": "400001", "stock": 70, "unit": "kg",
                 "supplier": "Grain House", "cost_price": 8.00, "mrp": 15.00, "reorder_level": 20, "tax_rate": 5,
                 "hsn_code": "1006", "expiry_date": "2027-01-15", "batch_no": "BT010", "manufacturer": "Grain House Ltd",
                 "created_at": datetime.utcnow().isoformat()},
                {"name": "T-Shirt", "price": 15.99, "category": "Clothing", "barcode": "500001", "stock": 60, "unit": "piece",
                 "supplier": "StyleWear", "cost_price": 8.00, "mrp": 20.00, "reorder_level": 15, "tax_rate": 5,
                 "hsn_code": "6109", "expiry_date": "2027-06-01", "batch_no": "BT012", "manufacturer": "StyleWear Inc",
                 "created_at": datetime.utcnow().isoformat()},
                {"name": "Running Shoes", "price": 45.99, "category": "Shoes", "barcode": "600001", "stock": 30, "unit": "pair",
                 "supplier": "FootFit", "cost_price": 25.00, "mrp": 55.00, "reorder_level": 8, "tax_rate": 5,
                 "hsn_code": "6403", "expiry_date": "2027-12-31", "batch_no": "BT014", "manufacturer": "FootFit Ltd",
                 "created_at": datetime.utcnow().isoformat()},
                {"name": "Chicken", "price": 9.99, "category": "Meat", "barcode": "700001", "stock": 25, "unit": "kg",
                 "supplier": "Farm Fresh", "cost_price": 6.00, "mrp": 12.00, "reorder_level": 8, "tax_rate": 5,
                 "hsn_code": "0207", "expiry_date": "2026-06-15", "batch_no": "BT016", "manufacturer": "Farm Fresh Co",
                 "created_at": datetime.utcnow().isoformat()},
                {"name": "Coke", "price": 1.99, "category": "Beverages", "barcode": "800001", "stock": 200, "unit": "bottle",
                 "supplier": "Beverage World", "cost_price": 0.80, "mrp": 2.50, "reorder_level": 50, "tax_rate": 12,
                 "hsn_code": "2201", "expiry_date": "2026-10-15", "batch_no": "BT019", "manufacturer": "Beverage World",
                 "created_at": datetime.utcnow().isoformat()},
            ]
            db.products.insert_many(sample_products)
            print("Sample products initialized")

        if db.categories.count_documents({}) == 0:
            categories = ["Fruits", "Dairy", "Bakery", "Grains", "Meat", "Beverages", "Vegetables", "Snacks", "Clothing", "Shoes", "Grocery"]
            db.categories.insert_many([{"name": cat} for cat in categories])
            print("Categories initialized")

        if db.vendors.count_documents({}) == 0:
            vendors = [
                {"name": "Fresh Farms", "contact": "John Doe", "phone": "+1234567890", "email": "fresh@farms.com",
                 "address": "123 Farm Road", "status": "active", "payment_terms": "Net 30",
                 "created_at": datetime.utcnow().isoformat()},
                {"name": "DairyBest", "contact": "Jane Smith", "phone": "+1234567891", "email": "contact@dairybest.com",
                 "address": "456 Dairy Lane", "status": "active", "payment_terms": "Net 15",
                 "created_at": datetime.utcnow().isoformat()},
                {"name": "StyleWear", "contact": "Mike Brown", "phone": "+1234567892", "email": "info@stylewear.com",
                 "address": "789 Fashion Ave", "status": "active", "payment_terms": "Net 45",
                 "created_at": datetime.utcnow().isoformat()},
            ]
            db.vendors.insert_many(vendors)
            print("Vendors initialized")

        if db.vendors.count_documents({}) == 0:
            vendors = [
                {"name": "Fresh Farms", "contact": "John Doe", "phone": "+1234567890", "email": "fresh@farms.com",
                 "address": "123 Farm Road", "status": "active", "payment_terms": "Net 30",
                 "created_at": datetime.utcnow().isoformat()},
                {"name": "DairyBest", "contact": "Jane Smith", "phone": "+1234567891", "email": "contact@dairybest.com",
                 "address": "456 Dairy Lane", "status": "active", "payment_terms": "Net 15",
                 "created_at": datetime.utcnow().isoformat()},
                {"name": "StyleWear", "contact": "Mike Brown", "phone": "+1234567892", "email": "info@stylewear.com",
                 "address": "789 Fashion Ave", "status": "active", "payment_terms": "Net 45",
                 "created_at": datetime.utcnow().isoformat()},
            ]
            db.vendors.insert_many(vendors)
            print("Vendors initialized")

        # Initialize product groups
        if USE_MEMORY_DB:
            if len(db.product_groups) == 0:
                groups = [
                    {"name": "Electronics", "description": "Electronic devices and accessories"},
                    {"name": "Clothing", "description": "Apparel and fashion items"},
                    {"name": "Food & Beverages", "description": "Edible items and drinks"},
                    {"name": "Home & Kitchen", "description": "Household and kitchen products"},
                    {"name": "Personal Care", "description": "Health and beauty products"},
                ]
                for i, g in enumerate(groups):
                    g['_id'] = str(i + 1)
                db.product_groups.extend(groups)
                print("Product groups initialized")
        else:
            if db.product_groups.count_documents({}) == 0:
                groups = [
                    {"name": "Electronics", "description": "Electronic devices and accessories"},
                    {"name": "Clothing", "description": "Apparel and fashion items"},
                    {"name": "Food & Beverages", "description": "Edible items and drinks"},
                    {"name": "Home & Kitchen", "description": "Household and kitchen products"},
                    {"name": "Personal Care", "description": "Health and beauty products"},
                ]
                db.product_groups.insert_many(groups)
                print("Product groups initialized")

    # Initialize admin user
    if USE_MEMORY_DB:
        if len(db.users) == 0:
            admin_hash = bcrypt.hashpw("admin123".encode('utf-8'), bcrypt.gensalt())
            db.insert_one('users', {
                "username": "admin", "password": admin_hash.decode('utf-8'),
                "role": "admin", "full_name": "System Admin", "email": "admin@pos.com",
                "created_at": datetime.utcnow().isoformat()
            })
            print("Default admin user created (admin/admin123)")
    else:
        if db.users.count_documents({}) == 0:
            admin_hash = bcrypt.hashpw("admin123".encode('utf-8'), bcrypt.gensalt())
            db.users.insert_one({
                "username": "admin", "password": admin_hash.decode('utf-8'),
                "role": "admin", "full_name": "System Admin", "email": "admin@pos.com",
                "created_at": datetime.utcnow().isoformat()
            })
            print("Default admin user created (admin/admin123)")

        if not db.users.find_one({'username': 'finance'}):
            hashed = bcrypt.hashpw('finance123'.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')
            db.users.insert_one({
                'username': 'finance',
                'password': hashed,
                'role': 'finance_manager',
                'full_name': 'Finance Manager',
                'email': 'finance@pos.com',
                'status': 'active',
                'created_at': datetime.utcnow().isoformat()
            })
            print("Finance manager user created (finance/finance123)")

    # Create indexes for procurement collections
    if not USE_MEMORY_DB:
        db.purchase_requisitions.create_index('status')
        db.purchase_requisitions.create_index('created_at')
        db.purchase_orders_v2.create_index('pr_id')
        db.purchase_orders_v2.create_index('status')
        db.goods_receipts.create_index('po_id')
        db.goods_receipts.create_index('status')
        db.vendor_invoices.create_index('gr_id')
        db.vendor_invoices.create_index('status')
        db.vendor_payments_v2.create_index('invoice_id')
        db.counters.create_index('_id')
        print("Procurement indexes created")

def serialize_doc(doc):
    if USE_MEMORY_DB:
        return doc
    if doc:
        doc['_id'] = str(doc['_id'])
    return doc


def generate_reset_token():
    """Generate a secure password reset token"""
    return secrets.token_urlsafe(32)


def send_password_reset_email(user_email, token, username, flask_request):
    """Send password reset email via SMTP"""
    reset_url = f"{flask_request.scheme}://{flask_request.host}/reset-password?token={token}"

    if not all([SMTP_SERVER, SMTP_USERNAME, SMTP_PASSWORD]):
        print(f"[DEMO MODE] No SMTP configured.")
        print(f"[DEMO MODE] Reset URL for {user_email}: {reset_url}")
        return False, None

    html = f"""
    <html>
    <body style="font-family: Segoe UI, Arial, sans-serif; max-width: 600px; margin: 0 auto; padding: 24px; background: #f5f6fa;">
        <div style="background: white; padding: 32px; border-radius: 12px; box-shadow: 0 4px 12px rgba(0,0,0,0.1);">
            <div style="text-align: center; margin-bottom: 24px;">
                <h2 style="color: #2c3e50; margin-bottom: 8px;">&#128274; Password Reset</h2>
                <p style="color: #7f8c8d; font-size: 15px;">Hello <strong>{username}</strong>,</p>
            </div>
            <p style="color: #555; line-height: 1.7;">
                We received a request to reset your password for your POS System account.
                Click the button below to set a new password:
            </p>
            <div style="text-align: center; margin: 32px 0;">
                <a href="{reset_url}"
                   style="background: linear-gradient(135deg, #667eea 0%, #764ba2 100%); color: white; padding: 14px 40px; border-radius: 8px; text-decoration: none; font-weight: 600; display: inline-block; font-size: 15px;">
                    Reset My Password
                </a>
            </div>
            <p style="font-size: 12px; color: #999; line-height: 1.6;">
                If you didn't request this, you can safely ignore this email.<br>
                This link expires in <strong>1 hour</strong>.
            </p>
            <hr style="border: none; border-top: 1px solid #eee; margin: 24px 0;">
            <p style="font-size: 11px; color: #bbb; text-align: center;">
                POS System &mdash; Departmental Store Management
            </p>
        </div>
    </body>
    </html>
    """

    try:
        msg = MIMEMultipart('alternative')
        msg['Subject'] = 'POS System - Password Reset Request'
        msg['From'] = SMTP_SENDER
        msg['To'] = user_email
        msg.attach(MIMEText(html, 'html'))

        with smtplib.SMTP(SMTP_SERVER, SMTP_PORT, timeout=15) as server:
            server.ehlo()
            server.starttls()
            server.ehlo()
            server.login(SMTP_USERNAME, SMTP_PASSWORD)
            server.sendmail(SMTP_SENDER, user_email, msg.as_string())

        print(f"[EMAIL OK] Reset email sent to {user_email}")
        return True, None

    except Exception as e:
        error_msg = f"{type(e).__name__}: {e}"
        print(f"[EMAIL ERROR] Failed to send reset email: {error_msg}")
        return False, error_msg


def get_next_id(collection_name):
    if USE_MEMORY_DB:
        collection = getattr(db, collection_name, [])
        if not collection:
            return "1"
        ids = [int(str(item.get('_id', '0'))) for item in collection]
        return str(max(ids) + 1)
    else:
        collection = getattr(db, collection_name)
        last = collection.find_one(sort=[('_id', -1)])
        if last:
            try:
                return str(int(str(last['_id'])) + 1)
            except:
                pass
        return "1"


# =============================================================================
# AUTH ROUTES
# =============================================================================

@app.before_request
def require_auth():
    """Protect dashboard and API routes that require authentication"""
    protected_paths = [
        '/dashboard', '/admin-dashboard', '/procurement-dashboard', '/sales-dashboard',
        '/api/users', '/api/products', '/api/categories', '/api/vendors',
        '/api/purchase-orders', '/api/prices', '/api/write-offs',
        '/api/reports', '/pos', '/reports',
        '/api/product-groups', '/api/comments', '/api/generate-barcode', '/api/colors',
        '/api/departments', '/api/payment-terms', '/api/units',
        '/products'
    ]

    protected_paths += ['/api/procurement/', '/procurement/pr', '/procurement/po',
                        '/procurement/gr', '/procurement/invoices', '/procurement/payments']

    protected_paths += ['/finance-dashboard', '/finance/', '/api/gl/', '/api/finance/']

    for path in protected_paths:
        if request.path.startswith(path):
            if 'user_id' not in session:
                if request.is_json:
                    return jsonify({'error': 'Authentication required'}), 401
                return redirect(url_for('login_page'))


@app.route('/login')
def login_page():
    if 'user_id' in session:
        return redirect_to_dashboard(session.get('role'))
    return render_template('login.html')


@app.route('/api/auth/login', methods=['POST'])
def login_api():
    data = request.json
    username = data.get('username', '')
    password = data.get('password', '')

    if USE_MEMORY_DB:
        user = db.find_one('users', {'username': username})
    else:
        user = db.users.find_one({'username': username})

    if not user:
        return jsonify({'success': False, 'message': 'Invalid username or password'}), 401

    stored_hash = user.get('password', '')
    if bcrypt.checkpw(password.encode('utf-8'), stored_hash.encode('utf-8')):
        session['user_id'] = str(user['_id'])
        session['username'] = user['username']
        session['role'] = user['role']
        session['full_name'] = user.get('full_name', user['username'])
        return jsonify({
            'success': True,
            'user': {
                'id': str(user['_id']),
                'username': user['username'],
                'role': user['role'],
                'full_name': user.get('full_name', user['username'])
            }
        })
    return jsonify({'success': False, 'message': 'Invalid username or password'}), 401


@app.route('/api/auth/logout', methods=['POST'])
def logout():
    session.clear()
    if request.is_json:
        return jsonify({'success': True})
    return jsonify({'success': True})


@app.route('/logout')
def logout_get():
    session.clear()
    return redirect(url_for('login_page'))


@app.route('/forgot-password')
def forgot_password_page():
    return render_template('forgot_password.html')


@app.route('/api/auth/forgot-password', methods=['POST'])
def forgot_password():
    data = request.json
    username = data.get('username', '').strip()
    email = data.get('email', '').strip().lower()

    if not username or not email:
        return jsonify({'success': False, 'message': 'Username and email address are required'}), 400

    # Find user by username first, then verify email matches
    if USE_MEMORY_DB:
        user = None
        for u in db.users:
            if u.get('username', '').lower() == username.lower():
                user = u
                break
    else:
        user = db.users.find_one({'username': {'$regex': f'^{username}$', '$options': 'i'}})

    # Verify both username and email match the same account
    if not user or user.get('email', '').lower() != email:
        return jsonify({'success': False, 'message': 'No account found with this username and email combination'}), 404

    # Generate reset token
    reset_token = generate_reset_token()
    reset_expiry = datetime.utcnow() + timedelta(hours=1)

    # Save token to user record
    if USE_MEMORY_DB:
        user['reset_token'] = reset_token
        user['reset_token_expiry'] = reset_expiry.isoformat()
    else:
        db.users.update_one(
            {'_id': user['_id']},
            {'$set': {
                'reset_token': reset_token,
                'reset_token_expiry': reset_expiry
            }}
        )

    # Send email
    email_sent, email_error = send_password_reset_email(email, reset_token, username, request)

    if email_sent:
        return jsonify({'success': True, 'message': 'A password reset link has been sent to your email address.'})

    # SMTP not configured — demo mode
    if not SMTP_SERVER:
        return jsonify({
            'success': True,
            'message': 'SMTP not configured. Use the token below to reset your password.',
            'demo_token': reset_token,
            'demo_url': f"/reset-password?token={reset_token}"
        })

    # SMTP configured but sending failed — return the actual error
    return jsonify({
        'success': False,
        'message': f'Account found but failed to send email. Error: {email_error}'
    }), 500


@app.route('/reset-password')
def reset_password_page():
    token = request.args.get('token', '')
    return render_template('reset_password.html', token=token)


@app.route('/api/auth/reset-password', methods=['POST'])
def reset_password():
    data = request.json
    token = data.get('token', '').strip()
    username = data.get('username', '').strip()
    email = data.get('email', '').strip().lower()
    new_password = data.get('password', '')

    if not token or not new_password:
        return jsonify({'success': False, 'message': 'Token and new password are required'}), 400

    if not username or not email:
        return jsonify({'success': False, 'message': 'Username and email are required for verification'}), 400

    if len(new_password) < 6:
        return jsonify({'success': False, 'message': 'Password must be at least 6 characters'}), 400

    # Find user by token
    if USE_MEMORY_DB:
        user = None
        for u in db.users:
            if u.get('reset_token') == token:
                user = u
                break
    else:
        user = db.users.find_one({'reset_token': token})

    if not user:
        return jsonify({'success': False, 'message': 'Invalid or expired reset link. Please request a new one.'}), 400

    # Verify username and email match the token owner
    if user.get('username', '').lower() != username.lower() or user.get('email', '').lower() != email:
        return jsonify({'success': False, 'message': 'Username or email does not match. Please check and try again.'}), 400

    # Check token expiry
    expiry_str = user.get('reset_token_expiry')
    if expiry_str:
        if isinstance(expiry_str, str):
            expiry = datetime.fromisoformat(expiry_str)
        else:
            expiry = expiry_str
        if datetime.utcnow() > expiry:
            return jsonify({'success': False, 'message': 'Reset link has expired. Please request a new one.'}), 400

    # Update password
    new_hash = bcrypt.hashpw(new_password.encode('utf-8'), bcrypt.gensalt())

    if USE_MEMORY_DB:
        user['password'] = new_hash.decode('utf-8')
        user['reset_token'] = None
        user['reset_token_expiry'] = None
    else:
        db.users.update_one(
            {'_id': user['_id']},
            {'$set': {'password': new_hash.decode('utf-8')},
             '$unset': {'reset_token': '', 'reset_token_expiry': ''}}
        )

    return jsonify({'success': True, 'message': 'Password has been reset successfully. You can now log in.'})


@app.route('/register')
def register_page():
    return render_template('register.html')


@app.route('/api/auth/register', methods=['POST'])
def register():
    data = request.json
    username = data.get('username', '').strip()
    password = data.get('password', '')
    full_name = data.get('full_name', '').strip()
    email = data.get('email', '').strip()
    role = data.get('role', 'sales_person')

    if not username or not password or not full_name:
        return jsonify({'success': False, 'message': 'Username, password and full name are required'}), 400

    if USE_MEMORY_DB:
        existing = db.find_one('users', {'username': username})
    else:
        existing = db.users.find_one({'username': username})

    if existing:
        return jsonify({'success': False, 'message': 'Username already exists'}), 400

    password_hash = bcrypt.hashpw(password.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')
    user_doc = {
        'username': username, 'password': password_hash, 'role': role,
        'full_name': full_name, 'email': email,
        'created_at': datetime.utcnow().isoformat(),
        'created_by': session.get('user_id', 'system')
    }

    if USE_MEMORY_DB:
        result = db.insert_one('users', user_doc)
    else:
        result = db.users.insert_one(user_doc)
        user_doc['_id'] = str(result.inserted_id)

    return jsonify({
        'success': True, 'message': 'User created successfully',
        'user': {'id': str(result.inserted_id), 'username': username, 'role': role, 'full_name': full_name}
    }), 201


@app.route('/api/users', methods=['GET'])
def get_users():
    if USE_MEMORY_DB:
        users = db.find('users')
    else:
        users = list(db.users.find())
    safe_users = []
    for u in users:
        safe_u = {}
        for k, v in u.items():
            if k == '_id':
                safe_u[k] = str(v)
            elif k != 'password':
                safe_u[k] = v
        safe_users.append(safe_u)
    return jsonify(safe_users)


@app.route('/api/users/<user_id>', methods=['DELETE'])
def delete_user(user_id):
    if USE_MEMORY_DB:
        result = db.delete_one('users', {'_id': user_id})
    else:
        try:
            result = db.users.delete_one({'_id': ObjectId(user_id)})
        except Exception:
            return jsonify({'success': False, 'message': 'Invalid user ID'}), 400
    if result.deleted_count > 0:
        return jsonify({'success': True, 'message': 'User deleted'})
    return jsonify({'success': False, 'message': 'User not found'}), 404


# =============================================================================
# DASHBOARD ROUTES
# =============================================================================

def redirect_to_dashboard(role):
    if role == 'admin':
        return redirect(url_for('admin_dashboard'))
    elif role == 'procurement_manager':
        return redirect(url_for('procurement_dashboard'))
    elif role == 'sales_manager':
        return redirect(url_for('sales_dashboard'))
    elif role == 'sales_person':
        return redirect(url_for('pos'))
    elif role == 'finance_manager':
        return redirect(url_for('finance_dashboard'))
    elif role == 'manual_viewer':
        return redirect(url_for('doc_user_guide'))
    else:
        return redirect(url_for('login_page'))


@app.route('/')
def index():
    if 'user_id' in session:
        return redirect_to_dashboard(session.get('role'))
    return redirect(url_for('login_page'))


@app.route('/admin-dashboard')
def admin_dashboard():
    return render_template('admin_dashboard.html', user=session.get('full_name'), role=session.get('role'))


@app.route('/procurement-dashboard')
def procurement_dashboard():
    return render_template('procurement_dashboard.html', user=session.get('full_name'), role=session.get('role'))


@app.route('/sales-dashboard')
def sales_dashboard():
    if 'user_id' not in session:
        return redirect(url_for('login_page'))
    if session.get('role') not in ['admin', 'sales_manager']:
        return redirect_to_dashboard(session.get('role'))
    return render_template('sales_dashboard.html', user=session.get('full_name'), role=session.get('role'))


@app.route('/finance-dashboard')
def finance_dashboard():
    if 'user_id' not in session:
        return redirect(url_for('login_page'))
    if session.get('role') not in ['admin', 'finance_manager']:
        return redirect_to_dashboard(session.get('role'))
    return render_template('finance_dashboard.html')

@app.route('/finance/coa')
def finance_coa():
    if 'user_id' not in session:
        return redirect(url_for('login_page'))
    if session.get('role') not in ['admin', 'finance_manager']:
        return redirect_to_dashboard(session.get('role'))
    return render_template('finance_coa.html')

@app.route('/finance/gl')
def finance_gl():
    if 'user_id' not in session:
        return redirect(url_for('login_page'))
    if session.get('role') not in ['admin', 'finance_manager']:
        return redirect_to_dashboard(session.get('role'))
    return render_template('finance_gl.html')

@app.route('/finance/reports')
def finance_reports():
    if 'user_id' not in session:
        return redirect(url_for('login_page'))
    if session.get('role') not in ['admin', 'finance_manager']:
        return redirect_to_dashboard(session.get('role'))
    return render_template('finance_reports.html')

@app.route('/finance/journal')
def finance_journal():
    if 'user_id' not in session:
        return redirect(url_for('login_page'))
    if session.get('role') not in ['admin', 'finance_manager']:
        return redirect_to_dashboard(session.get('role'))
    return render_template('finance_journal.html')


@app.route('/pos')
def pos():
    if 'user_id' not in session:
        return redirect(url_for('login_page'))
    if session.get('role') not in ['admin', 'sales_manager', 'sales_person']:
        return redirect_to_dashboard(session.get('role'))
    return render_template('index.html', current_time=datetime.now().strftime('%a, %d %b %Y %H:%M'))


@app.route('/products')
def products_page():
    if 'user_id' not in session:
        return redirect(url_for('login_page'))
    role = session.get('role')
    if role not in ['admin', 'procurement_manager']:
        return redirect_to_dashboard(role)
    return render_template('products.html', user=session.get('full_name'), role=session.get('role'))


# =============================================================================
# PRODUCTS API (Procurement Accessible)
# =============================================================================

@app.route('/api/products', methods=['GET'])
def get_products():
    category = request.args.get('category')
    search = request.args.get('search')
    query = {}
    if category:
        query['category'] = category
    if search:
        query['name'] = {'$regex': search, '$options': 'i'} if not USE_MEMORY_DB else {}

    if USE_MEMORY_DB:
        results = db.find('products', query) if query else db.find('products', {})
        if search and not query:
            results = [p for p in results if search.lower() in p['name'].lower()]
        return jsonify([serialize_doc(p) for p in results])

    products = list(db.products.find(query))
    return jsonify([serialize_doc(p) for p in products])


@app.route('/api/products/<id>', methods=['GET'])
def get_product(id):
    oid = id if USE_MEMORY_DB else ObjectId(id)
    if USE_MEMORY_DB and not any(c.isalnum() for c in id):
        oid = id
    product = db.products.find_one({'_id': oid}) if not USE_MEMORY_DB else db.find_one('products', {'_id': id})
    if product:
        return jsonify(serialize_doc(product))
    return jsonify({'error': 'Product not found'}), 404


@app.route('/api/products', methods=['POST'])
def add_product():
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.json
    required = ['name', 'price', 'category']
    for field in required:
        if field not in data:
            return jsonify({'error': f'Missing required field: {field}'}), 400

    product = {
        'name': data['name'],
        'code': data.get('code', ''),
        'barcode': data.get('barcode', ''),
        'qr_code': data.get('qr_code', secrets.token_hex(8)),
        'unit': data.get('unit', 'piece'),
        'category': data['category'],
        'group_id': data.get('group_id', ''),
        'group_name': data.get('group_name', ''),
        'supplier': data.get('supplier', ''),
        'cost_price': float(data.get('cost_price', 0)),
        'markup_percent': float(data.get('markup_percent', 0)),
        'price': float(data.get('price', 0)),
        'mrp': float(data.get('mrp', 0)),
        'price_includes_tax': data.get('price_includes_tax', False),
        'tax_rate': float(data.get('tax_rate', 0)),
        'price_change_allowed': data.get('price_change_allowed', True),
        'stock': int(data.get('stock', 0)),
        'default_quantity': int(data.get('default_quantity', 1)),
        'is_service': data.get('is_service', False),
        'low_stock_enabled': data.get('low_stock_enabled', False),
        'low_stock_quantity': int(data.get('low_stock_quantity', 0)),
        'reorder_level': int(data.get('reorder_level', 10)),
        'preferred_quantity': int(data.get('preferred_quantity', 0)),
        'age_restriction_years': int(data.get('age_restriction_years', 0)),
        'color': data.get('color', ''),
        'image_url': data.get('image_url', ''),
        'description': data.get('description', ''),
        'status': data.get('status', 'active'),
        'created_at': datetime.utcnow().isoformat(),
        'updated_at': datetime.utcnow().isoformat()
    }

    if USE_MEMORY_DB:
        result = db.insert_one('products', product)
    else:
        result = db.products.insert_one(product)
    product['_id'] = str(result.inserted_id)
    return jsonify(product), 201


@app.route('/api/products/<id>', methods=['PUT'])
def update_product(id):
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.json
    update_data = {}
    editable_fields = [
        'name', 'code', 'barcode', 'unit', 'category', 'group_id', 'group_name',
        'supplier', 'cost_price', 'markup_percent', 'price', 'mrp',
        'price_includes_tax', 'tax_rate', 'price_change_allowed',
        'stock', 'default_quantity', 'is_service', 'low_stock_enabled',
        'low_stock_quantity', 'reorder_level', 'preferred_quantity',
        'age_restriction_years', 'color', 'image_url', 'description', 'status'
    ]
    for field in editable_fields:
        if field in data:
            if field in ['price', 'cost_price', 'markup_percent', 'mrp', 'weight', 'discount_percent']:
                update_data[field] = float(data[field])
            elif field in ['stock', 'default_quantity', 'low_stock_quantity', 'reorder_level', 'preferred_quantity', 'age_restriction_years']:
                update_data[field] = int(data[field])
            elif field in ['price_includes_tax', 'price_change_allowed', 'is_service', 'low_stock_enabled']:
                update_data[field] = bool(data[field])
            else:
                update_data[field] = data[field]
    update_data['updated_at'] = datetime.utcnow().isoformat()

    if not update_data:
        return jsonify({'error': 'No fields to update'}), 400

    if USE_MEMORY_DB:
        result = db.update_one('products', {'_id': id}, {'$set': update_data})
        if result.modified_count > 0:
            product = db.find_one('products', {'_id': id})
            return jsonify(product)
        return jsonify({'error': 'Product not found'}), 404

    result = db.products.update_one({'_id': ObjectId(id)}, {'$set': update_data})
    if result.modified_count > 0:
        product = db.products.find_one({'_id': ObjectId(id)})
        return jsonify(serialize_doc(product))
    return jsonify({'error': 'Product not found'}), 404


@app.route('/api/products/upload-image', methods=['POST'])
def upload_product_image():
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    if 'image' not in request.files:
        return jsonify({'error': 'No image provided'}), 400
    file = request.files['image']
    if not file or file.filename == '':
        return jsonify({'error': 'No file selected'}), 400
    allowed = {'png', 'jpg', 'jpeg', 'gif', 'webp'}
    ext = file.filename.rsplit('.', 1)[-1].lower() if '.' in file.filename else ''
    if ext not in allowed:
        return jsonify({'error': 'Only PNG, JPG, GIF, WEBP allowed'}), 400
    filename = secrets.token_hex(12) + '.' + ext
    upload_dir = os.path.join(app.static_folder, 'uploads', 'products')
    os.makedirs(upload_dir, exist_ok=True)
    file.save(os.path.join(upload_dir, filename))
    url = '/static/uploads/products/' + filename
    return jsonify({'url': url}), 200


@app.route('/api/products/<id>', methods=['DELETE'])
def delete_product(id):
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    comment = request.args.get('comment', 'No comment')
    doc_id = id if USE_MEMORY_DB else ObjectId(id)
    product = db.products.find_one({'_id': doc_id}) if not USE_MEMORY_DB else db.find_one('products', {'_id': id})
    if not product:
        return jsonify({'error': 'Product not found'}), 404

    # Record write-off
    write_off_doc = {
        'product_id': id,
        'product_name': product.get('name', ''),
        'quantity': product.get('stock', 0),
        'reason': comment,
        'written_off_by': session.get('username', 'unknown'),
        'timestamp': datetime.utcnow().isoformat()
    }
    if USE_MEMORY_DB:
        db.write_offs.append(write_off_doc)
        result = db.delete_one('products', {'_id': id})
    else:
        db.write_offs.insert_one(write_off_doc)
        result = db.products.delete_one({'_id': ObjectId(id)})

    if result.deleted_count > 0:
        return jsonify({'message': 'Product deleted', 'write_off_recorded': True})
    return jsonify({'error': 'Product not found'}), 404


# =============================================================================
# VENDORS API
# =============================================================================

@app.route('/api/vendors', methods=['GET'])
def get_vendors():
    if USE_MEMORY_DB:
        vendors = db.find('vendors', {})
        return jsonify([serialize_doc(v) for v in vendors])
    vendors = list(db.vendors.find())
    return jsonify([serialize_doc(v) for v in vendors])


@app.route('/api/vendors/<vendor_id>', methods=['GET'])
def get_vendor(vendor_id):
    if USE_MEMORY_DB:
        vendor = db.find_one('vendors', {'_id': vendor_id})
    else:
        vendor = db.vendors.find_one({'_id': ObjectId(vendor_id)})
    if vendor:
        return jsonify(serialize_doc(vendor))
    return jsonify({'error': 'Vendor not found'}), 404


@app.route('/api/vendors', methods=['POST'])
def add_vendor():
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.json
    required = ['name']
    for field in required:
        if field not in data:
            return jsonify({'error': f'Missing required field: {field}'}), 400

    vendor = {
        'name': data['name'],
        'code': data.get('code', ''),
        'vendor_type': data.get('vendor_type', 'local_supplier'),
        'vendor_tier': data.get('vendor_tier', 'approved'),
        'tax_id': data.get('tax_id', ''),
        'contact_person': data.get('contact_person', ''),
        'phone': data.get('phone', ''),
        'email': data.get('email', ''),
        'whatsapp': data.get('whatsapp', ''),
        'address': data.get('address', ''),
        'city': data.get('city', ''),
        'country': data.get('country', 'Pakistan'),
        'status': data.get('status', 'active'),
        'payment_terms': data.get('payment_terms', ''),
        'category': data.get('category', ''),
        'department': data.get('department', ''),
        'trade_discount': float(data.get('trade_discount', 0)),
        'credit_limit': float(data.get('credit_limit', 0)),
        'bank_name': data.get('bank_name', ''),
        'bank_branch': data.get('bank_branch', ''),
        'account_title': data.get('account_title', ''),
        'account_number': data.get('account_number', ''),
        'notes': data.get('notes', ''),
        'created_at': datetime.utcnow().isoformat()
    }

    if USE_MEMORY_DB:
        result = db.insert_one('vendors', vendor)
    else:
        result = db.vendors.insert_one(vendor)
    vendor['_id'] = str(result.inserted_id)
    return jsonify(vendor), 201


@app.route('/api/vendors/<vendor_id>', methods=['PUT'])
def update_vendor(vendor_id):
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.json
    update_data = {}
    for field in ['name', 'code', 'vendor_type', 'vendor_tier', 'tax_id', 'contact_person',
                  'phone', 'email', 'whatsapp', 'address', 'city', 'country', 'status',
                  'payment_terms', 'category', 'department',
                  'bank_name', 'bank_branch', 'account_title', 'account_number', 'notes']:
        if field in data:
            update_data[field] = data[field]
    for field in ['trade_discount', 'credit_limit']:
        if field in data:
            update_data[field] = float(data[field])

    oid = vendor_id if USE_MEMORY_DB else ObjectId(vendor_id)
    if USE_MEMORY_DB:
        result = db.update_one('vendors', {'_id': vendor_id}, {'$set': update_data})
    else:
        result = db.vendors.update_one({'_id': ObjectId(vendor_id)}, {'$set': update_data})

    if result.modified_count > 0:
        if USE_MEMORY_DB:
            vendor = db.find_one('vendors', {'_id': vendor_id})
        else:
            vendor = db.vendors.find_one({'_id': ObjectId(vendor_id)})
        return jsonify(serialize_doc(vendor))
    return jsonify({'error': 'Vendor not found'}), 404


@app.route('/api/vendors/<vendor_id>', methods=['DELETE'])
def delete_vendor(vendor_id):
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    if USE_MEMORY_DB:
        result = db.delete_one('vendors', {'_id': vendor_id})
    else:
        result = db.vendors.delete_one({'_id': ObjectId(vendor_id)})
    if result.deleted_count > 0:
        return jsonify({'success': True, 'message': 'Vendor deleted'})
    return jsonify({'error': 'Vendor not found'}), 404


# =============================================================================
# PURCHASE ORDERS API
# =============================================================================

@app.route('/api/purchase-orders', methods=['GET'])
def get_purchase_orders():
    if USE_MEMORY_DB:
        orders = getattr(db, 'purchase_orders', [])
        return jsonify([serialize_doc(o) for o in orders])
    orders = list(db.purchase_orders.find().sort('order_date', -1))
    return jsonify([serialize_doc(o) for o in orders])


@app.route('/api/purchase-orders', methods=['POST'])
def create_purchase_order():
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.json
    required = ['vendor_id', 'items']
    for field in required:
        if field not in data:
            return jsonify({'error': f'Missing required field: {field}'}), 400

    total = sum(item.get('quantity', 0) * item.get('unit_price', 0) for item in data['items'])

    order_doc = {
        'vendor_id': data['vendor_id'],
        'vendor_name': data.get('vendor_name', ''),
        'items': data['items'],
        'total_amount': total,
        'paid_amount': data.get('paid_amount', 0),
        'pending_amount': total - data.get('paid_amount', 0),
        'payment_status': data.get('payment_status', 'pending'),
        'order_date': data.get('order_date', datetime.utcnow().isoformat()),
        'expected_delivery': data.get('expected_delivery', ''),
        'notes': data.get('notes', ''),
        'created_by': session.get('username', 'unknown'),
        'created_at': datetime.utcnow().isoformat()
    }

    if USE_MEMORY_DB:
        if not hasattr(db, 'purchase_orders'):
            db.purchase_orders = []
        order_doc['_id'] = str(len(db.purchase_orders) + 1)
        db.purchase_orders.append(order_doc)
        result_id = order_doc['_id']
    else:
        result = db.purchase_orders.insert_one(order_doc)
        result_id = str(result.inserted_id)

    return jsonify({'message': 'Purchase order created', 'order_id': result_id}), 201


# =============================================================================
# PRODUCT GROUPS API
# =============================================================================

@app.route('/api/product-groups', methods=['GET'])
def get_product_groups():
    if USE_MEMORY_DB:
        groups = getattr(db, 'product_groups', [])
        return jsonify([serialize_doc(g) for g in groups])
    groups = list(db.product_groups.find())
    return jsonify([serialize_doc(g) for g in groups])


@app.route('/api/product-groups', methods=['POST'])
def add_product_group():
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.json
    name = data.get('name', '').strip()
    if not name:
        return jsonify({'error': 'Group name is required'}), 400
    group = {
        'name': name,
        'description': data.get('description', ''),
        'created_at': datetime.utcnow().isoformat()
    }
    if USE_MEMORY_DB:
        if not hasattr(db, 'product_groups'):
            db.product_groups = []
        group['_id'] = str(len(db.product_groups) + 1)
        db.product_groups.append(group)
    else:
        result = db.product_groups.insert_one(group)
        group['_id'] = str(result.inserted_id)
    return jsonify(group), 201


@app.route('/api/product-groups/<group_id>', methods=['PUT'])
def update_product_group(group_id):
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.json
    update_data = {}
    for field in ['name', 'description']:
        if field in data:
            update_data[field] = data[field]
    if not update_data:
        return jsonify({'error': 'No fields to update'}), 400
    if USE_MEMORY_DB:
        groups = getattr(db, 'product_groups', [])
        for i, g in enumerate(groups):
            if str(g.get('_id')) == str(group_id):
                groups[i].update(update_data)
                return jsonify(serialize_doc(groups[i]))
        return jsonify({'error': 'Group not found'}), 404
    result = db.product_groups.update_one({'_id': ObjectId(group_id)}, {'$set': update_data})
    if result.modified_count > 0:
        group = db.product_groups.find_one({'_id': ObjectId(group_id)})
        return jsonify(serialize_doc(group))
    return jsonify({'error': 'Group not found'}), 404


@app.route('/api/product-groups/<group_id>', methods=['DELETE'])
def delete_product_group(group_id):
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    if USE_MEMORY_DB:
        groups = getattr(db, 'product_groups', [])
        for i, g in enumerate(groups):
            if str(g.get('_id')) == str(group_id):
                del groups[i]
                return jsonify({'success': True, 'message': 'Group deleted'})
        return jsonify({'error': 'Group not found'}), 404
    result = db.product_groups.delete_one({'_id': ObjectId(group_id)})
    if result.deleted_count > 0:
        return jsonify({'success': True, 'message': 'Group deleted'})
    return jsonify({'error': 'Group not found'}), 404


# =============================================================================
# PRODUCT COMMENTS API
# =============================================================================

@app.route('/api/products/<product_id>/comments', methods=['GET'])
def get_product_comments(product_id):
    if USE_MEMORY_DB:
        comments = getattr(db, 'product_comments', [])
        product_comments = [c for c in comments if str(c.get('product_id')) == str(product_id)]
        return jsonify([serialize_doc(c) for c in product_comments])
    product_comments = list(db.product_comments.find({'product_id': product_id}).sort('created_at', -1))
    return jsonify([serialize_doc(c) for c in product_comments])


@app.route('/api/products/<product_id>/comments', methods=['POST'])
def add_product_comment(product_id):
    if 'user_id' not in session:
        return jsonify({'error': 'Authentication required'}), 401
    data = request.json
    comment_text = data.get('comment', '').strip()
    if not comment_text:
        return jsonify({'error': 'Comment text is required'}), 400
    comment = {
        'product_id': product_id,
        'comment': comment_text,
        'created_by': session.get('username', 'unknown'),
        'created_at': datetime.utcnow().isoformat()
    }
    if USE_MEMORY_DB:
        if not hasattr(db, 'product_comments'):
            db.product_comments = []
        comment['_id'] = str(len(db.product_comments) + 1)
        db.product_comments.append(comment)
    else:
        result = db.product_comments.insert_one(comment)
        comment['_id'] = str(result.inserted_id)
    return jsonify(comment), 201


@app.route('/api/comments/<comment_id>', methods=['DELETE'])
def delete_comment(comment_id):
    if 'user_id' not in session:
        return jsonify({'error': 'Authentication required'}), 401
    if USE_MEMORY_DB:
        comments = getattr(db, 'product_comments', [])
        for i, c in enumerate(comments):
            if str(c.get('_id')) == str(comment_id):
                del comments[i]
                return jsonify({'success': True, 'message': 'Comment deleted'})
        return jsonify({'error': 'Comment not found'}), 404
    result = db.product_comments.delete_one({'_id': ObjectId(comment_id)})
    if result.deleted_count > 0:
        return jsonify({'success': True, 'message': 'Comment deleted'})
    return jsonify({'error': 'Comment not found'}), 404


# =============================================================================
# PRODUCT PRICES API (Date-wise pricing)
# =============================================================================
# PRODUCT PRICES API (Date-wise pricing)
# =============================================================================

@app.route('/api/prices', methods=['GET'])
def get_price_history():
    product_id = request.args.get('product_id')
    if USE_MEMORY_DB:
        prices = db.find_many('product_prices', {'product_id': product_id} if product_id else {})
        return jsonify([serialize_doc(p) for p in prices])

    query = {}
    if product_id:
        query['product_id'] = product_id
    prices = list(db.product_prices.find(query).sort('effective_date', -1))
    return jsonify([serialize_doc(p) for p in prices])


@app.route('/api/prices', methods=['POST'])
def add_product_price():
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.json
    required = ['product_id', 'price', 'effective_date']
    for field in required:
        if field not in data:
            return jsonify({'error': f'Missing required field: {field}'}), 400

    price_doc = {
        'product_id': data['product_id'],
        'product_name': data.get('product_name', ''),
        'price': float(data['price']),
        'old_price': float(data.get('old_price', 0)),
        'effective_date': data['effective_date'],
        'reason': data.get('reason', ''),
        'changed_by': session.get('username', 'unknown'),
        'created_at': datetime.utcnow().isoformat()
    }

    if USE_MEMORY_DB:
        db.product_prices.append(price_doc)
        price_doc['_id'] = str(len(db.product_prices))
    else:
        result = db.product_prices.insert_one(price_doc)
        price_doc['_id'] = str(result.inserted_id)

    return jsonify(price_doc), 201


# =============================================================================
# WRITE-OFFS API
# =============================================================================

@app.route('/api/write-offs', methods=['GET'])
def get_write_offs():
    if USE_MEMORY_DB:
        write_offs = getattr(db, 'write_offs', [])
        return jsonify(write_offs)
    write_offs = list(db.write_offs.find().sort('timestamp', -1))
    return jsonify([serialize_doc(w) for w in write_offs])


# =============================================================================
# UTILITY API
# =============================================================================

@app.route('/api/generate-barcode', methods=['GET'])
def generate_barcode():
    """Generate a unique numeric barcode (EAN-13 format)"""
    import random
    digits = [random.randint(0, 9) for _ in range(12)]
    total = sum(d if i % 2 == 0 else d * 3 for i, d in enumerate(digits))
    check_digit = (10 - (total % 10)) % 10
    barcode = ''.join(str(d) for d in digits) + str(check_digit)
    return jsonify({'barcode': barcode})


@app.route('/api/colors', methods=['GET'])
def get_colors():
    """Return list of colors for product color dropdown"""
    if USE_MEMORY_DB:
        colors = getattr(db, 'colors', [])
        return jsonify([serialize_doc(c) for c in colors])
    colors = list(db.colors.find())
    return jsonify([serialize_doc(c) for c in colors])


@app.route('/api/colors', methods=['POST'])
def add_color():
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.json
    if not data.get('value') or not data.get('label'):
        return jsonify({'error': 'Value and label are required'}), 400
    if USE_MEMORY_DB:
        colors = getattr(db, 'colors', [])
        if any(c['value'] == data['value'] for c in colors):
            return jsonify({'error': 'Color value already exists'}), 400
        new_color = {'_id': str(len(colors) + 1), 'value': data['value'], 'label': data['label']}
        colors.append(new_color)
        return jsonify(serialize_doc(new_color)), 201
    else:
        existing = db.colors.find_one({'value': data['value']})
        if existing:
            return jsonify({'error': 'Color value already exists'}), 400
        result = db.colors.insert_one({'value': data['value'], 'label': data['label']})
        new_color = db.colors.find_one({'_id': result.inserted_id})
        return jsonify(serialize_doc(new_color)), 201


@app.route('/api/colors/<color_id>', methods=['PUT'])
def update_color(color_id):
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.json
    if not data.get('value') or not data.get('label'):
        return jsonify({'error': 'Value and label are required'}), 400
    update_data = {'value': data['value'], 'label': data['label']}
    if USE_MEMORY_DB:
        colors = getattr(db, 'colors', [])
        for i, c in enumerate(colors):
            if str(c.get('_id')) == str(color_id):
                if any(other['value'] == data['value'] and str(other.get('_id')) != str(color_id) for other in colors):
                    return jsonify({'error': 'Color value already exists'}), 400
                colors[i].update(update_data)
                return jsonify(serialize_doc(colors[i]))
        return jsonify({'error': 'Color not found'}), 404
    else:
        existing = db.colors.find_one({'value': data['value'], '_id': {'$ne': ObjectId(color_id)}})
        if existing:
            return jsonify({'error': 'Color value already exists'}), 400
        result = db.colors.update_one({'_id': ObjectId(color_id)}, {'$set': update_data})
        if result.modified_count > 0:
            color = db.colors.find_one({'_id': ObjectId(color_id)})
            return jsonify(serialize_doc(color))
        return jsonify({'error': 'Color not found'}), 404


@app.route('/api/colors/<color_id>', methods=['DELETE'])
def delete_color(color_id):
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    if USE_MEMORY_DB:
        colors = getattr(db, 'colors', [])
        for i, c in enumerate(colors):
            if str(c.get('_id')) == str(color_id):
                del colors[i]
                return jsonify({'success': True, 'message': 'Color deleted'})
        return jsonify({'error': 'Color not found'}), 404
    result = db.colors.delete_one({'_id': ObjectId(color_id)})
    if result.deleted_count > 0:
        return jsonify({'success': True, 'message': 'Color deleted'})
    return jsonify({'error': 'Color not found'}), 404


# =============================================================================
# DEPARTMENTS API
# =============================================================================

@app.route('/api/departments', methods=['GET'])
def get_departments():
    if USE_MEMORY_DB:
        return jsonify([serialize_doc(d) for d in getattr(db, 'departments', [])])
    return jsonify([serialize_doc(d) for d in db.departments.find()])


@app.route('/api/departments', methods=['POST'])
def add_department():
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    name = (request.json or {}).get('name', '').strip()
    if not name:
        return jsonify({'error': 'Name is required'}), 400
    if USE_MEMORY_DB:
        items = getattr(db, 'departments', [])
        if any(d['name'] == name for d in items):
            return jsonify({'error': 'Department already exists'}), 400
        new = {'_id': str(len(items) + 1), 'name': name}
        items.append(new)
        return jsonify(serialize_doc(new)), 201
    if db.departments.find_one({'name': name}):
        return jsonify({'error': 'Department already exists'}), 400
    result = db.departments.insert_one({'name': name})
    return jsonify(serialize_doc(db.departments.find_one({'_id': result.inserted_id}))), 201


@app.route('/api/departments/<item_id>', methods=['PUT'])
def update_department(item_id):
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    name = (request.json or {}).get('name', '').strip()
    if not name:
        return jsonify({'error': 'Name is required'}), 400
    if USE_MEMORY_DB:
        items = getattr(db, 'departments', [])
        for i, d in enumerate(items):
            if str(d.get('_id')) == str(item_id):
                if any(o['name'] == name and str(o.get('_id')) != str(item_id) for o in items):
                    return jsonify({'error': 'Name already exists'}), 400
                items[i]['name'] = name
                return jsonify(serialize_doc(items[i]))
        return jsonify({'error': 'Not found'}), 404
    if db.departments.find_one({'name': name, '_id': {'$ne': ObjectId(item_id)}}):
        return jsonify({'error': 'Name already exists'}), 400
    db.departments.update_one({'_id': ObjectId(item_id)}, {'$set': {'name': name}})
    return jsonify(serialize_doc(db.departments.find_one({'_id': ObjectId(item_id)})))


@app.route('/api/departments/<item_id>', methods=['DELETE'])
def delete_department(item_id):
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    if USE_MEMORY_DB:
        items = getattr(db, 'departments', [])
        for i, d in enumerate(items):
            if str(d.get('_id')) == str(item_id):
                del items[i]
                return jsonify({'success': True})
        return jsonify({'error': 'Not found'}), 404
    result = db.departments.delete_one({'_id': ObjectId(item_id)})
    if result.deleted_count:
        return jsonify({'success': True})
    return jsonify({'error': 'Not found'}), 404


# =============================================================================
# PAYMENT TERMS API
# =============================================================================

@app.route('/api/payment-terms', methods=['GET'])
def get_payment_terms():
    if USE_MEMORY_DB:
        return jsonify([serialize_doc(d) for d in getattr(db, 'payment_terms', [])])
    return jsonify([serialize_doc(d) for d in db.payment_terms.find()])


@app.route('/api/payment-terms', methods=['POST'])
def add_payment_term():
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.json or {}
    name = data.get('name', '').strip()
    desc = data.get('description', '').strip()
    if not name:
        return jsonify({'error': 'Name is required'}), 400
    if USE_MEMORY_DB:
        items = getattr(db, 'payment_terms', [])
        if any(d['name'] == name for d in items):
            return jsonify({'error': 'Payment term already exists'}), 400
        new = {'_id': str(len(items) + 1), 'name': name, 'description': desc}
        items.append(new)
        return jsonify(serialize_doc(new)), 201
    if db.payment_terms.find_one({'name': name}):
        return jsonify({'error': 'Payment term already exists'}), 400
    result = db.payment_terms.insert_one({'name': name, 'description': desc})
    return jsonify(serialize_doc(db.payment_terms.find_one({'_id': result.inserted_id}))), 201


@app.route('/api/payment-terms/<item_id>', methods=['PUT'])
def update_payment_term(item_id):
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.json or {}
    name = data.get('name', '').strip()
    desc = data.get('description', '').strip()
    if not name:
        return jsonify({'error': 'Name is required'}), 400
    if USE_MEMORY_DB:
        items = getattr(db, 'payment_terms', [])
        for i, d in enumerate(items):
            if str(d.get('_id')) == str(item_id):
                if any(o['name'] == name and str(o.get('_id')) != str(item_id) for o in items):
                    return jsonify({'error': 'Name already exists'}), 400
                items[i].update({'name': name, 'description': desc})
                return jsonify(serialize_doc(items[i]))
        return jsonify({'error': 'Not found'}), 404
    if db.payment_terms.find_one({'name': name, '_id': {'$ne': ObjectId(item_id)}}):
        return jsonify({'error': 'Name already exists'}), 400
    db.payment_terms.update_one({'_id': ObjectId(item_id)}, {'$set': {'name': name, 'description': desc}})
    return jsonify(serialize_doc(db.payment_terms.find_one({'_id': ObjectId(item_id)})))


@app.route('/api/payment-terms/<item_id>', methods=['DELETE'])
def delete_payment_term(item_id):
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    if USE_MEMORY_DB:
        items = getattr(db, 'payment_terms', [])
        for i, d in enumerate(items):
            if str(d.get('_id')) == str(item_id):
                del items[i]
                return jsonify({'success': True})
        return jsonify({'error': 'Not found'}), 404
    result = db.payment_terms.delete_one({'_id': ObjectId(item_id)})
    if result.deleted_count:
        return jsonify({'success': True})
    return jsonify({'error': 'Not found'}), 404


# =============================================================================
# UNITS OF MEASURE API
# =============================================================================

@app.route('/api/units', methods=['GET'])
def get_units():
    if USE_MEMORY_DB:
        units = getattr(db, 'units', [])
        return jsonify([serialize_doc(u) for u in units])
    units = list(db.units.find())
    return jsonify([serialize_doc(u) for u in units])


@app.route('/api/units', methods=['POST'])
def add_unit():
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.json
    name = data.get('name', '').strip()
    if not name:
        return jsonify({'error': 'Unit name is required'}), 400
    if USE_MEMORY_DB:
        units = getattr(db, 'units', [])
        if any(u['name'] == name for u in units):
            return jsonify({'error': 'Unit already exists'}), 400
        new_unit = {'_id': str(len(units) + 1), 'name': name}
        units.append(new_unit)
        return jsonify(serialize_doc(new_unit)), 201
    else:
        existing = db.units.find_one({'name': name})
        if existing:
            return jsonify({'error': 'Unit already exists'}), 400
        result = db.units.insert_one({'name': name})
        new_unit = db.units.find_one({'_id': result.inserted_id})
        return jsonify(serialize_doc(new_unit)), 201


@app.route('/api/units/<unit_id>', methods=['PUT'])
def update_unit(unit_id):
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.json
    name = data.get('name', '').strip()
    if not name:
        return jsonify({'error': 'Unit name is required'}), 400
    if USE_MEMORY_DB:
        units = getattr(db, 'units', [])
        for i, u in enumerate(units):
            if str(u.get('_id')) == str(unit_id):
                if any(other['name'] == name and str(other.get('_id')) != str(unit_id) for other in units):
                    return jsonify({'error': 'Unit name already exists'}), 400
                units[i]['name'] = name
                return jsonify(serialize_doc(units[i]))
        return jsonify({'error': 'Unit not found'}), 404
    else:
        existing = db.units.find_one({'name': name, '_id': {'$ne': ObjectId(unit_id)}})
        if existing:
            return jsonify({'error': 'Unit name already exists'}), 400
        result = db.units.update_one({'_id': ObjectId(unit_id)}, {'$set': {'name': name}})
        if result.modified_count > 0:
            unit = db.units.find_one({'_id': ObjectId(unit_id)})
            return jsonify(serialize_doc(unit))
        return jsonify({'error': 'Unit not found'}), 404


@app.route('/api/units/<unit_id>', methods=['DELETE'])
def delete_unit(unit_id):
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    if USE_MEMORY_DB:
        units = getattr(db, 'units', [])
        for i, u in enumerate(units):
            if str(u.get('_id')) == str(unit_id):
                del units[i]
                return jsonify({'success': True, 'message': 'Unit deleted'})
        return jsonify({'error': 'Unit not found'}), 404
    result = db.units.delete_one({'_id': ObjectId(unit_id)})
    if result.deleted_count > 0:
        return jsonify({'success': True, 'message': 'Unit deleted'})
    return jsonify({'error': 'Unit not found'}), 404




# =============================================================================
# REPORTS API
# =============================================================================

@app.route('/api/reports/daily-products', methods=['GET'])
def daily_products_report():
    """Daily products registered report"""
    date_filter = request.args.get('date') or datetime.utcnow().strftime('%Y-%m-%d')
    vendor_filter = request.args.get('vendor')
    category_filter = request.args.get('category')

    if USE_MEMORY_DB:
        products = db.find('products', {})
        if vendor_filter:
            products = [p for p in products if vendor_filter.lower() in (p.get('supplier', '') or '').lower()]
        if category_filter:
            products = [p for p in products if p.get('category') == category_filter]
        return jsonify(products)

    query = {}
    if vendor_filter:
        query['supplier'] = {'$regex': vendor_filter, '$options': 'i'}
    if category_filter:
        query['category'] = category_filter
    products = list(db.products.find(query))
    return jsonify([serialize_doc(p) for p in products])


@app.route('/api/reports/stock-report', methods=['GET'])
def stock_report():
    """Remaining stock report"""
    vendor = request.args.get('vendor')
    product_id = request.args.get('product')
    sort_by_stock = request.args.get('sort', 'stock')

    if USE_MEMORY_DB:
        products = db.find('products', {})
        if vendor:
            products = [p for p in products if vendor.lower() in (p.get('supplier', '') or '').lower()]
        if product_id:
            products = [p for p in products if str(p.get('_id')) == str(product_id)]
        products = sorted(products, key=lambda x: x.get('stock', 0))
        return jsonify([serialize_doc(p) for p in products])

    query = {}
    if vendor:
        query['supplier'] = {'$regex': vendor, '$options': 'i'}
    if product_id:
        query['_id'] = ObjectId(product_id)

    products = list(db.products.find(query).sort('stock', 1))
    return jsonify([serialize_doc(p) for p in products])


@app.route('/api/reports/expiry-report', methods=['GET'])
def expiry_report():
    """Expiry report"""
    vendor = request.args.get('vendor')
    date_from = request.args.get('from')
    date_to = request.args.get('to')

    if USE_MEMORY_DB:
        products = db.find('products', {})
        results = []
        for p in products:
            exp = p.get('expiry_date', '')
            if exp:
                if vendor and vendor.lower() not in (p.get('supplier', '') or '').lower():
                    continue
                if date_from and exp < date_from:
                    continue
                if date_to and exp > date_to:
                    continue
                results.append(p)
        return jsonify(results)

    query = {'expiry_date': {'$exists': True, '$ne': ''}}
    if vendor:
        query['supplier'] = {'$regex': vendor, '$options': 'i'}
    if date_from and date_to:
        query['expiry_date'] = {'$gte': date_from, '$lte': date_to}
    elif date_from:
        query['expiry_date'] = {'$gte': date_from}
    elif date_to:
        query['expiry_date'] = {'$lte': date_to}

    products = list(db.products.find(query))
    return jsonify([serialize_doc(p) for p in products])


@app.route('/api/reports/price-list', methods=['GET'])
def price_list_report():
    """Product price list"""
    vendor = request.args.get('vendor')

    if USE_MEMORY_DB:
        products = db.find('products', {})
        if vendor:
            products = [p for p in products if vendor.lower() in (p.get('supplier', '') or '').lower()]
        result = []
        for p in products:
            result.append({
                'name': p['name'],
                'barcode': p.get('barcode', ''),
                'category': p.get('category', ''),
                'supplier': p.get('supplier', ''),
                'selling_price': p['price'],
                'cost_price': p.get('cost_price', 0),
                'mrp': p.get('mrp', 0),
                'discount_percent': p.get('discount_percent', 0),
                'effective_price': round(p['price'] * (1 - p.get('discount_percent', 0) / 100), 2),
            })
        return jsonify(result)

    query = {}
    if vendor:
        query['supplier'] = {'$regex': vendor, '$options': 'i'}
    products = list(db.products.find(query))
    result = []
    for p in products:
        result.append({
            'name': p['name'],
            'barcode': p.get('barcode', ''),
            'category': p.get('category', ''),
            'supplier': p.get('supplier', ''),
            'selling_price': p['price'],
            'cost_price': p.get('cost_price', 0),
            'mrp': p.get('mrp', 0),
            'discount_percent': p.get('discount_percent', 0),
            'effective_price': round(p['price'] * (1 - p.get('discount_percent', 0) / 100), 2),
        })
    return jsonify(result)


@app.route('/api/reports/slow-moving', methods=['GET'])
def slow_moving_report():
    """Slow moving items report"""
    data = slow_moving_report_data()
    # Serialize for JSON response
    if USE_MEMORY_DB:
        return jsonify(data)
    else:
        # Convert ObjectId in _id field to string for each item
        serialized = []
        for item in data:
            if '_id' in item and not isinstance(item['_id'], str):
                item['_id'] = str(item['_id'])
            serialized.append(item)
        return jsonify(serialized)


@app.route('/api/reports/vendor-payment', methods=['GET'])
def vendor_payment_report():
    """Vendor payment report"""
    vendor_id = request.args.get('vendor_id')
    date_from = request.args.get('from')
    date_to = request.args.get('to')

    if USE_MEMORY_DB:
        orders = getattr(db, 'purchase_orders', [])
        if vendor_id:
            orders = [o for o in orders if o.get('vendor_id') == vendor_id]
        return jsonify(orders)

    query = {}
    if vendor_id:
        query['vendor_id'] = vendor_id
    orders = list(db.purchase_orders.find(query).sort('order_date', -1))
    return jsonify([serialize_doc(o) for o in orders])


# =============================================================================
# EXTENDED REPORTS API
# =============================================================================

@app.route('/api/reports/sales-report', methods=['GET'])
def sales_report_api():
    date_from = request.args.get('from')
    date_to   = request.args.get('to')
    method    = request.args.get('method')
    query = {}
    if date_from:
        try: query.setdefault('timestamp', {})['$gte'] = datetime.fromisoformat(date_from)
        except: pass
    if date_to:
        try: query.setdefault('timestamp', {})['$lte'] = datetime.fromisoformat(date_to + 'T23:59:59')
        except: pass
    if method:
        query['payment_method'] = {'$regex': method, '$options': 'i'}
    sales = list(db.sales.find(query).sort('timestamp', -1).limit(500))
    result = []
    for s in sales:
        ts = s.get('timestamp')
        date_str = ts.strftime('%Y-%m-%d') if isinstance(ts, datetime) else str(ts or '')[:10]
        time_str = ts.strftime('%H:%M') if isinstance(ts, datetime) else ''
        result.append({
            'receipt_no': str(s['_id'])[:8].upper(),
            'date': date_str,
            'time': time_str,
            'cashier': s.get('cashier_name', ''),
            'customer': s.get('customer_name', '') or 'Walk-in',
            'items_count': len(s.get('items', [])),
            'payment_method': s.get('payment_method', ''),
            'tax': round(float(s.get('tax_amount', 0)), 2),
            'discount': round(float(s.get('cart_discount_amt', 0)) + float(s.get('promo_discount', 0)), 2),
            'total': round(float(s.get('total', 0)), 2),
        })
    return jsonify(result)


@app.route('/api/reports/sales-by-product', methods=['GET'])
def sales_by_product_api():
    date_from = request.args.get('from')
    date_to   = request.args.get('to')
    match = {}
    if date_from:
        try: match.setdefault('timestamp', {})['$gte'] = datetime.fromisoformat(date_from)
        except: pass
    if date_to:
        try: match.setdefault('timestamp', {})['$lte'] = datetime.fromisoformat(date_to + 'T23:59:59')
        except: pass
    pipeline = [
        {'$match': match},
        {'$unwind': '$items'},
        {'$group': {
            '_id': '$items.name',
            'qty_sold': {'$sum': '$items.quantity'},
            'revenue': {'$sum': {'$ifNull': ['$items.subtotal', {'$multiply': ['$items.price', '$items.quantity']}]}},
            'cost': {'$sum': {'$multiply': [{'$ifNull': ['$items.cost_price', 0]}, '$items.quantity']}}
        }},
        {'$project': {
            'product': '$_id',
            'qty_sold': 1,
            'revenue': 1,
            'cost': 1,
            'profit': {'$subtract': ['$revenue', '$cost']}
        }},
        {'$sort': {'revenue': -1}}
    ]
    rows = list(db.sales.aggregate(pipeline))
    result = []
    for r in rows:
        rev = float(r.get('revenue') or 0)
        cost = float(r.get('cost') or 0)
        result.append({
            'product': r.get('product', ''),
            'qty_sold': r.get('qty_sold', 0),
            'revenue': round(rev, 2),
            'cost': round(cost, 2),
            'profit': round(rev - cost, 2),
            'margin_pct': round((rev - cost) / rev * 100, 1) if rev > 0 else 0,
        })
    return jsonify(result)


@app.route('/api/reports/stock-valuation', methods=['GET'])
def stock_valuation_api():
    category = request.args.get('category')
    query = {}
    if category:
        query['category'] = category
    products = list(db.products.find(query).sort('name', 1))
    result = []
    for p in products:
        stock = float(p.get('stock', 0))
        cost  = float(p.get('cost_price', 0))
        price = float(p.get('price', 0))
        result.append({
            'product': p.get('name', ''),
            'barcode': p.get('barcode', ''),
            'category': p.get('category', ''),
            'supplier': p.get('supplier', ''),
            'stock': stock,
            'unit': p.get('unit', 'pcs'),
            'cost_price': round(cost, 2),
            'selling_price': round(price, 2),
            'stock_value': round(stock * cost, 2),
            'retail_value': round(stock * price, 2),
        })
    return jsonify(result)


@app.route('/api/reports/reorder-alert', methods=['GET'])
def reorder_alert_api():
    products = list(db.products.find().sort('stock', 1))
    result = []
    for p in products:
        stock   = float(p.get('stock', 0))
        reorder = float(p.get('reorder_level', 0))
        if stock <= max(reorder, 1):
            shortage = max(0, reorder - stock)
            if stock == 0:
                status = 'Out of Stock'
            elif reorder > 0 and stock <= reorder * 0.5:
                status = 'Critical'
            else:
                status = 'Low'
            result.append({
                'product': p.get('name', ''),
                'barcode': p.get('barcode', ''),
                'category': p.get('category', ''),
                'supplier': p.get('supplier', ''),
                'current_stock': stock,
                'reorder_level': reorder,
                'shortage': shortage,
                'status': status,
            })
    return jsonify(result)


@app.route('/api/reports/purchase-orders-report', methods=['GET'])
def purchase_orders_report_api():
    date_from = request.args.get('from')
    date_to   = request.args.get('to')
    status    = request.args.get('status')
    query = {}
    if date_from:
        query.setdefault('created_at', {})['$gte'] = date_from
    if date_to:
        query.setdefault('created_at', {})['$lte'] = date_to + 'T23:59:59'
    if status:
        query['status'] = status
    pos = list(db.purchase_orders_v2.find(query).sort('created_at', -1).limit(500))
    result = []
    for po in pos:
        result.append({
            'po_number': po.get('po_number', str(po['_id'])[:8]),
            'vendor': po.get('vendor_name', ''),
            'created_at': str(po.get('created_at', ''))[:10],
            'expected_delivery': str(po.get('expected_delivery', ''))[:10],
            'items_count': len(po.get('items', [])),
            'total_amount': round(float(po.get('total_amount', 0)), 2),
            'status': po.get('status', ''),
        })
    return jsonify(result)


@app.route('/api/reports/invoice-aging', methods=['GET'])
def invoice_aging_api():
    invoices = list(db.vendor_invoices.find().sort('created_at', -1))
    today = datetime.utcnow()
    result = []
    for inv in invoices:
        created = inv.get('created_at')
        if isinstance(created, str):
            try: created = datetime.fromisoformat(created[:10])
            except: created = today
        elif not isinstance(created, datetime):
            created = today
        age = (today - created).days
        if age <= 30:
            bucket = '0-30 days'
        elif age <= 60:
            bucket = '31-60 days'
        elif age <= 90:
            bucket = '61-90 days'
        else:
            bucket = '90+ days'
        amount = float(inv.get('amount', 0))
        paid   = float(inv.get('paid_amount', 0))
        result.append({
            'invoice_no': inv.get('invoice_number', str(inv['_id'])[:8]),
            'vendor': inv.get('vendor_name', ''),
            'invoice_date': str(inv.get('created_at', ''))[:10],
            'due_date': str(inv.get('due_date', ''))[:10],
            'amount': round(amount, 2),
            'paid': round(paid, 2),
            'outstanding': round(amount - paid, 2),
            'status': inv.get('status', ''),
            'age_days': age,
            'age_bucket': bucket,
        })
    return jsonify(result)


@app.route('/api/reports/profit-loss', methods=['GET'])
def profit_loss_report_api():
    try:
        rev_data = {r['_id']: float(r.get('revenue', 0))
                    for r in db.sales.aggregate([
                        {'$group': {'_id': {'$dateToString': {'format': '%Y-%m', 'date': '$timestamp'}},
                                    'revenue': {'$sum': '$total'}}},
                        {'$sort': {'_id': 1}}
                    ])}
    except:
        rev_data = {}
    try:
        cost_data = {r['_id']: float(r.get('purchases', 0))
                     for r in db.purchase_orders_v2.aggregate([
                         {'$group': {'_id': {'$substr': ['$created_at', 0, 7]},
                                     'purchases': {'$sum': '$total_amount'}}},
                         {'$sort': {'_id': 1}}
                     ])}
    except:
        cost_data = {}
    try:
        payroll_data = {r['_id']: float(r.get('payroll', 0))
                        for r in db.payroll.aggregate([
                            {'$group': {'_id': '$month', 'payroll': {'$sum': '$net_salary'}}},
                            {'$sort': {'_id': 1}}
                        ])}
    except:
        payroll_data = {}
    all_months = sorted(set(list(rev_data.keys()) + list(cost_data.keys()) + list(payroll_data.keys())))
    result = []
    for month in all_months[-24:]:
        revenue  = round(rev_data.get(month, 0), 2)
        purchases = round(cost_data.get(month, 0), 2)
        payroll  = round(payroll_data.get(month, 0), 2)
        expenses = round(purchases + payroll, 2)
        profit   = round(revenue - expenses, 2)
        margin   = round(profit / revenue * 100, 1) if revenue > 0 else 0.0
        result.append({
            'month': month,
            'revenue': revenue,
            'purchases': purchases,
            'payroll': payroll,
            'total_expenses': expenses,
            'net_profit': profit,
            'margin_pct': margin,
        })
    return jsonify(result)


@app.route('/api/reports/sales-returns', methods=['GET'])
def sales_returns_report_api():
    date_from = request.args.get('from')
    date_to   = request.args.get('to')
    query = {}
    if date_from:
        try: query.setdefault('timestamp', {})['$gte'] = datetime.fromisoformat(date_from)
        except: pass
    if date_to:
        try: query.setdefault('timestamp', {})['$lte'] = datetime.fromisoformat(date_to + 'T23:59:59')
        except: pass
    returns = list(db.sales_returns.find(query).sort('timestamp', -1).limit(500))
    result = []
    for r in returns:
        ts = r.get('timestamp')
        result.append({
            'return_id': str(r['_id'])[:8].upper(),
            'original_sale': str(r.get('original_sale_id', ''))[:8],
            'date': ts.strftime('%Y-%m-%d') if isinstance(ts, datetime) else str(ts or '')[:10],
            'cashier': r.get('cashier_name', ''),
            'customer': r.get('customer_name', '') or 'Walk-in',
            'reason': r.get('reason', ''),
            'items_count': len(r.get('items', [])),
            'refund_amount': round(float(r.get('refund_amount', 0)), 2),
        })
    return jsonify(result)


@app.route('/api/reports/payroll-summary', methods=['GET'])
def payroll_summary_report_api():
    month_filter = request.args.get('month')
    query = {}
    if month_filter:
        query['month'] = month_filter
    recs = list(db.payroll.find(query).sort('month', -1).limit(500))
    result = []
    for r in recs:
        result.append({
            'payroll_id': r.get('payroll_id', str(r['_id'])[:8]),
            'employee': r.get('employee_name', ''),
            'department': r.get('department', ''),
            'month': r.get('month', ''),
            'basic_salary': round(float(r.get('basic_salary', 0)), 2),
            'allowances': round(float(r.get('allowances', 0)), 2),
            'deductions': round(float(r.get('deductions', 0)), 2),
            'net_salary': round(float(r.get('net_salary', 0)), 2),
            'status': r.get('status', ''),
        })
    return jsonify(result)


@app.route('/api/reports/credit-sales', methods=['GET'])
def credit_sales_report_api():
    date_from = request.args.get('from')
    date_to   = request.args.get('to')
    query = {'payment_method': {'$regex': 'credit|account', '$options': 'i'}}
    if date_from:
        try: query.setdefault('timestamp', {})['$gte'] = datetime.fromisoformat(date_from)
        except: pass
    if date_to:
        try: query.setdefault('timestamp', {})['$lte'] = datetime.fromisoformat(date_to + 'T23:59:59')
        except: pass
    sales = list(db.sales.find(query).sort('timestamp', -1).limit(500))
    result = []
    for s in sales:
        ts = s.get('timestamp')
        result.append({
            'receipt_no': str(s['_id'])[:8].upper(),
            'date': ts.strftime('%Y-%m-%d') if isinstance(ts, datetime) else str(ts or '')[:10],
            'customer': s.get('customer_name', '') or 'Walk-in',
            'cashier': s.get('cashier_name', ''),
            'items_count': len(s.get('items', [])),
            'payment_method': s.get('payment_method', ''),
            'total': round(float(s.get('total', 0)), 2),
        })
    return jsonify(result)


# =============================================================================
# QR CODE ENDPOINT
# =============================================================================

@app.route('/api/products/qr/<product_id>', methods=['GET'])
def get_product_qr(product_id):
    """Generate QR code info for a product"""
    if USE_MEMORY_DB:
        product = db.find_one('products', {'_id': product_id})
    else:
        product = db.products.find_one({'_id': ObjectId(product_id)})

    if not product:
        return jsonify({'error': 'Product not found'}), 404

    qr_data = {
        'product_id': str(product['_id']),
        'name': product['name'],
        'price': product['price'],
        'barcode': product.get('barcode', ''),
        'qr_code': product.get('qr_code', ''),
        'discount_percent': product.get('discount_percent', 0),
        'discounted_price': round(product['price'] * (1 - product.get('discount_percent', 0) / 100), 2)
    }
    return jsonify(qr_data)


@app.route('/api/products/import-template', methods=['GET'])
def products_import_template():
    """Generate and download the Excel import template for products"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from io import BytesIO

        wb = Workbook()
        ws = wb.active
        ws.title = 'Products'

        headers = [
            ('name', 'Product Name *', True),
            ('barcode', 'Barcode', False),
            ('code', 'Product Code', False),
            ('category', 'Category', False),
            ('unit', 'Unit (kg/piece/liter…)', False),
            ('supplier', 'Supplier / Vendor', False),
            ('cost_price', 'Cost Price (PKR)', False),
            ('markup_percent', 'Markup %', False),
            ('price', 'Sale Price (PKR) *', True),
            ('mrp', 'MRP (Max Retail Price)', False),
            ('discount_percent', 'Discount %', False),
            ('tax_rate', 'Tax Rate %', False),
            ('stock', 'Opening Stock', False),
            ('reorder_level', 'Reorder Level', False),
            ('description', 'Description', False),
            ('status', 'Status (active/inactive)', False),
        ]

        header_fill = PatternFill('solid', fgColor='2980B9')
        req_fill   = PatternFill('solid', fgColor='E74C3C')
        header_font = Font(bold=True, color='FFFFFF', size=11)
        thin = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        for col_idx, (field, label, required) in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.fill = req_fill if required else header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = thin

        # Sample row
        sample = ['Apple', '100001', 'PROD-001', 'Fruits', 'kg', 'Fresh Farms',
                  150.00, 20, 180.00, 200.00, 0, 0, 100, 20, 'Fresh red apples', 'active']
        for col_idx, val in enumerate(sample, 1):
            cell = ws.cell(row=2, column=col_idx, value=val)
            cell.border = thin
            cell.alignment = Alignment(horizontal='center')

        # Second sample row
        sample2 = ['Mineral Water 500ml', '100002', 'PROD-002', 'Beverages', 'piece', 'Nestle Pakistan',
                   25.00, 40, 35.00, 40.00, 0, 0, 200, 50, '500ml mineral water bottle', 'active']
        for col_idx, val in enumerate(sample2, 1):
            cell = ws.cell(row=3, column=col_idx, value=val)
            cell.border = thin
            cell.alignment = Alignment(horizontal='center')

        # Instructions sheet
        ws2 = wb.create_sheet('Instructions')
        ws2['A1'] = 'PRODUCT IMPORT TEMPLATE - INSTRUCTIONS'
        ws2['A1'].font = Font(bold=True, size=13, color='2C3E50')
        instructions = [
            '', 'REQUIRED FIELDS (marked in Red):',
            '  • name — Product name (must be unique)',
            '  • price — Sale price in PKR',
            '',
            'OPTIONAL FIELDS:',
            '  • barcode — EAN/UPC barcode (leave blank to auto-generate)',
            '  • category — Must match an existing category in the system',
            '  • unit — Must match an existing unit (kg, piece, liter, etc.)',
            '  • supplier — Vendor name exactly as registered',
            '  • cost_price — Purchase cost (used for profit margin calculation)',
            '  • markup_percent — If provided, sale price = cost × (1 + markup/100)',
            '  • mrp — Maximum retail price printed on packaging',
            '  • discount_percent — Default discount at POS (0-100)',
            '  • tax_rate — GST/tax percentage (0-100)',
            '  • stock — Opening stock quantity (default 0)',
            '  • reorder_level — Alert when stock falls below this (default 10)',
            '  • status — "active" or "inactive" (default: active)',
            '',
            'NOTES:',
            '  • Delete the sample rows (rows 2-3) before importing',
            '  • Do not change the header row (row 1)',
            '  • Duplicate barcodes will be skipped or merged per your choice',
            '  • Date format: YYYY-MM-DD',
        ]
        for i, text in enumerate(instructions, 2):
            ws2[f'A{i}'] = text

        ws2.column_dimensions['A'].width = 70

        # Column widths in Products sheet
        col_widths = [28, 18, 15, 18, 18, 22, 18, 14, 18, 18, 14, 12, 15, 15, 30, 20]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[1].height = 35

        bio = BytesIO()
        wb.save(bio)
        bio.seek(0)
        return send_file(
            bio,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='product_import_template.xlsx'
        )
    except ImportError:
        # Fallback to CSV template if openpyxl not available
        import csv
        from io import StringIO
        output = StringIO()
        writer = csv.writer(output)
        writer.writerow(['name', 'barcode', 'code', 'category', 'unit', 'supplier',
                         'cost_price', 'markup_percent', 'price', 'mrp', 'discount_percent',
                         'tax_rate', 'stock', 'reorder_level', 'description', 'status'])
        writer.writerow(['Apple', '100001', 'PROD-001', 'Fruits', 'kg', 'Fresh Farms',
                         150, 20, 180, 200, 0, 0, 100, 20, 'Fresh red apples', 'active'])
        bio = BytesIO(output.getvalue().encode('utf-8-sig'))
        bio.seek(0)
        return send_file(bio, mimetype='text/csv', as_attachment=True,
                         download_name='product_import_template.csv')


# =============================================================================
# CATEGORIES API
# =============================================================================

@app.route('/api/categories', methods=['GET'])
def get_categories():
    if USE_MEMORY_DB:
        cats = db.find('categories', {})
        return jsonify([{'_id': c.get('_id'), 'name': c['name']} for c in cats])
    categories = list(db.categories.find())
    return jsonify([serialize_doc(c) for c in categories])


@app.route('/api/categories', methods=['POST'])
def add_category():
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.json
    if 'name' not in data or not data['name'].strip():
        return jsonify({'error': 'Category name required'}), 400
    name = data['name'].strip()

    if USE_MEMORY_DB:
        categories = getattr(db, 'categories', [])
        if any(c['name'] == name for c in categories):
            return jsonify({'error': 'Category already exists'}), 400
        result = db.insert_one('categories', {'name': name})
    else:
        existing = db.categories.find_one({'name': name})
        if existing:
            return jsonify({'error': 'Category already exists'}), 400
        result = db.categories.insert_one({'name': name})
    return jsonify({'_id': str(result.inserted_id), 'name': name}), 201

@app.route('/api/categories/<category_id>', methods=['PUT'])
def update_category(category_id):
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.json
    if 'name' not in data or not data['name'].strip():
        return jsonify({'error': 'Category name is required'}), 400
    new_name = data['name'].strip()
    if USE_MEMORY_DB:
        categories = getattr(db, 'categories', [])
        for i, cat in enumerate(categories):
            if str(cat.get('_id')) == str(category_id):
                if any(c['name'] == new_name and str(c.get('_id')) != str(category_id) for c in categories):
                    return jsonify({'error': 'Category name already exists'}), 400
                categories[i]['name'] = new_name
                return jsonify(categories[i])
        return jsonify({'error': 'Category not found'}), 404
    else:
        existing = db.categories.find_one({'name': new_name, '_id': {'$ne': ObjectId(category_id)}})
        if existing:
            return jsonify({'error': 'Category name already exists'}), 400
        result = db.categories.update_one({'_id': ObjectId(category_id)}, {'$set': {'name': new_name}})
        if result.modified_count > 0:
            cat = db.categories.find_one({'_id': ObjectId(category_id)})
            return jsonify(serialize_doc(cat))
        return jsonify({'error': 'Category not found'}), 404

@app.route('/api/categories/<category_id>', methods=['DELETE'])
def delete_category(category_id):
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    if USE_MEMORY_DB:
        categories = getattr(db, 'categories', [])
        for i, cat in enumerate(categories):
            if str(cat.get('_id')) == str(category_id):
                del categories[i]
                return jsonify({'success': True, 'message': 'Category deleted'})
        return jsonify({'error': 'Category not found'}), 404
    result = db.categories.delete_one({'_id': ObjectId(category_id)})
    if result.deleted_count > 0:
        return jsonify({'success': True, 'message': 'Category deleted'})
    return jsonify({'error': 'Category not found'}), 404

# =============================================================================
# SALES API
# =============================================================================

@app.route('/api/sales', methods=['POST'])
def create_sale():
    if session.get('role') not in ['admin', 'procurement_manager', 'sales_person', 'cashier']:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.json or {}
    if 'items' not in data or not data['items']:
        return jsonify({'error': 'No items in cart'}), 400

    # ── Fetch products ───────────────────────────────────────────────────────
    product_ids = [str(item['product_id']) for item in data['items']]
    if USE_MEMORY_DB:
        prods_map = {p['_id']: p for p in db.find('products', {}) if p['_id'] in product_ids}
    else:
        prods_map = {str(p['_id']): p for p in db.products.find({'_id': {'$in': [ObjectId(pid) for pid in product_ids]}})}

    # ── Build sale items with per-item discounts ─────────────────────────────
    sale_items = []
    items_subtotal = 0.0
    for item in data['items']:
        pid = str(item['product_id'])
        if pid not in prods_map: continue
        product  = prods_map[pid]
        qty      = int(item.get('quantity', 1))
        price    = float(product['price'])
        # item-level discount (PKR or %)
        item_disc_pct = float(item.get('discount_pct', 0))
        item_disc_amt = float(item.get('discount_amt', 0))
        if item_disc_pct > 0:
            item_disc_amt = round(price * qty * item_disc_pct / 100, 2)
        subtotal = round(price * qty - item_disc_amt, 2)
        items_subtotal += subtotal
        sale_items.append({
            'product_id': pid, 'product_name': product['name'],
            'price': price, 'quantity': qty,
            'discount_pct': item_disc_pct, 'discount_amt': item_disc_amt,
            'subtotal': subtotal,
        })
        # Deduct stock
        if USE_MEMORY_DB:
            p = db.find_one('products', {'_id': pid})
            if p: p['stock'] = p.get('stock', 0) - qty
        else:
            db.products.update_one({'_id': ObjectId(pid)}, {'$inc': {'stock': -qty}})

    # ── Cart-level discount ──────────────────────────────────────────────────
    cart_disc_pct = float(data.get('cart_discount_pct', 0))
    cart_disc_amt = float(data.get('cart_discount_amt', 0))
    if cart_disc_pct > 0:
        cart_disc_amt = round(items_subtotal * cart_disc_pct / 100, 2)

    # ── Promo code discount ──────────────────────────────────────────────────
    promo_code     = (data.get('promo_code','') or '').strip().upper()
    promo_discount = float(data.get('promo_discount', 0))

    # ── Loyalty points redemption ────────────────────────────────────────────
    loyalty_used = int(data.get('loyalty_points_used', 0))
    loyalty_value = 0.0
    cfg = {}
    if not USE_MEMORY_DB:
        cfg_doc = db.store_config.find_one({})
        cfg = cfg_doc or {}
    redeem_rate = float(cfg.get('loyalty_redeem_rate', 0.5))
    if loyalty_used > 0:
        loyalty_value = round(loyalty_used * redeem_rate, 2)

    # ── Tax ──────────────────────────────────────────────────────────────────
    tax_rate = float(data.get('tax_rate', cfg.get('tax_rate', 0)))
    pre_tax  = max(0, items_subtotal - cart_disc_amt - promo_discount - loyalty_value)
    tax_amt  = round(pre_tax * tax_rate / 100, 2)
    total    = round(pre_tax + tax_amt, 2)

    # ── Split payment ────────────────────────────────────────────────────────
    payments = data.get('payments', [])  # [{'method':'cash','amount':500},...]
    if not payments:
        payments = [{'method': data.get('payment_method','cash'), 'amount': total}]
    payments_total = sum(float(p.get('amount',0)) for p in payments)
    cash_tendered  = sum(float(p.get('amount',0)) for p in payments if p.get('method')=='cash')
    change_given   = max(0, round(cash_tendered - (total - (payments_total - cash_tendered)), 2))
    primary_method = payments[0]['method'] if len(payments)==1 else 'split'

    # ── Customer ──────────────────────────────────────────────────────────────
    customer_id   = data.get('customer_id', '')
    customer_name = data.get('customer_name', '')
    loyalty_rate  = float(cfg.get('loyalty_rate', 1))
    loyalty_earned = int(total * loyalty_rate / 10)  # 1 pt per PKR 10

    # ── Shift ─────────────────────────────────────────────────────────────────
    shift_id = data.get('shift_id', '')

    # ── Build sale document ──────────────────────────────────────────────────
    sale_doc = {
        'items': sale_items, 'items_subtotal': items_subtotal,
        'cart_discount_pct': cart_disc_pct, 'cart_discount_amt': cart_disc_amt,
        'promo_code': promo_code, 'promo_discount': promo_discount,
        'loyalty_points_used': loyalty_used, 'loyalty_value': loyalty_value,
        'tax_rate': tax_rate, 'tax_amount': tax_amt,
        'total': total, 'payments': payments, 'payment_method': primary_method,
        'cash_tendered': cash_tendered, 'change_given': change_given,
        'customer_id': customer_id, 'customer_name': customer_name,
        'loyalty_points_earned': loyalty_earned,
        'shift_id': shift_id, 'cashier_id': session.get('user_id',''),
        'cashier_name': session.get('full_name',''), 'timestamp': datetime.utcnow(),
    }

    if USE_MEMORY_DB:
        sale_doc['_id'] = get_next_id('sales')
        db.sales.append(sale_doc); result_id = str(sale_doc['_id'])
    else:
        result  = db.sales.insert_one(sale_doc); result_id = str(result.inserted_id)

        # GL Auto-Post
        if total > 0:
            gl_lines = [{'account_code':'1000','account_name':'Cash & Cash Equivalents','dr':total,'cr':0},
                        {'account_code':'4000','account_name':'Sales Revenue','dr':0,'cr':total}]
            total_cost = sum(float(i.get('cost_price',0))*i.get('quantity',1) for i in sale_items)
            if total_cost > 0:
                gl_lines += [{'account_code':'5000','account_name':'Cost of Goods Sold','dr':total_cost,'cr':0},
                             {'account_code':'1200','account_name':'Inventory','dr':0,'cr':total_cost}]
            post_journal_entry(description=f"Sale — {result_id[:8]}", lines=gl_lines,
                               reference_type='sale', reference_id=result_id, reference_number=result_id[:8])

        # Update customer loyalty
        if customer_id:
            try:
                db.customers.update_one({'_id': ObjectId(customer_id)}, {
                    '$inc': {'loyalty_points': loyalty_earned - loyalty_used,
                             'total_spent': total, 'visit_count': 1}
                })
            except: pass

        # Update promo code usage
        if promo_code and promo_discount > 0:
            db.promo_codes.update_one({'code': promo_code}, {'$inc': {'uses': 1}})

        # Update shift totals
        if shift_id:
            pb_inc = {f'payment_breakdown.{p["method"]}': float(p["amount"]) for p in payments}
            try:
                db.shifts.update_one({'_id': ObjectId(shift_id)}, {
                    '$inc': {'sales_count': 1, 'sales_total': total,
                             'discount_total': cart_disc_amt + promo_discount, **pb_inc}
                })
            except: pass

    resp = sale_doc.copy(); resp['_id'] = result_id
    if isinstance(resp.get('timestamp'), datetime): resp['timestamp'] = resp['timestamp'].isoformat()
    return jsonify({'message': 'Sale completed', 'sale': resp, 'receipt_number': result_id[:8]}), 201


@app.route('/api/sales', methods=['GET'])
def get_sales():
    if session.get('role') not in ['admin', 'procurement_manager', 'sales_person', 'cashier']:
        return jsonify({'error': 'Unauthorized'}), 403
    limit = int(request.args.get('limit', 50))
    if USE_MEMORY_DB:
        sales = sorted(getattr(db, 'sales', []), key=lambda x: x.get('timestamp', ''), reverse=True)[:limit]
        for s in sales:
            if '_id' in s and not isinstance(s['_id'], str):
                s['_id'] = str(s['_id'])
            if isinstance(s.get('timestamp'), datetime):
                s['timestamp'] = s['timestamp'].isoformat()
        return jsonify(sales)
    sales = list(db.sales.find().sort('timestamp', -1).limit(limit))
    for s in sales:
        s['_id'] = str(s['_id'])
        if isinstance(s.get('timestamp'), datetime):
            s['timestamp'] = s['timestamp'].isoformat()
    return jsonify(sales)


@app.route('/api/stats/today', methods=['GET'])
def today_stats():
    if USE_MEMORY_DB:
        today_start = datetime.combine(datetime.utcnow().date(), datetime.min.time())
        sales = getattr(db, 'sales', [])
        today_sales = [s for s in sales if isinstance(s.get('timestamp'), datetime) and s['timestamp'] >= today_start]
        total_sales = sum(s['total'] for s in today_sales)
        transactions = len(today_sales)
        product_count = db.count_documents('products', {})
        low_stock = len([p for p in db.find('products', {}) if p.get('stock', 0) < (p.get('reorder_level') or 10)])
        return jsonify({
            'today_sales': total_sales,
            'today_transactions': transactions,
            'total_products': product_count,
            'low_stock_products': low_stock
        })
    today_start = datetime.combine(datetime.utcnow().date(), datetime.min.time())
    today_end = datetime.combine(datetime.utcnow().date(), datetime.max.time())
    pipeline = [
        {'$match': {'timestamp': {'$gte': today_start, '$lte': today_end}}},
        {'$group': {
            '_id': None,
            'total_sales': {'$sum': '$total'},
            'total_transactions': {'$sum': 1}
        }}
    ]
    summary = list(db.sales.aggregate(pipeline))
    product_count = db.products.count_documents({})
    low_stock = db.products.count_documents({
        '$expr': {
            '$lt': ['$stock', {'$ifNull': ['$reorder_level', 10]}]
        }
    })
    if summary:
        return jsonify({
            'today_sales': summary[0]['total_sales'],
            'today_transactions': summary[0]['total_transactions'],
            'total_products': product_count,
            'low_stock_products': low_stock
        })
    return jsonify({
        'today_sales': 0,
        'today_transactions': 0,
        'total_products': product_count,
        'low_stock_products': low_stock
    })


@app.route('/api/sales/by-category', methods=['GET'])
def sales_by_category():
    if USE_MEMORY_DB:
        product_sales = {}
        for sale in getattr(db, 'sales', []):
            for item in sale.get('items', []):
                pid = item.get('product_id')
                prod = db.find_one('products', {'_id': pid})
                name = (prod.get('name') if prod else 'Unknown')
                if name not in product_sales:
                    product_sales[name] = {'quantity_sold': 0, 'revenue': 0}
                product_sales[name]['quantity_sold'] += item.get('quantity', 0)
                product_sales[name]['revenue'] += item.get('subtotal', 0)
        result = [{'name': k, 'quantity_sold': v['quantity_sold'], 'revenue': v['revenue']} for k, v in product_sales.items()]
        result.sort(key=lambda x: x['revenue'], reverse=True)
        return jsonify(result[:20])
    pipeline = [
        {'$unwind': '$items'},
        {'$group': {
            '_id': '$items.product_name',
            'quantity_sold': {'$sum': '$items.quantity'},
            'revenue': {'$sum': '$items.subtotal'}
        }},
        {'$sort': {'revenue': -1}},
        {'$limit': 20}
    ]
    results = list(db.sales.aggregate(pipeline))
    for r in results:
        r['name'] = r.pop('_id')
    return jsonify(results)


# =============================================================================
# UPLOAD PRODUCTS (Existing)
# =============================================================================

@app.route('/api/upload/products', methods=['POST'])
def upload_products():
    """Upload products from CSV/JSON file"""
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403

    if 'file' not in request.files:
        return jsonify({'error': 'No file provided'}), 400

    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No file selected'}), 400

    try:
        content = file.read().decode('utf-8')
        file_format = file.filename.rsplit('.', 1)[-1].lower()

        if file_format == 'csv':
            import csv
            from io import StringIO
            reader = csv.DictReader(StringIO(content))
            products_data = list(reader)
        elif file_format in ('json', 'txt'):
            products_data = json.loads(content)
            if isinstance(products_data, dict) and 'products' in products_data:
                products_data = products_data['products']
        else:
            return jsonify({'error': 'Unsupported file format. Use CSV or JSON'}), 400

        inserted = 0
        errors = []
        for i, item in enumerate(products_data):
            try:
                product = {
                    'name': item.get('name', f'Product {i+1}'),
                    'price': float(item.get('price', 0)),
                    'category': item.get('category', 'General'),
                    'barcode': item.get('barcode', ''),
                    'stock': int(item.get('stock', 0)),
                    'unit': item.get('unit', 'piece'),
                    'supplier': item.get('supplier', ''),
                    'cost_price': float(item.get('cost_price', 0)),
                    'mrp': float(item.get('mrp', 0)),
                    'discount_percent': float(item.get('discount_percent', 0)),
                    'reorder_level': int(item.get('reorder_level', 10)),
                    'tax_rate': float(item.get('tax_rate', 0)),
                    'hsn_code': item.get('hsn_code', ''),
                    'expiry_date': item.get('expiry_date', ''),
                    'batch_no': item.get('batch_no', ''),
                    'manufacturer': item.get('manufacturer', ''),
                    'qr_code': item.get('qr_code', secrets.token_hex(8)),
                    'created_at': datetime.utcnow().isoformat()
                }
                if USE_MEMORY_DB:
                    db.insert_one('products', product)
                else:
                    db.products.insert_one(product)
                inserted += 1
            except Exception as e:
                errors.append(f'Row {i+1}: {str(e)}')

        return jsonify({
            'message': f'Upload complete',
            'inserted': inserted,
            'errors': len(errors),
            'error_details': errors[:10]
        })
    except Exception as e:
        return jsonify({'error': f'Failed to process file: {str(e)}'}), 400


@app.route('/api/upload/excel-products', methods=['POST'])
def upload_excel_products():
    """Upload products from Excel file with duplicate/merge options"""
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Unauthorized'}), 403

    if 'file' not in request.files:
        return jsonify({'error': 'No file provided'}), 400

    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No file selected'}), 400

    # Get options from form data
    skip_duplicates = request.form.get('skip_duplicates', 'true').lower() == 'true'
    merge_duplicates = request.form.get('merge_duplicates', 'true').lower() == 'true'
    upload_type = request.form.get('upload_type', 'inventory')  # inventory or purchase
    default_category = request.form.get('category', '').strip()
    default_group = request.form.get('product_group', '').strip()
    default_quantity = request.form.get('default_quantity', '').strip()
    default_expiry = request.form.get('default_expiry', '').strip()
    default_color = request.form.get('color', '').strip()

    try:
        from openpyxl import load_workbook
        from io import BytesIO

        wb = load_workbook(filename=BytesIO(file.read()), read_only=True)
        ws = wb.active

        products_data = []
        headers = None

        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                headers = [str(h).strip() if h else f'col_{i}' for i, h in enumerate(row)]
                continue
            if not any(row):
                continue
            row_dict = {}
            for idx, val in enumerate(row):
                if idx < len(headers):
                    row_dict[headers[idx]] = val
            products_data.append(row_dict)

        wb.close()

        inserted = 0
        updated = 0
        skipped = 0
        errors = []

        for i, item in enumerate(products_data):
            try:
                product_name = str(item.get('Name', f'Product {i+1}')).strip()
                if not product_name or product_name == 'None':
                    continue

                cost = float(item.get('Cost') or item.get('cost') or 0)
                markup = float(item.get('Markup') or item.get('markup') or 0)
                price = float(item.get('Price') or item.get('price') or (cost * (1 + markup / 100)) or 0)
                mrp = float(item.get('Price') or item.get('mrp') or price)

                barcode_val = item.get('Barcode') or item.get('barcode') or ''
                barcode_str = str(barcode_val) if barcode_val else ''

                product = {
                    'name': product_name,
                    'code': str(item.get('SKU') or item.get('code') or ''),
                    'barcode': barcode_str,
                    'qr_code': secrets.token_hex(8),
                    'unit': str(item.get('Measurement unit') or item.get('unit') or 'piece'),
                    'category': str(item.get('Product group') or item.get('category') or default_category or 'General'),
                    'group_name': str(item.get('Product group') or default_group or ''),
                    'supplier': str(item.get('Supplier') or item.get('supplier') or ''),
                    'cost_price': cost,
                    'markup_percent': markup,
                    'price': price,
                    'mrp': mrp,
                    'price_includes_tax': str(item.get('Tax inclusive price') or '').lower() in ('true', '1', 'yes') if item.get('Tax inclusive price') else False,
                    'tax_rate': float(item.get('Tax') or item.get('tax_rate') or 0),
                    'price_change_allowed': str(item.get('Price change allowed') or '').lower() in ('true', '1', 'yes') if item.get('Price change allowed') else True,
                    'stock': int(item.get('Quantity') or item.get('stock') or (int(default_quantity) if default_quantity else 0)),
                    'default_quantity': int(item.get('Using default quantity') or item.get('default_quantity') or 1),
                    'expiry_date': str(item.get('Expiry date') or item.get('expiry_date') or default_expiry or ''),
                    'color': str(item.get('Color') or item.get('color') or default_color or ''),
                    'is_service': str(item.get('Service (not using stock)') or '').lower() in ('true', '1', 'yes') if item.get('Service (not using stock)') else False,
                    'enabled': str(item.get('Enabled') or '').lower() in ('true', '1', 'yes') if item.get('Enabled') else True,
                    'low_stock_enabled': str(item.get('Low stock warning') or '').lower() in ('true', '1', 'yes') if item.get('Low stock warning') else False,
                    'low_stock_quantity': int(item.get('Low stock warning quantity') or item.get('low_stock_quantity') or 0),
                    'reorder_level': int(item.get('Reorder point') or item.get('reorder_level') or 10),
                    'preferred_quantity': int(item.get('Preferred quantity') or item.get('preferred_quantity') or 0),
                    'description': str(item.get('Description') or ''),
                    'status': 'active' if item.get('Enabled') is None or str(item.get('Enabled')).lower() in ('true', '1', 'yes') else 'inactive',
                    'created_at': datetime.utcnow().isoformat(),
                    'updated_at': datetime.utcnow().isoformat()
                }

                existing = None
                if USE_MEMORY_DB:
                    for p in db.find('products', {}):
                        if p.get('barcode') and str(p.get('barcode')) == barcode_str:
                            existing = p
                            break
                        if str(p.get('name', '')).lower() == product_name.lower() and not barcode_str:
                            existing = p
                            break
                else:
                    if barcode_str:
                        existing = db.products.find_one({'barcode': barcode_str})
                    if not existing:
                        existing = db.products.find_one({'name': {'$regex': f'^{product_name}$', '$options': 'i'}})

                if existing:
                    if skip_duplicates and not merge_duplicates:
                        skipped += 1
                        continue
                    elif merge_duplicates:
                        update_data = {k: v for k, v in product.items() if k not in ('_id', 'created_at')}
                        update_data['updated_at'] = datetime.utcnow().isoformat()
                        if USE_MEMORY_DB:
                            db.update_one('products', {'_id': existing['_id']}, {'$set': update_data})
                        else:
                            db.products.update_one({'_id': ObjectId(existing['_id'])}, {'$set': update_data})
                        updated += 1
                        continue

                if USE_MEMORY_DB:
                    db.insert_one('products', product)
                else:
                    db.products.insert_one(product)
                inserted += 1

            except Exception as e:
                errors.append(f'Row {i+2}: {str(e)}')

        type_msg = 'Inventory' if upload_type == 'inventory' else 'Purchase'
        return jsonify({
            'message': f'{type_msg} upload complete',
            'inserted': inserted,
            'updated': updated,
            'skipped': skipped,
            'errors': len(errors),
            'error_details': errors[:10]
        })

    except Exception as e:
        return jsonify({'error': f'Failed to process Excel file: {str(e)}'}), 400


# =============================================================================
# LEGACY ROUTES (backward compatible)
# =============================================================================

@app.route('/dashboard')
def dashboard_legacy():
    return render_template('dashboard.html')


# =============================================================================
# REPORT VIEWER & EXPORT
# =============================================================================

@app.route('/reports/<report_type>')
def report_viewer(report_type):
    """HTML report viewer page"""
    if report_type == 'z-report':
        return redirect(url_for('z_report_page'))
    valid_reports = [
        'stock-report', 'expiry-report', 'price-list',
        'slow-moving', 'vendor-payment', 'daily-products',
        'sales-report', 'sales-by-product', 'stock-valuation',
        'reorder-alert', 'purchase-orders-report', 'invoice-aging',
        'profit-loss', 'sales-returns', 'payroll-summary', 'credit-sales',
    ]
    if report_type not in valid_reports:
        return "Invalid report type", 404
    return render_template('report_viewer.html', report_type=report_type)


def dicts_to_csv(data, headers):
    """Convert list of dicts to CSV string"""
    import io
    import csv
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=headers, extrasaction='ignore')
    writer.writeheader()
    for row in data:
        writer.writerow(row)
    return output.getvalue()


@app.route('/api/reports/stock-report/excel')
def stock_report_excel():
    data = stock_report_data()  # Will call existing function after defined
    headers = ['name', 'category', 'supplier', 'stock', 'unit', 'price']
    csv_data = dicts_to_csv(data, headers)
    return excel_response(csv_data, 'stock_report')


@app.route('/api/reports/expiry-report/excel')
def expiry_report_excel():
    data = expiry_report_data()
    headers = ['name', 'batch_no', 'expiry_date', 'stock']
    csv_data = dicts_to_csv(data, headers)
    return excel_response(csv_data, 'expiry_report')


@app.route('/api/reports/price-list/excel')
def price_list_excel():
    data = price_list_report_data()
    headers = ['name', 'barcode', 'category', 'supplier', 'selling_price', 'cost_price', 'mrp', 'discount_percent']
    csv_data = dicts_to_csv(data, headers)
    return excel_response(csv_data, 'price_list')


@app.route('/api/reports/slow-moving/excel')
def slow_moving_excel():
    data = slow_moving_report_data()
    headers = ['name', 'category', 'stock', 'last_sold', 'days_since']
    csv_data = dicts_to_csv(data, headers)
    return excel_response(csv_data, 'slow_moving')


@app.route('/api/reports/vendor-payment/excel')
def vendor_payment_excel():
    data = vendor_payment_report_data()
    headers = ['_id', 'vendor_name', 'order_date', 'total_amount', 'paid_amount', 'pending_amount', 'payment_status']
    csv_data = dicts_to_csv(data, headers)
    return excel_response(csv_data, 'vendor_payment')


@app.route('/api/reports/daily-products/excel')
def daily_products_excel():
    data = daily_products_report_data()
    headers = ['name', 'category', 'supplier', 'stock']
    csv_data = dicts_to_csv(data, headers)
    return excel_response(csv_data, 'daily_products')


def _xlsx_response(rows, col_headers, sheet_name, filename):
    """Generate a real .xlsx file using openpyxl and return as download."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from io import BytesIO
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    header_fill = PatternFill('solid', fgColor='1e3a5f')
    header_font = Font(bold=True, color='FFFFFF', size=11)
    border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC'),
    )
    alt_fill = PatternFill('solid', fgColor='F2F6FC')
    # Header row
    for ci, hdr in enumerate(col_headers, 1):
        cell = ws.cell(row=1, column=ci, value=hdr)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    ws.row_dimensions[1].height = 22
    # Data rows
    for ri, row in enumerate(rows, 2):
        fill = alt_fill if ri % 2 == 0 else PatternFill()
        for ci, key in enumerate(row.keys(), 1):
            cell = ws.cell(row=ri, column=ci, value=row[key])
            cell.border = border
            cell.fill = fill
            cell.alignment = Alignment(vertical='center')
            if isinstance(row[key], float):
                cell.number_format = '#,##0.00'
    # Auto column width
    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    from flask import send_file
    return send_file(
        bio,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'{filename}_{datetime.utcnow().strftime("%Y-%m-%d")}.xlsx'
    )


@app.route('/api/reports/sales-report/excel')
def sales_report_excel():
    with app.test_request_context(request.url):
        pass
    sales = list(db.sales.find().sort('timestamp', -1).limit(1000))
    rows = []
    for s in sales:
        ts = s.get('timestamp')
        rows.append({
            'Receipt No': str(s['_id'])[:8].upper(),
            'Date': ts.strftime('%Y-%m-%d') if isinstance(ts, datetime) else str(ts or '')[:10],
            'Time': ts.strftime('%H:%M') if isinstance(ts, datetime) else '',
            'Cashier': s.get('cashier_name', ''),
            'Customer': s.get('customer_name', '') or 'Walk-in',
            'Items': len(s.get('items', [])),
            'Payment Method': s.get('payment_method', ''),
            'Tax (PKR)': round(float(s.get('tax_amount', 0)), 2),
            'Discount (PKR)': round(float(s.get('cart_discount_amt', 0)) + float(s.get('promo_discount', 0)), 2),
            'Total (PKR)': round(float(s.get('total', 0)), 2),
        })
    return _xlsx_response(rows, list(rows[0].keys()) if rows else [], 'Sales Report', 'sales_report')


@app.route('/api/reports/sales-by-product/excel')
def sales_by_product_excel():
    pipeline = [
        {'$unwind': '$items'},
        {'$group': {
            '_id': '$items.name',
            'qty_sold': {'$sum': '$items.quantity'},
            'revenue': {'$sum': {'$ifNull': ['$items.subtotal', {'$multiply': ['$items.price', '$items.quantity']}]}},
            'cost': {'$sum': {'$multiply': [{'$ifNull': ['$items.cost_price', 0]}, '$items.quantity']}}
        }},
        {'$sort': {'revenue': -1}}
    ]
    data = list(db.sales.aggregate(pipeline))
    rows = []
    for r in data:
        rev = float(r.get('revenue') or 0)
        cost = float(r.get('cost') or 0)
        rows.append({
            'Product': r.get('_id', ''),
            'Qty Sold': r.get('qty_sold', 0),
            'Revenue (PKR)': round(rev, 2),
            'Cost (PKR)': round(cost, 2),
            'Profit (PKR)': round(rev - cost, 2),
            'Margin %': round((rev - cost) / rev * 100, 1) if rev > 0 else 0,
        })
    return _xlsx_response(rows, list(rows[0].keys()) if rows else [], 'Sales by Product', 'sales_by_product')


@app.route('/api/reports/stock-valuation/excel')
def stock_valuation_excel():
    products = list(db.products.find().sort('name', 1))
    rows = []
    for p in products:
        stock = float(p.get('stock', 0))
        cost  = float(p.get('cost_price', 0))
        price = float(p.get('price', 0))
        rows.append({
            'Product': p.get('name', ''),
            'Barcode': p.get('barcode', ''),
            'Category': p.get('category', ''),
            'Supplier': p.get('supplier', ''),
            'Stock': stock,
            'Unit': p.get('unit', 'pcs'),
            'Cost Price (PKR)': round(cost, 2),
            'Selling Price (PKR)': round(price, 2),
            'Stock Value (PKR)': round(stock * cost, 2),
            'Retail Value (PKR)': round(stock * price, 2),
        })
    return _xlsx_response(rows, list(rows[0].keys()) if rows else [], 'Stock Valuation', 'stock_valuation')


@app.route('/api/reports/reorder-alert/excel')
def reorder_alert_excel():
    products = list(db.products.find().sort('stock', 1))
    rows = []
    for p in products:
        stock   = float(p.get('stock', 0))
        reorder = float(p.get('reorder_level', 0))
        if stock <= max(reorder, 1):
            if stock == 0: status = 'Out of Stock'
            elif reorder > 0 and stock <= reorder * 0.5: status = 'Critical'
            else: status = 'Low'
            rows.append({
                'Product': p.get('name', ''),
                'Barcode': p.get('barcode', ''),
                'Category': p.get('category', ''),
                'Supplier': p.get('supplier', ''),
                'Current Stock': stock,
                'Reorder Level': reorder,
                'Shortage': max(0, reorder - stock),
                'Status': status,
            })
    return _xlsx_response(rows, list(rows[0].keys()) if rows else [], 'Reorder Alert', 'reorder_alert')


@app.route('/api/reports/purchase-orders-report/excel')
def purchase_orders_report_excel():
    pos = list(db.purchase_orders_v2.find().sort('created_at', -1).limit(500))
    rows = []
    for po in pos:
        rows.append({
            'PO Number': po.get('po_number', str(po['_id'])[:8]),
            'Vendor': po.get('vendor_name', ''),
            'Created': str(po.get('created_at', ''))[:10],
            'Expected Delivery': str(po.get('expected_delivery', ''))[:10],
            'Items': len(po.get('items', [])),
            'Total (PKR)': round(float(po.get('total_amount', 0)), 2),
            'Status': po.get('status', ''),
        })
    return _xlsx_response(rows, list(rows[0].keys()) if rows else [], 'Purchase Orders', 'purchase_orders')


@app.route('/api/reports/invoice-aging/excel')
def invoice_aging_excel():
    invoices = list(db.vendor_invoices.find().sort('created_at', -1))
    today = datetime.utcnow()
    rows = []
    for inv in invoices:
        created = inv.get('created_at')
        if isinstance(created, str):
            try: created = datetime.fromisoformat(created[:10])
            except: created = today
        elif not isinstance(created, datetime):
            created = today
        age = (today - created).days
        if age <= 30: bucket = '0-30 days'
        elif age <= 60: bucket = '31-60 days'
        elif age <= 90: bucket = '61-90 days'
        else: bucket = '90+ days'
        amount = float(inv.get('amount', 0))
        paid   = float(inv.get('paid_amount', 0))
        rows.append({
            'Invoice No': inv.get('invoice_number', str(inv['_id'])[:8]),
            'Vendor': inv.get('vendor_name', ''),
            'Invoice Date': str(inv.get('created_at', ''))[:10],
            'Due Date': str(inv.get('due_date', ''))[:10],
            'Amount (PKR)': round(amount, 2),
            'Paid (PKR)': round(paid, 2),
            'Outstanding (PKR)': round(amount - paid, 2),
            'Status': inv.get('status', ''),
            'Age (Days)': age,
            'Age Bucket': bucket,
        })
    return _xlsx_response(rows, list(rows[0].keys()) if rows else [], 'Invoice Aging', 'invoice_aging')


@app.route('/api/reports/profit-loss/excel')
def profit_loss_excel():
    try:
        rev_data = {r['_id']: float(r.get('revenue', 0)) for r in db.sales.aggregate([
            {'$group': {'_id': {'$dateToString': {'format': '%Y-%m', 'date': '$timestamp'}}, 'revenue': {'$sum': '$total'}}},
            {'$sort': {'_id': 1}}
        ])}
    except: rev_data = {}
    try:
        cost_data = {r['_id']: float(r.get('purchases', 0)) for r in db.purchase_orders_v2.aggregate([
            {'$group': {'_id': {'$substr': ['$created_at', 0, 7]}, 'purchases': {'$sum': '$total_amount'}}},
            {'$sort': {'_id': 1}}
        ])}
    except: cost_data = {}
    try:
        payroll_data = {r['_id']: float(r.get('payroll', 0)) for r in db.payroll.aggregate([
            {'$group': {'_id': '$month', 'payroll': {'$sum': '$net_salary'}}},
            {'$sort': {'_id': 1}}
        ])}
    except: payroll_data = {}
    all_months = sorted(set(list(rev_data.keys()) + list(cost_data.keys()) + list(payroll_data.keys())))
    rows = []
    for month in all_months[-24:]:
        revenue   = round(rev_data.get(month, 0), 2)
        purchases = round(cost_data.get(month, 0), 2)
        payroll   = round(payroll_data.get(month, 0), 2)
        expenses  = round(purchases + payroll, 2)
        profit    = round(revenue - expenses, 2)
        rows.append({
            'Month': month,
            'Revenue (PKR)': revenue,
            'Purchases (PKR)': purchases,
            'Payroll (PKR)': payroll,
            'Total Expenses (PKR)': expenses,
            'Net Profit (PKR)': profit,
            'Margin %': round(profit / revenue * 100, 1) if revenue > 0 else 0,
        })
    return _xlsx_response(rows, list(rows[0].keys()) if rows else [], 'Profit & Loss', 'profit_loss')


@app.route('/api/reports/sales-returns/excel')
def sales_returns_excel():
    returns = list(db.sales_returns.find().sort('timestamp', -1).limit(500))
    rows = []
    for r in returns:
        ts = r.get('timestamp')
        rows.append({
            'Return ID': str(r['_id'])[:8].upper(),
            'Original Sale': str(r.get('original_sale_id', ''))[:8],
            'Date': ts.strftime('%Y-%m-%d') if isinstance(ts, datetime) else str(ts or '')[:10],
            'Cashier': r.get('cashier_name', ''),
            'Customer': r.get('customer_name', '') or 'Walk-in',
            'Reason': r.get('reason', ''),
            'Items': len(r.get('items', [])),
            'Refund (PKR)': round(float(r.get('refund_amount', 0)), 2),
        })
    return _xlsx_response(rows, list(rows[0].keys()) if rows else [], 'Sales Returns', 'sales_returns')


@app.route('/api/reports/payroll-summary/excel')
def payroll_summary_excel():
    recs = list(db.payroll.find().sort('month', -1).limit(500))
    rows = []
    for r in recs:
        rows.append({
            'Payroll ID': r.get('payroll_id', str(r['_id'])[:8]),
            'Employee': r.get('employee_name', ''),
            'Department': r.get('department', ''),
            'Month': r.get('month', ''),
            'Basic Salary (PKR)': round(float(r.get('basic_salary', 0)), 2),
            'Allowances (PKR)': round(float(r.get('allowances', 0)), 2),
            'Deductions (PKR)': round(float(r.get('deductions', 0)), 2),
            'Net Salary (PKR)': round(float(r.get('net_salary', 0)), 2),
            'Status': r.get('status', ''),
        })
    return _xlsx_response(rows, list(rows[0].keys()) if rows else [], 'Payroll Summary', 'payroll_summary')


@app.route('/api/reports/credit-sales/excel')
def credit_sales_excel():
    sales = list(db.sales.find({'payment_method': {'$regex': 'credit|account', '$options': 'i'}}).sort('timestamp', -1).limit(500))
    rows = []
    for s in sales:
        ts = s.get('timestamp')
        rows.append({
            'Receipt No': str(s['_id'])[:8].upper(),
            'Date': ts.strftime('%Y-%m-%d') if isinstance(ts, datetime) else str(ts or '')[:10],
            'Customer': s.get('customer_name', '') or 'Walk-in',
            'Cashier': s.get('cashier_name', ''),
            'Items': len(s.get('items', [])),
            'Payment Method': s.get('payment_method', ''),
            'Total (PKR)': round(float(s.get('total', 0)), 2),
        })
    return _xlsx_response(rows, list(rows[0].keys()) if rows else [], 'Credit Sales', 'credit_sales')


def excel_response(csv_content, filename):
    """Return CSV as Excel downloadable response"""
    from io import BytesIO
    bio = BytesIO(csv_content.encode('utf-8-sig'))  # BOM for Excel
    bio.seek(0)
    from flask import send_file
    return send_file(
        bio,
        mimetype='text/csv',
        as_attachment=True,
        download_name=f'{filename}_{datetime.utcnow().strftime("%Y-%m-%d")}.csv'
    )


# Helper data functions for Excel exports (duplicate logic from above, but return raw data)
def stock_report_data():
    vendor = request.args.get('vendor')
    product_id = request.args.get('product')
    if USE_MEMORY_DB:
        products = db.find('products', {})
        if vendor:
            products = [p for p in products if vendor.lower() in (p.get('supplier', '') or '').lower()]
        if product_id:
            products = [p for p in products if str(p.get('_id')) == str(product_id)]
        return products
    query = {}
    if vendor:
        query['supplier'] = {'$regex': vendor, '$options': 'i'}
    if product_id:
        query['_id'] = ObjectId(product_id)
    return list(db.products.find(query).sort('stock', 1))


def expiry_report_data():
    vendor = request.args.get('vendor')
    date_from = request.args.get('from')
    date_to = request.args.get('to')
    if USE_MEMORY_DB:
        products = db.find('products', {})
        results = []
        for p in products:
            exp = p.get('expiry_date', '')
            if exp:
                if vendor and vendor.lower() not in (p.get('supplier', '') or '').lower():
                    continue
                if date_from and exp < date_from:
                    continue
                if date_to and exp > date_to:
                    continue
                results.append(p)
        return results
    query = {'expiry_date': {'$exists': True, '$ne': ''}}
    if vendor:
        query['supplier'] = {'$regex': vendor, '$options': 'i'}
    if date_from and date_to:
        query['expiry_date'] = {'$gte': date_from, '$lte': date_to}
    elif date_from:
        query['expiry_date'] = {'$gte': date_from}
    elif date_to:
        query['expiry_date'] = {'$lte': date_to}
    return list(db.products.find(query))


def price_list_report_data():
    vendor = request.args.get('vendor')
    if USE_MEMORY_DB:
        products = db.find('products', {})
        if vendor:
            products = [p for p in products if vendor.lower() in (p.get('supplier', '') or '').lower()]
        result = []
        for p in products:
            result.append({
                'name': p['name'],
                'barcode': p.get('barcode', ''),
                'category': p.get('category', ''),
                'supplier': p.get('supplier', ''),
                'selling_price': p['price'],
                'cost_price': p.get('cost_price', 0),
                'mrp': p.get('mrp', 0),
                'discount_percent': p.get('discount_percent', 0)
            })
        return result
    query = {}
    if vendor:
        query['supplier'] = {'$regex': vendor, '$options': 'i'}
    products = list(db.products.find(query))
    result = []
    for p in products:
        result.append({
            'name': p['name'],
            'barcode': p.get('barcode', ''),
            'category': p.get('category', ''),
            'supplier': p.get('supplier', ''),
            'selling_price': p['price'],
            'cost_price': p.get('cost_price', 0),
            'mrp': p.get('mrp', 0),
            'discount_percent': p.get('discount_percent', 0)
        })
    return result


def slow_moving_report_data():
    days = int(request.args.get('days', 30))
    cutoff = datetime.utcnow() - timedelta(days=days)
    if USE_MEMORY_DB:
        products = db.find('products', {})
        recent_sales = getattr(db, 'sales', [])
        product_ids_sold = set()
        for s in recent_sales:
            s_date = s.get('timestamp')
            if isinstance(s_date, str):
                try:
                    s_date = datetime.fromisoformat(s_date)
                except:
                    s_date = datetime.min
            if s_date >= cutoff:
                for item in s.get('items', []):
                    product_ids_sold.add(item.get('product_id', ''))
        slow_items = [p for p in products if str(p.get('_id')) not in product_ids_sold]
        for p in slow_items:
            p['days_since'] = days
            # Find last sale ever
            last_sold = None
            for s in recent_sales:
                s_date = s.get('timestamp')
                if isinstance(s_date, str):
                    try:
                        s_date = datetime.fromisoformat(s_date)
                    except:
                        continue
                if s_date:
                    for item in s.get('items', []):
                        if item.get('product_id') == str(p.get('_id')):
                            if not last_sold or s_date > last_sold:
                                last_sold = s_date
            p['last_sold'] = last_sold.strftime('%Y-%m-%d') if last_sold else 'Never'
        return slow_items
    pipeline = [
        {'$match': {'timestamp': {'$gte': cutoff}}},
        {'$unwind': '$items'},
        {'$group': {'_id': '$items.product_id'}},
        {'$project': {'_id': 1}}
    ]
    sold_products = {str(r['_id']) for r in db.sales.aggregate(pipeline)}
    unsold = list(db.products.find({'_id': {'$nin': [ObjectId(pid) for pid in sold_products]}}))
    for p in unsold:
        p['days_since'] = days
        # Find last sale ever (if any)
        last_sale = db.sales.find({'items.product_id': p['_id']}).sort('timestamp', -1).limit(1)
        last_list = list(last_sale)
        if last_list:
            ts = last_list[0]['timestamp']
            if isinstance(ts, datetime):
                p['last_sold'] = ts.strftime('%Y-%m-%d')
            else:
                p['last_sold'] = str(ts)
        else:
            p['last_sold'] = 'Never'
    return unsold


def vendor_payment_report_data():
    vendor_id = request.args.get('vendor_id')
    date_from = request.args.get('from')
    date_to = request.args.get('to')
    if USE_MEMORY_DB:
        orders = getattr(db, 'purchase_orders', [])
        if vendor_id:
            orders = [o for o in orders if o.get('vendor_id') == vendor_id]
        return orders
    query = {}
    if vendor_id:
        query['vendor_id'] = vendor_id
    return list(db.purchase_orders.find(query).sort('order_date', -1))


def daily_products_report_data():
    date_filter = request.args.get('date') or datetime.utcnow().strftime('%Y-%m-%d')
    vendor_filter = request.args.get('vendor')
    category_filter = request.args.get('category')
    if USE_MEMORY_DB:
        products = db.find('products', {})
        if vendor_filter:
            products = [p for p in products if vendor_filter.lower() in (p.get('supplier', '') or '').lower()]
        if category_filter:
            products = [p for p in products if p.get('category') == category_filter]
        return products
    query = {}
    if vendor_filter:
        query['supplier'] = {'$regex': vendor_filter, '$options': 'i'}
    if category_filter:
        query['category'] = category_filter
    return list(db.products.find(query))


# =============================================================================
# PROCUREMENT WORKFLOW — PR → PO → GR → INVOICE → PAYMENT
# =============================================================================

STANDARD_COA = [
    # Assets
    {'code':'1000','name':'Cash & Cash Equivalents','type':'asset','subtype':'current','normal_balance':'debit','description':'Physical cash and petty cash fund'},
    {'code':'1010','name':'Bank Account','type':'asset','subtype':'current','normal_balance':'debit','description':'Business bank current account'},
    {'code':'1100','name':'Accounts Receivable','type':'asset','subtype':'current','normal_balance':'debit','description':'Amounts owed by customers'},
    {'code':'1200','name':'Inventory','type':'asset','subtype':'current','normal_balance':'debit','description':'Stock of goods held for sale'},
    {'code':'1300','name':'Prepaid Expenses','type':'asset','subtype':'current','normal_balance':'debit','description':'Expenses paid in advance'},
    {'code':'1400','name':'Fixed Assets','type':'asset','subtype':'non_current','normal_balance':'debit','description':'Property, plant and equipment at cost'},
    {'code':'1500','name':'Accumulated Depreciation','type':'asset','subtype':'contra','normal_balance':'credit','description':'Accumulated depreciation on fixed assets'},
    # Liabilities
    {'code':'2000','name':'Accounts Payable','type':'liability','subtype':'current','normal_balance':'credit','description':'Amounts owed to suppliers/vendors'},
    {'code':'2100','name':'Accrued Liabilities','type':'liability','subtype':'current','normal_balance':'credit','description':'Expenses incurred but not yet paid'},
    {'code':'2200','name':'Sales Tax Payable','type':'liability','subtype':'current','normal_balance':'credit','description':'GST/Sales tax collected, due to government'},
    {'code':'2300','name':'Short-term Loans','type':'liability','subtype':'current','normal_balance':'credit','description':'Loans and overdrafts due within one year'},
    {'code':'2400','name':'Long-term Loans','type':'liability','subtype':'non_current','normal_balance':'credit','description':'Loans due after one year'},
    # Equity
    {'code':'3000','name':"Owner's Capital",'type':'equity','subtype':'capital','normal_balance':'credit','description':'Owner investment and paid-in capital'},
    {'code':'3100','name':'Retained Earnings','type':'equity','subtype':'retained','normal_balance':'credit','description':'Cumulative net profits retained in business'},
    {'code':'3200','name':"Owner's Drawings",'type':'equity','subtype':'drawings','normal_balance':'debit','description':'Owner withdrawals from the business'},
    # Revenue
    {'code':'4000','name':'Sales Revenue','type':'revenue','subtype':'operating','normal_balance':'credit','description':'Income from product sales at POS and invoices'},
    {'code':'4100','name':'Other Income','type':'revenue','subtype':'non_operating','normal_balance':'credit','description':'Miscellaneous income not from core sales'},
    # COGS
    {'code':'5000','name':'Cost of Goods Sold','type':'cogs','subtype':'direct','normal_balance':'debit','description':'Direct cost of products sold to customers'},
    {'code':'5100','name':'Purchase Returns','type':'cogs','subtype':'contra','normal_balance':'credit','description':'Goods returned to vendors (reduces COGS)'},
    # Expenses
    {'code':'6000','name':'Salaries & Wages','type':'expense','subtype':'operating','normal_balance':'debit','description':'Employee compensation'},
    {'code':'6100','name':'Rent Expense','type':'expense','subtype':'operating','normal_balance':'debit','description':'Store and office rent'},
    {'code':'6200','name':'Utilities Expense','type':'expense','subtype':'operating','normal_balance':'debit','description':'Electricity, water, gas bills'},
    {'code':'6300','name':'Marketing & Advertising','type':'expense','subtype':'operating','normal_balance':'debit','description':'Promotional and marketing costs'},
    {'code':'6400','name':'Depreciation Expense','type':'expense','subtype':'operating','normal_balance':'debit','description':'Depreciation on fixed assets'},
    {'code':'6500','name':'Bank Charges','type':'expense','subtype':'operating','normal_balance':'debit','description':'Bank fees, transfer charges'},
    {'code':'6900','name':'Miscellaneous Expense','type':'expense','subtype':'operating','normal_balance':'debit','description':'Other general operating expenses'},
]

def init_gl_accounts():
    """Seed standard COA if collection is empty. MongoDB only."""
    if USE_MEMORY_DB:
        return
    try:
        if db.gl_accounts.count_documents({}) == 0:
            now = datetime.utcnow().isoformat()
            for acct in STANDARD_COA:
                db.gl_accounts.insert_one({**acct, 'is_active': True, 'created_at': now})
    except Exception as e:
        print(f"GL init error: {e}")

def get_account_balance(code, date_from=None, date_to=None):
    """Return (total_dr, total_cr, net) for a GL account code, optionally filtered by date."""
    if USE_MEMORY_DB:
        return 0, 0, 0
    match = {'lines.account_code': code}
    if date_from or date_to:
        date_filter = {}
        if date_from: date_filter['$gte'] = date_from
        if date_to:   date_filter['$lte'] = date_to
        match['date'] = date_filter
    pipeline = [
        {'$match': match},
        {'$unwind': '$lines'},
        {'$match': {'lines.account_code': code}},
        {'$group': {'_id': None, 'total_dr': {'$sum': '$lines.dr'}, 'total_cr': {'$sum': '$lines.cr'}}}
    ]
    result = list(db.journal_entries.aggregate(pipeline))
    if not result:
        return 0, 0, 0
    dr = round(result[0]['total_dr'], 2)
    cr = round(result[0]['total_cr'], 2)
    acct = db.gl_accounts.find_one({'code': code})
    if acct and acct.get('normal_balance') == 'credit':
        net = round(cr - dr, 2)
    else:
        net = round(dr - cr, 2)
    return dr, cr, net

def post_journal_entry(description, lines, reference_type='manual', reference_id='', reference_number='', entry_date=None):
    """Post a balanced double-entry journal. lines = [{'account_code','account_name','dr','cr'}]. Returns inserted id or None."""
    if USE_MEMORY_DB:
        return None
    total_dr = round(sum(float(l.get('dr', 0)) for l in lines), 2)
    total_cr = round(sum(float(l.get('cr', 0)) for l in lines), 2)
    if abs(total_dr - total_cr) > 0.01:
        print(f"JOURNAL IMBALANCE: DR={total_dr} CR={total_cr} for '{description}'")
        return None
    try:
        entry = {
            'entry_number': generate_doc_number('JE'),
            'date': entry_date or datetime.utcnow().strftime('%Y-%m-%d'),
            'description': description,
            'reference_type': reference_type,
            'reference_id': str(reference_id),
            'reference_number': str(reference_number),
            'lines': [{'account_code': l['account_code'], 'account_name': l['account_name'],
                       'dr': round(float(l.get('dr', 0)), 2), 'cr': round(float(l.get('cr', 0)), 2)} for l in lines],
            'total_dr': total_dr,
            'total_cr': total_cr,
            'posted_by': session.get('full_name', 'System') if session else 'System',
            'created_at': datetime.utcnow().isoformat(),
        }
        result = db.journal_entries.insert_one(entry)
        return str(result.inserted_id)
    except Exception as e:
        print(f"Journal post error: {e}")
        return None


def generate_doc_number(prefix):
    """Auto-increment document number: PREFIX-YEAR-NNNN"""
    year = datetime.utcnow().year
    key = f"{prefix}-{year}"
    if USE_MEMORY_DB:
        if not hasattr(db, '_counters'):
            db._counters = {}
        db._counters[key] = db._counters.get(key, 0) + 1
        return f"{key}-{db._counters[key]:04d}"
    result = db.counters.find_one_and_update(
        {'_id': key},
        {'$inc': {'seq': 1}},
        upsert=True,
        return_document=True
    )
    return f"{key}-{result['seq']:04d}"


def proc_serialize(doc):
    """Serialize a procurement document (converts ObjectId to str)."""
    if not doc:
        return doc
    if USE_MEMORY_DB:
        return doc
    doc['_id'] = str(doc['_id'])
    for field in ('vendor_id', 'pr_id', 'po_id', 'gr_id', 'invoice_id', 'created_by'):
        if field in doc and doc[field] is not None:
            try:
                doc[field] = str(doc[field])
            except Exception:
                pass
    if 'items' in doc and isinstance(doc['items'], list):
        for item in doc['items']:
            for f in ('product_id', 'pr_item_id', 'po_item_id'):
                if f in item and item[f] is not None:
                    try:
                        item[f] = str(item[f])
                    except Exception:
                        pass
    return doc


def _update_pr_ordered_qty(pr_id):
    """Recalculate PR ordered quantities and status from all non-cancelled POs."""
    if USE_MEMORY_DB:
        return
    pr = db.purchase_requisitions.find_one({'_id': ObjectId(pr_id)})
    if not pr:
        return
    pos = list(db.purchase_orders_v2.find({'pr_id': ObjectId(pr_id), 'status': {'$ne': 'cancelled'}}))
    ordered_map = {}
    for po in pos:
        for item in po.get('items', []):
            pid = str(item.get('pr_item_id', ''))
            ordered_map[pid] = ordered_map.get(pid, 0) + item.get('qty_ordered', 0)
    new_items = []
    all_fully = True
    any_ordered = False
    for item in pr.get('items', []):
        key = str(item.get('_id', item.get('pr_item_id', '')))
        qty_ord = ordered_map.get(key, 0)
        item['qty_ordered'] = qty_ord
        if qty_ord < item.get('qty_requested', 0):
            all_fully = False
        if qty_ord > 0:
            any_ordered = True
        new_items.append(item)
    if all_fully and any_ordered:
        pr_status = 'fully_ordered'
    elif any_ordered:
        pr_status = 'partially_ordered'
    else:
        pr_status = pr.get('status', 'approved')
    db.purchase_requisitions.update_one(
        {'_id': ObjectId(pr_id)},
        {'$set': {'items': new_items, 'status': pr_status}}
    )


# ── Purchase Requisitions ────────────────────────────────────────────────────

@app.route('/procurement/pr')
def pr_page():
    if 'user_id' not in session:
        return redirect(url_for('login_page'))
    if session.get('role') not in ['admin', 'procurement_manager']:
        return redirect_to_dashboard(session.get('role'))
    return render_template('procurement_pr.html', user=session.get('full_name'), role=session.get('role'))


@app.route('/api/procurement/pr', methods=['GET'])
def get_prs():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    status_f = request.args.get('status')
    if USE_MEMORY_DB:
        docs = getattr(db, 'purchase_requisitions', [])
        if status_f:
            docs = [d for d in docs if d.get('status') == status_f]
        return jsonify([proc_serialize(dict(d)) for d in docs])
    query = {}
    if status_f:
        query['status'] = status_f
    docs = list(db.purchase_requisitions.find(query).sort('created_at', -1))
    return jsonify([proc_serialize(d) for d in docs])


@app.route('/api/procurement/pr', methods=['POST'])
def create_pr():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    data = request.get_json() or {}
    items = data.get('items', [])
    if not items:
        return jsonify({'success': False, 'message': 'At least one item is required'}), 400
    now = datetime.utcnow().isoformat()
    doc = {
        'doc_number': generate_doc_number('PR'),
        'title': data.get('title', '').strip() or 'Purchase Requisition',
        'department': data.get('department', '').strip(),
        'required_date': data.get('required_date', ''),
        'notes': data.get('notes', '').strip(),
        'status': 'draft',
        'created_by': session.get('user_id'),
        'created_by_name': session.get('full_name'),
        'created_at': now,
        'updated_at': now,
        'items': [
            {
                'product_id': item.get('product_id', ''),
                'product_name': item.get('product_name', ''),
                'qty_requested': float(item.get('qty_requested', 0)),
                'qty_ordered': 0.0,
                'unit': item.get('unit', ''),
                'estimated_unit_price': float(item.get('estimated_unit_price', 0)),
                'notes': item.get('notes', ''),
            }
            for item in items
        ]
    }
    if USE_MEMORY_DB:
        if not hasattr(db, 'purchase_requisitions'):
            db.purchase_requisitions = []
        doc['_id'] = str(len(db.purchase_requisitions) + 1)
        db.purchase_requisitions.append(doc)
        return jsonify({'success': True, 'id': doc['_id'], 'doc_number': doc['doc_number']}), 201
    result = db.purchase_requisitions.insert_one(doc)
    return jsonify({'success': True, 'id': str(result.inserted_id), 'doc_number': doc['doc_number']}), 201


@app.route('/api/procurement/pr/<pr_id>', methods=['GET'])
def get_pr(pr_id):
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        docs = getattr(db, 'purchase_requisitions', [])
        doc = next((d for d in docs if d.get('_id') == pr_id), None)
    else:
        doc = db.purchase_requisitions.find_one({'_id': ObjectId(pr_id)})
    if not doc:
        return jsonify({'error': 'PR not found'}), 404
    return jsonify(proc_serialize(doc))


@app.route('/api/procurement/pr/<pr_id>', methods=['PUT'])
def update_pr(pr_id):
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    data = request.get_json() or {}
    if USE_MEMORY_DB:
        docs = getattr(db, 'purchase_requisitions', [])
        doc = next((d for d in docs if d.get('_id') == pr_id), None)
        if not doc:
            return jsonify({'error': 'Not found'}), 404
        if doc.get('status') not in ('draft',):
            return jsonify({'success': False, 'message': 'Only draft PRs can be edited'}), 400
        doc.update({'title': data.get('title', doc['title']), 'notes': data.get('notes', doc.get('notes', '')),
                    'required_date': data.get('required_date', doc.get('required_date', '')),
                    'items': data.get('items', doc['items']), 'updated_at': datetime.utcnow().isoformat()})
        return jsonify({'success': True})
    doc = db.purchase_requisitions.find_one({'_id': ObjectId(pr_id)})
    if not doc:
        return jsonify({'error': 'Not found'}), 404
    if doc.get('status') not in ('draft',):
        return jsonify({'success': False, 'message': 'Only draft PRs can be edited'}), 400
    update_fields = {'updated_at': datetime.utcnow().isoformat()}
    for f in ('title', 'department', 'required_date', 'notes', 'items'):
        if f in data:
            update_fields[f] = data[f]
    db.purchase_requisitions.update_one({'_id': ObjectId(pr_id)}, {'$set': update_fields})
    return jsonify({'success': True})


@app.route('/api/procurement/pr/<pr_id>/submit', methods=['POST'])
def submit_pr(pr_id):
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        docs = getattr(db, 'purchase_requisitions', [])
        doc = next((d for d in docs if d.get('_id') == pr_id), None)
        if not doc:
            return jsonify({'error': 'Not found'}), 404
        if doc.get('status') != 'draft':
            return jsonify({'success': False, 'message': 'Only draft PRs can be submitted'}), 400
        doc['status'] = 'pending_approval'
        doc['submitted_at'] = datetime.utcnow().isoformat()
        return jsonify({'success': True})
    doc = db.purchase_requisitions.find_one({'_id': ObjectId(pr_id)})
    if not doc:
        return jsonify({'error': 'Not found'}), 404
    if doc.get('status') != 'draft':
        return jsonify({'success': False, 'message': 'Only draft PRs can be submitted'}), 400
    db.purchase_requisitions.update_one({'_id': ObjectId(pr_id)}, {
        '$set': {'status': 'pending_approval', 'submitted_at': datetime.utcnow().isoformat()}
    })
    return jsonify({'success': True})


@app.route('/api/procurement/pr/<pr_id>/approve', methods=['POST'])
def approve_pr(pr_id):
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if session.get('role') not in ('admin', 'procurement_manager'):
        return jsonify({'error': 'Forbidden'}), 403
    if USE_MEMORY_DB:
        docs = getattr(db, 'purchase_requisitions', [])
        doc = next((d for d in docs if d.get('_id') == pr_id), None)
        if not doc:
            return jsonify({'error': 'Not found'}), 404
        if doc.get('status') != 'pending_approval':
            return jsonify({'success': False, 'message': 'PR is not pending approval'}), 400
        doc['status'] = 'approved'
        doc['approved_by'] = session.get('full_name')
        doc['approved_at'] = datetime.utcnow().isoformat()
        return jsonify({'success': True})
    doc = db.purchase_requisitions.find_one({'_id': ObjectId(pr_id)})
    if not doc:
        return jsonify({'error': 'Not found'}), 404
    if doc.get('status') != 'pending_approval':
        return jsonify({'success': False, 'message': 'PR is not pending approval'}), 400
    db.purchase_requisitions.update_one({'_id': ObjectId(pr_id)}, {
        '$set': {'status': 'approved', 'approved_by': session.get('full_name'),
                 'approved_at': datetime.utcnow().isoformat()}
    })
    return jsonify({'success': True})


@app.route('/api/procurement/pr/<pr_id>/reject', methods=['POST'])
def reject_pr(pr_id):
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if session.get('role') not in ('admin', 'procurement_manager'):
        return jsonify({'error': 'Forbidden'}), 403
    data = request.get_json() or {}
    reason = data.get('reason', '').strip()
    if USE_MEMORY_DB:
        docs = getattr(db, 'purchase_requisitions', [])
        doc = next((d for d in docs if d.get('_id') == pr_id), None)
        if not doc:
            return jsonify({'error': 'Not found'}), 404
        doc['status'] = 'rejected'
        doc['rejection_reason'] = reason
        doc['rejected_at'] = datetime.utcnow().isoformat()
        return jsonify({'success': True})
    db.purchase_requisitions.update_one({'_id': ObjectId(pr_id)}, {
        '$set': {'status': 'rejected', 'rejection_reason': reason, 'rejected_at': datetime.utcnow().isoformat()}
    })
    return jsonify({'success': True})


@app.route('/api/procurement/pr/<pr_id>/cancel', methods=['POST'])
def cancel_pr(pr_id):
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if session.get('role') not in ('admin', 'procurement_manager'):
        return jsonify({'error': 'Forbidden'}), 403
    data = request.get_json() or {}
    reason = data.get('reason', '').strip()
    if USE_MEMORY_DB:
        docs = getattr(db, 'purchase_requisitions', [])
        doc = next((d for d in docs if d.get('_id') == pr_id), None)
        if not doc:
            return jsonify({'error': 'Not found'}), 404
        doc['status'] = 'cancelled'
        doc['cancel_reason'] = reason
        doc['cancelled_at'] = datetime.utcnow().isoformat()
        doc['cancelled_by'] = session.get('full_name')
        return jsonify({'success': True})
    db.purchase_requisitions.update_one({'_id': ObjectId(pr_id)}, {
        '$set': {'status': 'cancelled', 'cancel_reason': reason,
                 'cancelled_at': datetime.utcnow().isoformat(),
                 'cancelled_by': session.get('full_name')}
    })
    return jsonify({'success': True})


@app.route('/api/procurement/pr/<pr_id>/send_back', methods=['POST'])
def send_back_pr(pr_id):
    """Admin sends PR back to initiator — resets to draft so it can be edited and re-submitted."""
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if session.get('role') not in ('admin', 'procurement_manager'):
        return jsonify({'error': 'Forbidden'}), 403
    data = request.get_json() or {}
    remarks = data.get('remarks', '').strip()
    if USE_MEMORY_DB:
        docs = getattr(db, 'purchase_requisitions', [])
        doc = next((d for d in docs if d.get('_id') == pr_id), None)
        if not doc:
            return jsonify({'error': 'Not found'}), 404
        if doc.get('status') != 'pending_approval':
            return jsonify({'success': False, 'message': 'PR is not pending approval'}), 400
        doc['status'] = 'draft'
        doc['send_back_remarks'] = remarks
        doc['sent_back_at'] = datetime.utcnow().isoformat()
        doc['sent_back_by'] = session.get('full_name')
        return jsonify({'success': True})
    doc = db.purchase_requisitions.find_one({'_id': ObjectId(pr_id)})
    if not doc:
        return jsonify({'error': 'Not found'}), 404
    if doc.get('status') != 'pending_approval':
        return jsonify({'success': False, 'message': 'PR is not pending approval'}), 400
    db.purchase_requisitions.update_one({'_id': ObjectId(pr_id)}, {
        '$set': {'status': 'draft', 'send_back_remarks': remarks,
                 'sent_back_at': datetime.utcnow().isoformat(),
                 'sent_back_by': session.get('full_name')}
    })
    return jsonify({'success': True})


# ── Purchase Orders ──────────────────────────────────────────────────────────

@app.route('/procurement/po')
def po_page():
    if 'user_id' not in session:
        return redirect(url_for('login_page'))
    if session.get('role') not in ['admin', 'procurement_manager']:
        return redirect_to_dashboard(session.get('role'))
    return render_template('procurement_po.html', user=session.get('full_name'), role=session.get('role'))


@app.route('/api/procurement/po', methods=['GET'])
def get_pos():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    status_f = request.args.get('status')
    pr_id_f = request.args.get('pr_id')
    if USE_MEMORY_DB:
        docs = getattr(db, 'purchase_orders_v2', [])
        if status_f:
            docs = [d for d in docs if d.get('status') == status_f]
        if pr_id_f:
            docs = [d for d in docs if d.get('pr_id') == pr_id_f]
        return jsonify([proc_serialize(dict(d)) for d in docs])
    query = {}
    if status_f:
        query['status'] = status_f
    if pr_id_f:
        query['pr_id'] = ObjectId(pr_id_f)
    docs = list(db.purchase_orders_v2.find(query).sort('created_at', -1))
    return jsonify([proc_serialize(d) for d in docs])


@app.route('/api/procurement/po', methods=['POST'])
def create_po():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    data = request.get_json() or {}
    pr_id = data.get('pr_id', '')
    vendor_id = data.get('vendor_id', '')
    items = data.get('items', [])
    if not vendor_id:
        return jsonify({'success': False, 'message': 'Vendor is required'}), 400
    if not items:
        return jsonify({'success': False, 'message': 'At least one item required'}), 400
    now = datetime.utcnow().isoformat()
    po_items = []
    total_amount = 0.0
    for item in items:
        qty = float(item.get('qty_ordered', 0))
        unit_price = float(item.get('unit_price', 0))
        line_total = qty * unit_price
        total_amount += line_total
        po_items.append({
            'pr_item_id': item.get('pr_item_id', ''),
            'product_id': item.get('product_id', ''),
            'product_name': item.get('product_name', ''),
            'qty_ordered': qty,
            'qty_received': 0.0,
            'unit': item.get('unit', ''),
            'unit_price': unit_price,
            'line_total': line_total,
            'notes': item.get('notes', ''),
        })
    doc = {
        'doc_number': generate_doc_number('PO'),
        'pr_id': ObjectId(pr_id) if pr_id and not USE_MEMORY_DB else pr_id,
        'vendor_id': ObjectId(vendor_id) if not USE_MEMORY_DB else vendor_id,
        'vendor_name': data.get('vendor_name', ''),
        'expected_delivery': data.get('expected_delivery', ''),
        'payment_terms': data.get('payment_terms', ''),
        'notes': data.get('notes', '').strip(),
        'status': 'draft',
        'total_amount': round(total_amount, 2),
        'created_by': session.get('user_id'),
        'created_by_name': session.get('full_name'),
        'created_at': now,
        'updated_at': now,
        'items': po_items,
    }
    if USE_MEMORY_DB:
        if not hasattr(db, 'purchase_orders_v2'):
            db.purchase_orders_v2 = []
        doc['_id'] = str(len(db.purchase_orders_v2) + 1)
        db.purchase_orders_v2.append(doc)
        return jsonify({'success': True, 'id': doc['_id'], 'doc_number': doc['doc_number']}), 201
    result = db.purchase_orders_v2.insert_one(doc)
    po_id = str(result.inserted_id)
    if pr_id:
        _update_pr_ordered_qty(pr_id)
    return jsonify({'success': True, 'id': po_id, 'doc_number': doc['doc_number']}), 201


@app.route('/api/procurement/po/<po_id>', methods=['GET'])
def get_po(po_id):
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        docs = getattr(db, 'purchase_orders_v2', [])
        doc = next((d for d in docs if d.get('_id') == po_id), None)
    else:
        doc = db.purchase_orders_v2.find_one({'_id': ObjectId(po_id)})
    if not doc:
        return jsonify({'error': 'PO not found'}), 404
    return jsonify(proc_serialize(doc))


@app.route('/api/procurement/po/<po_id>/send', methods=['POST'])
def send_po(po_id):
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        docs = getattr(db, 'purchase_orders_v2', [])
        doc = next((d for d in docs if d.get('_id') == po_id), None)
        if not doc:
            return jsonify({'error': 'Not found'}), 404
        if doc.get('status') != 'draft':
            return jsonify({'success': False, 'message': 'Only draft POs can be sent'}), 400
        doc['status'] = 'sent'
        doc['sent_at'] = datetime.utcnow().isoformat()
        return jsonify({'success': True})
    doc = db.purchase_orders_v2.find_one({'_id': ObjectId(po_id)})
    if not doc:
        return jsonify({'error': 'Not found'}), 404
    if doc.get('status') != 'draft':
        return jsonify({'success': False, 'message': 'Only draft POs can be sent'}), 400
    db.purchase_orders_v2.update_one({'_id': ObjectId(po_id)}, {
        '$set': {'status': 'sent', 'sent_at': datetime.utcnow().isoformat()}
    })
    return jsonify({'success': True})


@app.route('/api/procurement/po/<po_id>/cancel', methods=['POST'])
def cancel_po(po_id):
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if session.get('role') not in ('admin', 'procurement_manager'):
        return jsonify({'error': 'Forbidden'}), 403
    data = request.get_json() or {}
    if USE_MEMORY_DB:
        docs = getattr(db, 'purchase_orders_v2', [])
        doc = next((d for d in docs if d.get('_id') == po_id), None)
        if not doc:
            return jsonify({'error': 'Not found'}), 404
        doc['status'] = 'cancelled'
        doc['cancel_reason'] = data.get('reason', '')
        return jsonify({'success': True})
    doc = db.purchase_orders_v2.find_one({'_id': ObjectId(po_id)})
    if not doc:
        return jsonify({'error': 'Not found'}), 404
    db.purchase_orders_v2.update_one({'_id': ObjectId(po_id)}, {
        '$set': {'status': 'cancelled', 'cancel_reason': data.get('reason', ''),
                 'cancelled_at': datetime.utcnow().isoformat()}
    })
    pr_id = doc.get('pr_id')
    if pr_id:
        _update_pr_ordered_qty(str(pr_id))
    return jsonify({'success': True})


# ── Goods Receipts ───────────────────────────────────────────────────────────

@app.route('/procurement/gr')
def gr_page():
    if 'user_id' not in session:
        return redirect(url_for('login_page'))
    if session.get('role') not in ['admin', 'procurement_manager']:
        return redirect_to_dashboard(session.get('role'))
    return render_template('procurement_gr.html', user=session.get('full_name'), role=session.get('role'))


@app.route('/api/procurement/gr', methods=['GET'])
def get_grs():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    status_f = request.args.get('status')
    po_id_f = request.args.get('po_id')
    if USE_MEMORY_DB:
        docs = getattr(db, 'goods_receipts', [])
        if status_f:
            docs = [d for d in docs if d.get('status') == status_f]
        if po_id_f:
            docs = [d for d in docs if d.get('po_id') == po_id_f]
        return jsonify([proc_serialize(dict(d)) for d in docs])
    query = {}
    if status_f:
        query['status'] = status_f
    if po_id_f:
        query['po_id'] = ObjectId(po_id_f)
    docs = list(db.goods_receipts.find(query).sort('created_at', -1))
    return jsonify([proc_serialize(d) for d in docs])


@app.route('/api/procurement/gr', methods=['POST'])
def create_gr():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    data = request.get_json() or {}
    po_id = data.get('po_id', '')
    items = data.get('items', [])
    if not po_id:
        return jsonify({'success': False, 'message': 'PO reference is required'}), 400
    if not items:
        return jsonify({'success': False, 'message': 'At least one item required'}), 400
    now = datetime.utcnow().isoformat()
    gr_items = []
    for item in items:
        qty = float(item.get('qty_received', 0))
        gr_items.append({
            'po_item_id': item.get('po_item_id', ''),
            'product_id': item.get('product_id', ''),
            'product_name': item.get('product_name', ''),
            'qty_received': qty,
            'unit': item.get('unit', ''),
            'unit_price': float(item.get('unit_price', 0)),
            'batch_no': item.get('batch_no', ''),
            'expiry_date': item.get('expiry_date', ''),
            'notes': item.get('notes', ''),
        })
    doc = {
        'doc_number': generate_doc_number('GR'),
        'po_id': ObjectId(po_id) if not USE_MEMORY_DB else po_id,
        'po_number': data.get('po_number', ''),
        'vendor_id': ObjectId(data.get('vendor_id', '')) if data.get('vendor_id') and not USE_MEMORY_DB else data.get('vendor_id', ''),
        'vendor_name': data.get('vendor_name', ''),
        'delivery_date': data.get('delivery_date', now[:10]),
        'delivery_note_no': data.get('delivery_note_no', ''),
        'notes': data.get('notes', '').strip(),
        'status': 'draft',
        'created_by': session.get('user_id'),
        'created_by_name': session.get('full_name'),
        'created_at': now,
        'updated_at': now,
        'items': gr_items,
    }
    if USE_MEMORY_DB:
        if not hasattr(db, 'goods_receipts'):
            db.goods_receipts = []
        doc['_id'] = str(len(db.goods_receipts) + 1)
        db.goods_receipts.append(doc)
        return jsonify({'success': True, 'id': doc['_id'], 'doc_number': doc['doc_number']}), 201
    result = db.goods_receipts.insert_one(doc)
    return jsonify({'success': True, 'id': str(result.inserted_id), 'doc_number': doc['doc_number']}), 201


@app.route('/api/procurement/gr/<gr_id>', methods=['GET'])
def get_gr(gr_id):
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        docs = getattr(db, 'goods_receipts', [])
        doc = next((d for d in docs if d.get('_id') == gr_id), None)
    else:
        doc = db.goods_receipts.find_one({'_id': ObjectId(gr_id)})
    if not doc:
        return jsonify({'error': 'GR not found'}), 404
    return jsonify(proc_serialize(doc))


@app.route('/api/procurement/gr/<gr_id>/confirm', methods=['POST'])
def confirm_gr(gr_id):
    """Confirm GR: update stock and mark PO items received."""
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        docs = getattr(db, 'goods_receipts', [])
        doc = next((d for d in docs if d.get('_id') == gr_id), None)
        if not doc:
            return jsonify({'error': 'Not found'}), 404
        if doc.get('status') != 'draft':
            return jsonify({'success': False, 'message': 'GR already confirmed'}), 400
        doc['status'] = 'confirmed'
        doc['confirmed_at'] = datetime.utcnow().isoformat()
        return jsonify({'success': True})
    doc = db.goods_receipts.find_one({'_id': ObjectId(gr_id)})
    if not doc:
        return jsonify({'error': 'Not found'}), 404
    if doc.get('status') != 'draft':
        return jsonify({'success': False, 'message': 'GR already confirmed'}), 400
    # Update stock for each received item
    for item in doc.get('items', []):
        pid = item.get('product_id')
        qty = float(item.get('qty_received', 0))
        if pid and qty > 0:
            try:
                db.products.update_one({'_id': ObjectId(pid)}, {'$inc': {'stock': qty}})
            except Exception:
                pass
    # Update PO received quantities
    po_id = doc.get('po_id')
    if po_id:
        po = db.purchase_orders_v2.find_one({'_id': po_id})
        if po:
            gr_map = {str(item.get('po_item_id', '')): item.get('qty_received', 0)
                      for item in doc.get('items', [])}
            new_po_items = []
            all_received = True
            any_received = False
            for pi in po.get('items', []):
                pi_key = str(pi.get('_id', pi.get('po_item_id', '')))
                pi['qty_received'] = pi.get('qty_received', 0) + gr_map.get(pi_key, 0)
                if pi['qty_received'] < pi.get('qty_ordered', 0):
                    all_received = False
                if pi['qty_received'] > 0:
                    any_received = True
                new_po_items.append(pi)
            if all_received and any_received:
                po_status = 'fully_received'
            elif any_received:
                po_status = 'partially_received'
            else:
                po_status = po.get('status', 'sent')
            db.purchase_orders_v2.update_one({'_id': po_id}, {
                '$set': {'items': new_po_items, 'status': po_status}
            })
    db.goods_receipts.update_one({'_id': ObjectId(gr_id)}, {
        '$set': {'status': 'confirmed', 'confirmed_at': datetime.utcnow().isoformat(),
                 'confirmed_by': session.get('full_name')}
    })
    # GL Auto-Post: DR Inventory, CR Accounts Payable
    total_cost = sum(float(i.get('qty_received',0)) * float(i.get('unit_price',0)) for i in doc.get('items',[]))
    if total_cost > 0:
        post_journal_entry(
            description=f"GR Confirmed — {doc.get('doc_number',gr_id)}",
            lines=[
                {'account_code':'1200','account_name':'Inventory','dr':total_cost,'cr':0},
                {'account_code':'2000','account_name':'Accounts Payable','dr':0,'cr':total_cost},
            ],
            reference_type='gr', reference_id=gr_id,
            reference_number=doc.get('doc_number','')
        )
    return jsonify({'success': True})


# ── Vendor Invoices ──────────────────────────────────────────────────────────

@app.route('/procurement/invoices')
def invoices_page():
    if 'user_id' not in session:
        return redirect(url_for('login_page'))
    if session.get('role') not in ['admin', 'procurement_manager']:
        return redirect_to_dashboard(session.get('role'))
    return render_template('procurement_invoices.html', user=session.get('full_name'), role=session.get('role'))


@app.route('/api/procurement/invoices', methods=['GET'])
def get_invoices():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    status_f = request.args.get('status')
    gr_id_f = request.args.get('gr_id')
    if USE_MEMORY_DB:
        docs = getattr(db, 'vendor_invoices', [])
        if status_f:
            docs = [d for d in docs if d.get('status') == status_f]
        if gr_id_f:
            docs = [d for d in docs if d.get('gr_id') == gr_id_f]
        return jsonify([proc_serialize(dict(d)) for d in docs])
    query = {}
    if status_f:
        query['status'] = status_f
    if gr_id_f:
        query['gr_id'] = ObjectId(gr_id_f)
    docs = list(db.vendor_invoices.find(query).sort('created_at', -1))
    return jsonify([proc_serialize(d) for d in docs])


@app.route('/api/procurement/invoices', methods=['POST'])
def create_invoice():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    data = request.get_json() or {}
    gr_id = data.get('gr_id', '')
    if not gr_id:
        return jsonify({'success': False, 'message': 'GR reference is required'}), 400
    now = datetime.utcnow().isoformat()
    invoice_amount = float(data.get('invoice_amount', 0))
    tax_amount = float(data.get('tax_amount', 0))
    total_amount = invoice_amount + tax_amount
    doc = {
        'doc_number': generate_doc_number('INV'),
        'gr_id': ObjectId(gr_id) if not USE_MEMORY_DB else gr_id,
        'gr_number': data.get('gr_number', ''),
        'po_id': ObjectId(data.get('po_id', '')) if data.get('po_id') and not USE_MEMORY_DB else data.get('po_id', ''),
        'po_number': data.get('po_number', ''),
        'vendor_id': ObjectId(data.get('vendor_id', '')) if data.get('vendor_id') and not USE_MEMORY_DB else data.get('vendor_id', ''),
        'vendor_name': data.get('vendor_name', ''),
        'vendor_invoice_no': data.get('vendor_invoice_no', '').strip(),
        'invoice_date': data.get('invoice_date', now[:10]),
        'due_date': data.get('due_date', ''),
        'invoice_amount': invoice_amount,
        'tax_amount': tax_amount,
        'total_amount': round(total_amount, 2),
        'amount_paid': 0.0,
        'amount_due': round(total_amount, 2),
        'status': 'pending',
        'notes': data.get('notes', '').strip(),
        'created_by': session.get('user_id'),
        'created_by_name': session.get('full_name'),
        'created_at': now,
        'updated_at': now,
    }
    if USE_MEMORY_DB:
        if not hasattr(db, 'vendor_invoices'):
            db.vendor_invoices = []
        doc['_id'] = str(len(db.vendor_invoices) + 1)
        db.vendor_invoices.append(doc)
        return jsonify({'success': True, 'id': doc['_id'], 'doc_number': doc['doc_number']}), 201
    result = db.vendor_invoices.insert_one(doc)
    return jsonify({'success': True, 'id': str(result.inserted_id), 'doc_number': doc['doc_number']}), 201


@app.route('/api/procurement/invoices/<inv_id>', methods=['GET'])
def get_invoice(inv_id):
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        docs = getattr(db, 'vendor_invoices', [])
        doc = next((d for d in docs if d.get('_id') == inv_id), None)
    else:
        doc = db.vendor_invoices.find_one({'_id': ObjectId(inv_id)})
    if not doc:
        return jsonify({'error': 'Invoice not found'}), 404
    return jsonify(proc_serialize(doc))


@app.route('/api/procurement/invoices/<inv_id>/cancel', methods=['POST'])
def cancel_invoice(inv_id):
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if session.get('role') not in ('admin', 'procurement_manager'):
        return jsonify({'error': 'Forbidden'}), 403
    if USE_MEMORY_DB:
        docs = getattr(db, 'vendor_invoices', [])
        doc = next((d for d in docs if d.get('_id') == inv_id), None)
        if not doc:
            return jsonify({'error': 'Not found'}), 404
        if doc.get('amount_paid', 0) > 0:
            return jsonify({'success': False, 'message': 'Cannot cancel a partially/fully paid invoice'}), 400
        doc['status'] = 'cancelled'
        return jsonify({'success': True})
    doc = db.vendor_invoices.find_one({'_id': ObjectId(inv_id)})
    if not doc:
        return jsonify({'error': 'Not found'}), 404
    if doc.get('amount_paid', 0) > 0:
        return jsonify({'success': False, 'message': 'Cannot cancel a partially/fully paid invoice'}), 400
    db.vendor_invoices.update_one({'_id': ObjectId(inv_id)}, {
        '$set': {'status': 'cancelled', 'cancelled_at': datetime.utcnow().isoformat()}
    })
    return jsonify({'success': True})


# ── Vendor Payments ──────────────────────────────────────────────────────────

@app.route('/procurement/payments')
def procurement_payments_page():
    if 'user_id' not in session:
        return redirect(url_for('login_page'))
    if session.get('role') not in ['admin', 'procurement_manager']:
        return redirect_to_dashboard(session.get('role'))
    return render_template('procurement_payments.html', user=session.get('full_name'), role=session.get('role'))


@app.route('/api/procurement/payments', methods=['GET'])
def get_payments():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    invoice_id_f = request.args.get('invoice_id')
    if USE_MEMORY_DB:
        docs = getattr(db, 'vendor_payments_v2', [])
        if invoice_id_f:
            docs = [d for d in docs if d.get('invoice_id') == invoice_id_f]
        return jsonify([proc_serialize(dict(d)) for d in docs])
    query = {}
    if invoice_id_f:
        query['invoice_id'] = ObjectId(invoice_id_f)
    docs = list(db.vendor_payments_v2.find(query).sort('created_at', -1))
    return jsonify([proc_serialize(d) for d in docs])


@app.route('/api/procurement/payments', methods=['POST'])
def create_payment():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    data = request.get_json() or {}
    invoice_id = data.get('invoice_id', '')
    amount = float(data.get('amount', 0))
    if not invoice_id:
        return jsonify({'success': False, 'message': 'Invoice reference is required'}), 400
    if amount <= 0:
        return jsonify({'success': False, 'message': 'Payment amount must be greater than zero'}), 400
    # Fetch invoice
    if USE_MEMORY_DB:
        inv_docs = getattr(db, 'vendor_invoices', [])
        invoice = next((d for d in inv_docs if d.get('_id') == invoice_id), None)
    else:
        invoice = db.vendor_invoices.find_one({'_id': ObjectId(invoice_id)})
    if not invoice:
        return jsonify({'success': False, 'message': 'Invoice not found'}), 404
    if invoice.get('status') == 'cancelled':
        return jsonify({'success': False, 'message': 'Cannot pay a cancelled invoice'}), 400
    amount_due = float(invoice.get('amount_due', 0))
    if amount > amount_due + 0.001:
        return jsonify({'success': False, 'message': f'Payment amount exceeds amount due ({amount_due:.2f})'}), 400
    now = datetime.utcnow().isoformat()
    doc = {
        'doc_number': generate_doc_number('PAY'),
        'invoice_id': ObjectId(invoice_id) if not USE_MEMORY_DB else invoice_id,
        'invoice_number': invoice.get('doc_number', ''),
        'vendor_id': invoice.get('vendor_id', ''),
        'vendor_name': invoice.get('vendor_name', ''),
        'amount': round(amount, 2),
        'payment_method': data.get('payment_method', 'bank_transfer'),
        'payment_date': data.get('payment_date', now[:10]),
        'reference_no': data.get('reference_no', '').strip(),
        'notes': data.get('notes', '').strip(),
        'created_by': session.get('user_id'),
        'created_by_name': session.get('full_name'),
        'created_at': now,
    }
    if USE_MEMORY_DB:
        if not hasattr(db, 'vendor_payments_v2'):
            db.vendor_payments_v2 = []
        doc['_id'] = str(len(db.vendor_payments_v2) + 1)
        db.vendor_payments_v2.append(doc)
        # Update invoice
        new_paid = float(invoice.get('amount_paid', 0)) + amount
        new_due = float(invoice.get('total_amount', 0)) - new_paid
        invoice['amount_paid'] = round(new_paid, 2)
        invoice['amount_due'] = round(new_due, 2)
        invoice['status'] = 'paid' if new_due <= 0.001 else 'partially_paid'
        return jsonify({'success': True, 'id': doc['_id'], 'doc_number': doc['doc_number']}), 201
    result = db.vendor_payments_v2.insert_one(doc)
    new_paid = float(invoice.get('amount_paid', 0)) + amount
    new_due = float(invoice.get('total_amount', 0)) - new_paid
    inv_status = 'paid' if new_due <= 0.001 else 'partially_paid'
    db.vendor_invoices.update_one({'_id': ObjectId(invoice_id)}, {
        '$set': {'amount_paid': round(new_paid, 2), 'amount_due': round(new_due, 2),
                 'status': inv_status, 'updated_at': now}
    })
    # GL Auto-Post: DR Accounts Payable, CR Bank
    pay_amount = float(doc.get('amount', 0))
    if pay_amount > 0:
        post_journal_entry(
            description=f"Vendor Payment — {doc.get('doc_number','')} to {doc.get('vendor_name','')}",
            lines=[
                {'account_code':'2000','account_name':'Accounts Payable','dr':pay_amount,'cr':0},
                {'account_code':'1010','account_name':'Bank Account','dr':0,'cr':pay_amount},
            ],
            reference_type='payment', reference_id=str(result.inserted_id),
            reference_number=doc.get('doc_number','')
        )
    return jsonify({'success': True, 'id': str(result.inserted_id), 'doc_number': doc['doc_number']}), 201


@app.route('/api/procurement/payments/<pay_id>', methods=['GET'])
def get_payment(pay_id):
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        docs = getattr(db, 'vendor_payments_v2', [])
        doc = next((d for d in docs if d.get('_id') == pay_id), None)
    else:
        doc = db.vendor_payments_v2.find_one({'_id': ObjectId(pay_id)})
    if not doc:
        return jsonify({'error': 'Payment not found'}), 404
    return jsonify(proc_serialize(doc))


# ── Procurement Summary ──────────────────────────────────────────────────────

@app.route('/api/procurement/summary', methods=['GET'])
def procurement_summary():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        pr_list = getattr(db, 'purchase_requisitions', [])
        po_list = getattr(db, 'purchase_orders_v2', [])
        gr_list = getattr(db, 'goods_receipts', [])
        inv_list = getattr(db, 'vendor_invoices', [])
        pay_list = getattr(db, 'vendor_payments_v2', [])
        return jsonify({
            'pr': {'total': len(pr_list), 'pending': sum(1 for d in pr_list if d.get('status') == 'pending_approval'),
                   'approved': sum(1 for d in pr_list if d.get('status') == 'approved')},
            'po': {'total': len(po_list), 'draft': sum(1 for d in po_list if d.get('status') == 'draft'),
                   'sent': sum(1 for d in po_list if d.get('status') == 'sent'),
                   'total_value': sum(d.get('total_amount', 0) for d in po_list)},
            'gr': {'total': len(gr_list), 'confirmed': sum(1 for d in gr_list if d.get('status') == 'confirmed')},
            'invoices': {'total': len(inv_list), 'pending': sum(1 for d in inv_list if d.get('status') == 'pending'),
                         'total_due': sum(d.get('amount_due', 0) for d in inv_list if d.get('status') != 'cancelled')},
            'payments': {'total': len(pay_list), 'total_paid': sum(d.get('amount', 0) for d in pay_list)},
        })
    pr_col = db.purchase_requisitions
    po_col = db.purchase_orders_v2
    gr_col = db.goods_receipts
    inv_col = db.vendor_invoices
    pay_col = db.vendor_payments_v2
    po_pipeline = [{'$group': {'_id': None, 'total_value': {'$sum': '$total_amount'}}}]
    po_total_val = list(po_col.aggregate(po_pipeline))
    inv_pipeline = [{'$match': {'status': {'$ne': 'cancelled'}}},
                    {'$group': {'_id': None, 'total_due': {'$sum': '$amount_due'}}}]
    inv_due = list(inv_col.aggregate(inv_pipeline))
    pay_pipeline = [{'$group': {'_id': None, 'total_paid': {'$sum': '$amount'}}}]
    pay_total = list(pay_col.aggregate(pay_pipeline))
    return jsonify({
        'pr': {
            'total': pr_col.count_documents({}),
            'pending': pr_col.count_documents({'status': 'pending_approval'}),
            'approved': pr_col.count_documents({'status': 'approved'}),
        },
        'po': {
            'total': po_col.count_documents({}),
            'draft': po_col.count_documents({'status': 'draft'}),
            'sent': po_col.count_documents({'status': 'sent'}),
            'total_value': po_total_val[0]['total_value'] if po_total_val else 0,
        },
        'gr': {
            'total': gr_col.count_documents({}),
            'confirmed': gr_col.count_documents({'status': 'confirmed'}),
        },
        'invoices': {
            'total': inv_col.count_documents({}),
            'pending': inv_col.count_documents({'status': 'pending'}),
            'total_due': inv_due[0]['total_due'] if inv_due else 0,
        },
        'payments': {
            'total': pay_col.count_documents({}),
            'total_paid': pay_total[0]['total_paid'] if pay_total else 0,
        },
    })


# =============================================================================
# PROCUREMENT MANAGEMENT PAGES
# =============================================================================

@app.route('/admin/users')
def admin_users():
    if 'user_id' not in session:
        return redirect(url_for('login_page'))
    if session.get('role') != 'admin':
        return redirect(url_for('login_page'))
    return render_template('admin_users.html', user=session.get('full_name'), role=session.get('role'))

@app.route('/procurement/vendors')
def procurement_vendors():
    if 'user_id' not in session:
        return redirect(url_for('login_page'))
    return render_template('procurement_vendors.html', user=session.get('full_name'), role=session.get('role'))

@app.route('/procurement/stock')
def procurement_stock():
    if 'user_id' not in session:
        return redirect(url_for('login_page'))
    return render_template('procurement_stock.html', user=session.get('full_name'), role=session.get('role'))

@app.route('/procurement/catalog')
def procurement_catalog():
    if 'user_id' not in session:
        return redirect(url_for('login_page'))
    return render_template('procurement_catalog.html', user=session.get('full_name'), role=session.get('role'))

@app.route('/procurement/qr')
def procurement_qr():
    if 'user_id' not in session:
        return redirect(url_for('login_page'))
    return render_template('procurement_qr.html', user=session.get('full_name'), role=session.get('role'))

@app.route('/procurement/reports')
def procurement_reports():
    if 'user_id' not in session:
        return redirect(url_for('login_page'))
    return render_template('procurement_reports.html', user=session.get('full_name'), role=session.get('role'))


# =============================================================================
# FINANCE / GL API
# =============================================================================

# ── Chart of Accounts ────────────────────────────────────────────────────────

@app.route('/api/gl/accounts', methods=['GET'])
def get_gl_accounts():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        return jsonify(STANDARD_COA)
    accts = list(db.gl_accounts.find({'is_active': {'$ne': False}}).sort('code', 1))
    return jsonify([serialize_doc(a) for a in accts])

@app.route('/api/gl/accounts', methods=['POST'])
def add_gl_account():
    if session.get('role') not in ['admin', 'finance_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.get_json() or {}
    code = data.get('code', '').strip()
    name = data.get('name', '').strip()
    if not code or not name:
        return jsonify({'error': 'Code and name are required'}), 400
    if not USE_MEMORY_DB and db.gl_accounts.find_one({'code': code}):
        return jsonify({'error': 'Account code already exists'}), 400
    acct = {
        'code': code, 'name': name,
        'type': data.get('type', 'asset'),
        'subtype': data.get('subtype', ''),
        'normal_balance': data.get('normal_balance', 'debit'),
        'description': data.get('description', ''),
        'is_active': True,
        'created_at': datetime.utcnow().isoformat()
    }
    if USE_MEMORY_DB:
        return jsonify(acct), 201
    result = db.gl_accounts.insert_one(acct)
    acct['_id'] = str(result.inserted_id)
    return jsonify(acct), 201

@app.route('/api/gl/accounts/<acct_id>', methods=['PUT'])
def update_gl_account(acct_id):
    if session.get('role') not in ['admin', 'finance_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.get_json() or {}
    update = {}
    for f in ['name', 'type', 'subtype', 'normal_balance', 'description', 'is_active']:
        if f in data:
            update[f] = data[f]
    if not update:
        return jsonify({'error': 'Nothing to update'}), 400
    if USE_MEMORY_DB:
        return jsonify({'success': True})
    db.gl_accounts.update_one({'_id': ObjectId(acct_id)}, {'$set': update})
    return jsonify(serialize_doc(db.gl_accounts.find_one({'_id': ObjectId(acct_id)})))

@app.route('/api/gl/accounts/<acct_id>', methods=['DELETE'])
def delete_gl_account(acct_id):
    if session.get('role') not in ['admin', 'finance_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    if USE_MEMORY_DB:
        return jsonify({'success': True})
    db.gl_accounts.update_one({'_id': ObjectId(acct_id)}, {'$set': {'is_active': False}})
    return jsonify({'success': True})

# ── Journal Entries ──────────────────────────────────────────────────────────

@app.route('/api/gl/journal', methods=['GET'])
def get_journal_entries():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        return jsonify([])
    date_from = request.args.get('from')
    date_to   = request.args.get('to')
    ref_type  = request.args.get('type')
    filt = {}
    if date_from or date_to:
        df = {}
        if date_from: df['$gte'] = date_from
        if date_to:   df['$lte'] = date_to
        filt['date'] = df
    if ref_type:
        filt['reference_type'] = ref_type
    entries = list(db.journal_entries.find(filt).sort('created_at', -1).limit(500))
    return jsonify([serialize_doc(e) for e in entries])

@app.route('/api/gl/journal', methods=['POST'])
def post_manual_journal():
    if session.get('role') not in ['admin', 'finance_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.get_json() or {}
    description = data.get('description', '').strip()
    lines = data.get('lines', [])
    if not description:
        return jsonify({'error': 'Description required'}), 400
    if len(lines) < 2:
        return jsonify({'error': 'At least 2 journal lines required'}), 400
    je_id = post_journal_entry(
        description=description,
        lines=lines,
        reference_type='manual',
        entry_date=data.get('date')
    )
    if not je_id:
        return jsonify({'error': 'Journal imbalanced — DR must equal CR'}), 400
    return jsonify({'success': True, 'id': je_id}), 201

@app.route('/api/gl/journal/<je_id>', methods=['GET'])
def get_journal_entry(je_id):
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        return jsonify({'error': 'Not available in memory mode'}), 404
    entry = db.journal_entries.find_one({'_id': ObjectId(je_id)})
    if not entry:
        return jsonify({'error': 'Not found'}), 404
    return jsonify(serialize_doc(entry))

# ── General Ledger (per account with running balance) ────────────────────────

@app.route('/api/gl/ledger', methods=['GET'])
def get_ledger():
    """Return all journal lines grouped by account with running balance."""
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        return jsonify([])
    code      = request.args.get('account')
    date_from = request.args.get('from')
    date_to   = request.args.get('to')
    match_stage = {}
    if date_from or date_to:
        df = {}
        if date_from: df['$gte'] = date_from
        if date_to:   df['$lte'] = date_to
        match_stage['date'] = df
    if code:
        match_stage['lines.account_code'] = code
    pipeline = [
        {'$match': match_stage},
        {'$unwind': '$lines'},
        {'$match': {'lines.account_code': code} if code else {'lines.account_code': {'$exists': True}}},
        {'$project': {
            'entry_number': 1, 'date': 1, 'description': 1,
            'reference_type': 1, 'reference_number': 1,
            'account_code': '$lines.account_code',
            'account_name': '$lines.account_name',
            'dr': '$lines.dr', 'cr': '$lines.cr'
        }},
        {'$sort': {'date': 1}}
    ]
    rows = list(db.journal_entries.aggregate(pipeline))
    running = 0
    result = []
    acct = db.gl_accounts.find_one({'code': code}) if code else None
    normal = acct.get('normal_balance', 'debit') if acct else 'debit'
    for r in rows:
        dr = float(r.get('dr', 0))
        cr = float(r.get('cr', 0))
        if normal == 'debit':
            running += dr - cr
        else:
            running += cr - dr
        result.append({
            'entry_number': r.get('entry_number', ''),
            'date': r.get('date', ''),
            'description': r.get('description', ''),
            'reference_type': r.get('reference_type', ''),
            'reference_number': r.get('reference_number', ''),
            'account_code': r.get('account_code', ''),
            'account_name': r.get('account_name', ''),
            'dr': round(dr, 2),
            'cr': round(cr, 2),
            'balance': round(running, 2)
        })
    return jsonify(result)

# ── Financial Reports ─────────────────────────────────────────────────────────

@app.route('/api/finance/trial-balance', methods=['GET'])
def trial_balance():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        return jsonify([])
    date_from = request.args.get('from')
    date_to   = request.args.get('to')
    accounts = list(db.gl_accounts.find({'is_active': {'$ne': False}}).sort('code', 1))
    result = []
    total_dr = total_cr = 0
    for acct in accounts:
        code = acct['code']
        filt = {'lines.account_code': code}
        if date_from or date_to:
            df = {}
            if date_from: df['$gte'] = date_from
            if date_to:   df['$lte'] = date_to
            filt['date'] = df
        pipeline = [
            {'$match': filt},
            {'$unwind': '$lines'},
            {'$match': {'lines.account_code': code}},
            {'$group': {'_id': None, 'dr': {'$sum': '$lines.dr'}, 'cr': {'$sum': '$lines.cr'}}}
        ]
        res = list(db.journal_entries.aggregate(pipeline))
        if not res:
            continue
        dr = round(res[0]['dr'], 2)
        cr = round(res[0]['cr'], 2)
        if acct.get('normal_balance') == 'debit':
            bal = dr - cr
            row_dr = round(bal, 2) if bal >= 0 else 0
            row_cr = round(-bal, 2) if bal < 0 else 0
        else:
            bal = cr - dr
            row_dr = round(-bal, 2) if bal < 0 else 0
            row_cr = round(bal, 2) if bal >= 0 else 0
        total_dr += row_dr
        total_cr += row_cr
        result.append({
            'code': code, 'name': acct['name'], 'type': acct['type'],
            'dr': row_dr, 'cr': row_cr
        })
    return jsonify({'rows': result, 'total_dr': round(total_dr, 2), 'total_cr': round(total_cr, 2)})

@app.route('/api/finance/income-statement', methods=['GET'])
def income_statement():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        return jsonify({})
    date_from = request.args.get('from')
    date_to   = request.args.get('to')

    def account_net(code_prefix, credit_normal=False):
        filt = {'lines.account_code': {'$regex': f'^{code_prefix}'}}
        if date_from or date_to:
            df = {}
            if date_from: df['$gte'] = date_from
            if date_to: df['$lte'] = date_to
            filt['date'] = df
        pipeline = [
            {'$match': filt},
            {'$unwind': '$lines'},
            {'$match': {'lines.account_code': {'$regex': f'^{code_prefix}'}}},
            {'$group': {'_id': None, 'dr': {'$sum': '$lines.dr'}, 'cr': {'$sum': '$lines.cr'}}}
        ]
        res = list(db.journal_entries.aggregate(pipeline))
        if not res: return 0
        dr = res[0]['dr']; cr = res[0]['cr']
        return round(cr - dr if credit_normal else dr - cr, 2)

    sales_revenue  = account_net('4000', credit_normal=True)
    other_income   = account_net('4100', credit_normal=True)
    total_revenue  = round(sales_revenue + other_income, 2)

    cogs           = account_net('5000')
    purchase_ret   = account_net('5100', credit_normal=True)
    net_cogs       = round(cogs - purchase_ret, 2)

    gross_profit   = round(total_revenue - net_cogs, 2)

    expense_accounts = list(db.gl_accounts.find({'type': 'expense', 'is_active': {'$ne': False}}).sort('code', 1))
    expenses = []
    total_expenses = 0
    for ea in expense_accounts:
        net = account_net(ea['code'])
        if net != 0:
            expenses.append({'code': ea['code'], 'name': ea['name'], 'amount': net})
            total_expenses += net
    total_expenses = round(total_expenses, 2)

    net_profit = round(gross_profit - total_expenses, 2)

    return jsonify({
        'revenue': {
            'sales_revenue': sales_revenue,
            'other_income': other_income,
            'total': total_revenue
        },
        'cogs': {
            'gross_cogs': cogs,
            'purchase_returns': purchase_ret,
            'net_cogs': net_cogs
        },
        'gross_profit': gross_profit,
        'expenses': {
            'items': expenses,
            'total': total_expenses
        },
        'net_profit': net_profit
    })

@app.route('/api/finance/cash-flow', methods=['GET'])
def cash_flow():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        return jsonify({})
    date_from = request.args.get('from')
    date_to   = request.args.get('to')

    def account_change(code_prefix, credit_normal=False):
        filt = {'lines.account_code': {'$regex': f'^{code_prefix}'}}
        if date_from or date_to:
            df = {}
            if date_from: df['$gte'] = date_from
            if date_to: df['$lte'] = date_to
            filt['date'] = df
        pipeline = [
            {'$match': filt},
            {'$unwind': '$lines'},
            {'$match': {'lines.account_code': {'$regex': f'^{code_prefix}'}}},
            {'$group': {'_id': None, 'dr': {'$sum': '$lines.dr'}, 'cr': {'$sum': '$lines.cr'}}}
        ]
        res = list(db.journal_entries.aggregate(pipeline))
        if not res: return 0
        dr = res[0]['dr']; cr = res[0]['cr']
        return round(cr - dr if credit_normal else dr - cr, 2)

    # Get net profit via income_statement logic inline
    def account_net_cf(code_prefix, credit_normal=False):
        filt = {'lines.account_code': {'$regex': f'^{code_prefix}'}}
        if date_from or date_to:
            df = {}
            if date_from: df['$gte'] = date_from
            if date_to: df['$lte'] = date_to
            filt['date'] = df
        pipeline = [
            {'$match': filt},
            {'$unwind': '$lines'},
            {'$match': {'lines.account_code': {'$regex': f'^{code_prefix}'}}},
            {'$group': {'_id': None, 'dr': {'$sum': '$lines.dr'}, 'cr': {'$sum': '$lines.cr'}}}
        ]
        res = list(db.journal_entries.aggregate(pipeline))
        if not res: return 0
        dr = res[0]['dr']; cr = res[0]['cr']
        return round(cr - dr if credit_normal else dr - cr, 2)

    revenue = account_net_cf('4', credit_normal=True)
    cogs = account_net_cf('5')
    expenses_t = account_net_cf('6')
    net_profit = round(revenue - cogs - expenses_t, 2)

    inv_change     = account_change('1200')
    ar_change      = account_change('1100')
    ap_change      = account_change('2000', credit_normal=True)
    accruals_change= account_change('2100', credit_normal=True)
    operating_cf   = round(net_profit - inv_change - ar_change + ap_change + accruals_change, 2)

    fixed_change   = -account_change('1400')
    investing_cf   = round(fixed_change, 2)

    loans_change   = account_change('2300', credit_normal=True) + account_change('2400', credit_normal=True)
    capital_change = account_change('3000', credit_normal=True)
    drawings       = -account_change('3200')
    financing_cf   = round(loans_change + capital_change + drawings, 2)

    cash_pipeline = [
        {'$match': {'lines.account_code': {'$in': ['1000','1010']}}},
        {'$unwind': '$lines'},
        {'$match': {'lines.account_code': {'$in': ['1000','1010']}}},
        {'$group': {'_id': None, 'dr': {'$sum': '$lines.dr'}, 'cr': {'$sum': '$lines.cr'}}}
    ]
    cash_res = list(db.journal_entries.aggregate(cash_pipeline))
    cash_net = 0
    if cash_res:
        cash_net = round(cash_res[0]['dr'] - cash_res[0]['cr'], 2)

    net_change = round(operating_cf + investing_cf + financing_cf, 2)

    return jsonify({
        'operating': {
            'net_profit': net_profit,
            'adjustments': {
                'inventory_change': -inv_change,
                'ar_change': -ar_change,
                'ap_change': ap_change,
                'accruals_change': accruals_change
            },
            'total': operating_cf
        },
        'investing': {
            'fixed_asset_purchases': fixed_change,
            'total': investing_cf
        },
        'financing': {
            'loans_net': loans_change,
            'capital_injected': capital_change,
            'drawings': drawings,
            'total': financing_cf
        },
        'net_change': net_change,
        'closing_cash': round(cash_net, 2)
    })

@app.route('/api/finance/equity-statement', methods=['GET'])
def equity_statement():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        return jsonify({})
    date_from = request.args.get('from')
    date_to   = request.args.get('to')

    def credit_net(code_prefix):
        filt = {'lines.account_code': {'$regex': f'^{code_prefix}'}}
        if date_from or date_to:
            df = {}
            if date_from: df['$gte'] = date_from
            if date_to: df['$lte'] = date_to
            filt['date'] = df
        pipeline = [
            {'$match': filt},
            {'$unwind': '$lines'},
            {'$match': {'lines.account_code': {'$regex': f'^{code_prefix}'}}},
            {'$group': {'_id': None, 'dr': {'$sum': '$lines.dr'}, 'cr': {'$sum': '$lines.cr'}}}
        ]
        res = list(db.journal_entries.aggregate(pipeline))
        if not res: return 0
        return round(res[0]['cr'] - res[0]['dr'], 2)

    def account_net_eq(code_prefix, credit_normal=False):
        filt = {'lines.account_code': {'$regex': f'^{code_prefix}'}}
        if date_from or date_to:
            df = {}
            if date_from: df['$gte'] = date_from
            if date_to: df['$lte'] = date_to
            filt['date'] = df
        pipeline = [
            {'$match': filt},
            {'$unwind': '$lines'},
            {'$match': {'lines.account_code': {'$regex': f'^{code_prefix}'}}},
            {'$group': {'_id': None, 'dr': {'$sum': '$lines.dr'}, 'cr': {'$sum': '$lines.cr'}}}
        ]
        res = list(db.journal_entries.aggregate(pipeline))
        if not res: return 0
        dr = res[0]['dr']; cr = res[0]['cr']
        return round(cr - dr if credit_normal else dr - cr, 2)

    opening_capital   = credit_net('3000')
    retained_earnings = credit_net('3100')
    drawings          = -credit_net('3200')
    revenue = account_net_eq('4', credit_normal=True)
    cogs = account_net_eq('5')
    expenses_t = account_net_eq('6')
    net_profit = round(revenue - cogs - expenses_t, 2)
    opening_equity    = round(opening_capital + retained_earnings, 2)
    closing_equity    = round(opening_equity + net_profit - drawings, 2)

    return jsonify({
        'opening_capital': opening_capital,
        'retained_earnings': retained_earnings,
        'opening_equity': opening_equity,
        'net_profit': net_profit,
        'drawings': drawings,
        'closing_equity': closing_equity
    })

@app.route('/api/finance/dashboard-kpis', methods=['GET'])
def finance_dashboard_kpis():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        return jsonify({'revenue':0,'cogs':0,'gross_profit':0,'net_profit':0,'cash':0,'ar':0,'ap':0,'inventory':0,'journal_count':0})
    today = datetime.utcnow()
    month_from = today.strftime('%Y-%m-01')
    month_to   = today.strftime('%Y-%m-%d')

    def net(code_prefix, credit_normal=False, df=None, dt=None):
        filt = {'lines.account_code': {'$regex': f'^{code_prefix}'}}
        if df or dt:
            drange = {}
            if df: drange['$gte'] = df
            if dt: drange['$lte'] = dt
            filt['date'] = drange
        pipeline = [
            {'$match': filt},
            {'$unwind': '$lines'},
            {'$match': {'lines.account_code': {'$regex': f'^{code_prefix}'}}},
            {'$group': {'_id': None, 'dr': {'$sum': '$lines.dr'}, 'cr': {'$sum': '$lines.cr'}}}
        ]
        res = list(db.journal_entries.aggregate(pipeline))
        if not res: return 0
        dr = res[0]['dr']; cr = res[0]['cr']
        return round(cr - dr if credit_normal else dr - cr, 2)

    revenue     = net('4', credit_normal=True, df=month_from, dt=month_to)
    cogs        = net('5', df=month_from, dt=month_to)
    gross       = round(revenue - cogs, 2)
    expenses_t  = net('6', df=month_from, dt=month_to)
    net_profit  = round(gross - expenses_t, 2)
    cash        = net('1000') + net('1010')
    ar          = net('1100')
    ap          = net('2000', credit_normal=True)
    inventory   = net('1200')
    je_count    = db.journal_entries.count_documents({})

    return jsonify({
        'revenue': revenue, 'cogs': cogs, 'gross_profit': gross,
        'net_profit': net_profit, 'cash': cash, 'ar': ar, 'ap': ap,
        'inventory': inventory, 'journal_count': je_count,
        'period': f"{month_from} to {month_to}"
    })


# =============================================================================
# FINANCE — SAMPLE DATA SEED
# =============================================================================

@app.route('/api/finance/seed-sample', methods=['POST'])
def seed_sample_transactions():
    """Post one complete procurement cycle as sample GL journal entries."""
    if session.get('role') not in ['admin', 'finance_manager']:
        return jsonify({'error': 'Unauthorized'}), 403
    if USE_MEMORY_DB:
        return jsonify({'error': 'MongoDB required for GL posting'}), 400

    today = datetime.utcnow()
    d = lambda offset: (today.replace(day=today.day - offset) if today.day > offset
                        else today).strftime('%Y-%m-%d')

    posted = []

    # ── Step 1: Capital Injection (Owner invests PKR 500,000) ────────────────
    je1 = post_journal_entry(
        description='Owner capital injection — initial investment',
        lines=[
            {'account_code': '1010', 'account_name': 'Bank Account',        'dr': 500000, 'cr': 0},
            {'account_code': '3000', 'account_name': "Owner's Capital",      'dr': 0,      'cr': 500000},
        ],
        reference_type='manual', reference_number='DEMO-001',
        entry_date=d(10)
    )
    posted.append({'step': 1, 'desc': 'Capital Injection', 'je_id': je1})

    # ── Step 2: Inventory Purchase via PO → GR Confirmed ────────────────────
    # Vendor supplies goods worth PKR 120,000
    je2 = post_journal_entry(
        description='GR Confirmed — GR-DEMO-001 (Goods received from vendor)',
        lines=[
            {'account_code': '1200', 'account_name': 'Inventory',           'dr': 120000, 'cr': 0},
            {'account_code': '2000', 'account_name': 'Accounts Payable',    'dr': 0,      'cr': 120000},
        ],
        reference_type='gr', reference_number='GR-DEMO-001',
        entry_date=d(8)
    )
    posted.append({'step': 2, 'desc': 'GR Confirmation → DR Inventory / CR AP', 'je_id': je2})

    # ── Step 3: Vendor Invoice received (tax 17% on purchase) ───────────────
    tax_on_purchase = round(120000 * 0.17, 2)
    je3 = post_journal_entry(
        description='Vendor Invoice — INV-DEMO-001 (GST 17% on purchase)',
        lines=[
            {'account_code': '2100', 'account_name': 'Accrued Liabilities', 'dr': 0,           'cr': tax_on_purchase},
            {'account_code': '1200', 'account_name': 'Inventory',           'dr': tax_on_purchase, 'cr': 0},
        ],
        reference_type='manual', reference_number='INV-DEMO-001',
        entry_date=d(7)
    )
    posted.append({'step': 3, 'desc': 'Vendor Tax Invoice — GST added to Inventory cost', 'je_id': je3})

    # ── Step 4: Vendor Payment (full payment via bank) ───────────────────────
    je4 = post_journal_entry(
        description='Vendor Payment — PAY-DEMO-001 to supplier (full settlement)',
        lines=[
            {'account_code': '2000', 'account_name': 'Accounts Payable',   'dr': 120000, 'cr': 0},
            {'account_code': '1010', 'account_name': 'Bank Account',       'dr': 0,      'cr': 120000},
        ],
        reference_type='payment', reference_number='PAY-DEMO-001',
        entry_date=d(5)
    )
    posted.append({'step': 4, 'desc': 'Vendor Payment → DR AP / CR Bank', 'je_id': je4})

    # ── Step 5: POS Sale (sell goods for PKR 175,000 cash) ───────────────────
    sale_revenue  = 175000
    cogs_amount   = 100000   # cost of the goods sold (portion of inventory)
    je5 = post_journal_entry(
        description='POS Sale — SALE-DEMO-001 (cash sales)',
        lines=[
            {'account_code': '1000', 'account_name': 'Cash & Cash Equivalents', 'dr': sale_revenue, 'cr': 0},
            {'account_code': '4000', 'account_name': 'Sales Revenue',           'dr': 0, 'cr': sale_revenue},
        ],
        reference_type='sale', reference_number='SALE-DEMO-001',
        entry_date=d(3)
    )
    posted.append({'step': 5, 'desc': 'POS Sale → DR Cash / CR Revenue', 'je_id': je5})

    # ── Step 6: COGS Recognition (goods delivered to customers) ─────────────
    je6 = post_journal_entry(
        description='COGS — SALE-DEMO-001 (cost of goods sold recognised)',
        lines=[
            {'account_code': '5000', 'account_name': 'Cost of Goods Sold', 'dr': cogs_amount, 'cr': 0},
            {'account_code': '1200', 'account_name': 'Inventory',          'dr': 0,           'cr': cogs_amount},
        ],
        reference_type='sale', reference_number='SALE-DEMO-001',
        entry_date=d(3)
    )
    posted.append({'step': 6, 'desc': 'COGS Recognition → DR COGS / CR Inventory', 'je_id': je6})

    # ── Step 7: Operating Expenses (rent + salaries) ─────────────────────────
    je7 = post_journal_entry(
        description='Monthly operating expenses — Rent PKR 25,000 + Salaries PKR 45,000',
        lines=[
            {'account_code': '6100', 'account_name': 'Rent Expense',       'dr': 25000, 'cr': 0},
            {'account_code': '6000', 'account_name': 'Salaries & Wages',   'dr': 45000, 'cr': 0},
            {'account_code': '1010', 'account_name': 'Bank Account',       'dr': 0,     'cr': 70000},
        ],
        reference_type='manual', reference_number='EXP-DEMO-001',
        entry_date=d(1)
    )
    posted.append({'step': 7, 'desc': 'Operating Expenses → DR Rent+Salaries / CR Bank', 'je_id': je7})

    # ── Summary ──────────────────────────────────────────────────────────────
    successful = [p for p in posted if p['je_id']]
    failed     = [p for p in posted if not p['je_id']]

    return jsonify({
        'success': True,
        'posted': len(successful),
        'failed': len(failed),
        'entries': posted,
        'summary': {
            'capital_injected':   500000,
            'inventory_purchased': 120000,
            'sales_revenue':       sale_revenue,
            'cogs':                cogs_amount,
            'gross_profit':        sale_revenue - cogs_amount,
            'operating_expenses':  70000,
            'net_profit':          sale_revenue - cogs_amount - 70000,
        }
    }), 201


# =============================================================================
# PHASE 4 — POS FEATURES
# =============================================================================

def get_next_shift_number():
    if USE_MEMORY_DB:
        return f"SHF-{str(len(getattr(db,'shifts',[]))+1).zfill(4)}"
    last = db.shifts.find_one(sort=[('_id', -1)])
    if last and last.get('shift_number'):
        try: return f"SHF-{str(int(last['shift_number'].split('-')[1])+1).zfill(4)}"
        except: pass
    return "SHF-0001"

def get_next_customer_number():
    if USE_MEMORY_DB:
        return f"CUST-{str(len(getattr(db,'customers',[]))+1).zfill(4)}"
    last = db.customers.find_one(sort=[('_id', -1)])
    if last and last.get('customer_number'):
        try: return f"CUST-{str(int(last['customer_number'].split('-')[1])+1).zfill(4)}"
        except: pass
    return "CUST-0001"


# ── Store Configuration ──────────────────────────────────────────────────────

STORE_DEFAULTS = {
    'store_name': 'My Supermarket', 'address': '123 Main Street, City',
    'phone': '+92-XXX-XXXXXXX', 'email': 'store@example.com', 'ntn': '',
    'tax_rate': 0, 'loyalty_rate': 1, 'loyalty_redeem_rate': 0.5,
    'receipt_footer': 'Thank you for shopping with us!', 'currency': 'PKR'
}

@app.route('/api/store-config', methods=['GET'])
def get_store_config():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        return jsonify(getattr(db, 'store_config', STORE_DEFAULTS.copy()))
    cfg = db.store_config.find_one({})
    if cfg: cfg.pop('_id', None); return jsonify(cfg)
    return jsonify(STORE_DEFAULTS.copy())

@app.route('/api/store-config', methods=['PUT'])
def update_store_config():
    if session.get('role') != 'admin': return jsonify({'error': 'Unauthorized'}), 403
    data = request.json or {}
    allowed = ['store_name','address','phone','email','ntn','tax_rate','loyalty_rate','loyalty_redeem_rate','receipt_footer','currency']
    update = {k: data[k] for k in allowed if k in data}
    if USE_MEMORY_DB: db.store_config = update; return jsonify({'success': True})
    db.store_config.replace_one({}, update, upsert=True)
    return jsonify({'success': True})


# ── Customers ────────────────────────────────────────────────────────────────

@app.route('/api/customers', methods=['GET'])
def get_customers():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    search = request.args.get('search', '')
    limit  = min(int(request.args.get('limit', 50)), 200)
    if USE_MEMORY_DB:
        custs = getattr(db, 'customers', [])
        if search:
            s = search.lower()
            custs = [c for c in custs if s in (c.get('name','') + c.get('phone','')).lower()]
        return jsonify([{**dict(c), '_id': str(c.get('_id',''))} for c in custs[:limit]])
    query = {}
    if search:
        query['$or'] = [{'name': {'$regex': search, '$options': 'i'}},
                        {'phone': {'$regex': search, '$options': 'i'}},
                        {'customer_number': {'$regex': search, '$options': 'i'}}]
    custs = list(db.customers.find(query).limit(limit))
    for c in custs: c['_id'] = str(c['_id'])
    return jsonify(custs)

@app.route('/api/customers', methods=['POST'])
def create_customer():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    data = request.json or {}
    if not data.get('name'): return jsonify({'error': 'Customer name required'}), 400
    cust = {
        'customer_number': get_next_customer_number(),
        'name': data.get('name','').strip(), 'phone': data.get('phone','').strip(),
        'email': data.get('email','').strip(), 'address': data.get('address','').strip(),
        'loyalty_points': 0, 'total_spent': 0, 'visit_count': 0,
        'created_at': datetime.utcnow().isoformat(), 'notes': data.get('notes','').strip(),
    }
    if USE_MEMORY_DB:
        cust['_id'] = get_next_id('customers')
        if not hasattr(db, 'customers'): db.customers = []
        db.customers.append(cust); cust['_id'] = str(cust['_id'])
        return jsonify({'success': True, 'customer': cust}), 201
    result = db.customers.insert_one(cust); cust['_id'] = str(result.inserted_id)
    return jsonify({'success': True, 'customer': cust}), 201

@app.route('/api/customers/<cid>', methods=['GET'])
def get_customer(cid):
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        c = next((x for x in getattr(db,'customers',[]) if str(x.get('_id'))==cid), None)
        if not c: return jsonify({'error': 'Not found'}), 404
        return jsonify({**dict(c), '_id': str(c['_id'])})
    try: c = db.customers.find_one({'_id': ObjectId(cid)})
    except: c = None
    if not c: return jsonify({'error': 'Not found'}), 404
    c['_id'] = str(c['_id']); return jsonify(c)

@app.route('/api/customers/<cid>', methods=['PUT'])
def update_customer(cid):
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    data = request.json or {}
    upd = {k: data[k] for k in ['name','phone','email','address','notes'] if k in data}
    if USE_MEMORY_DB:
        for c in getattr(db,'customers',[]):
            if str(c.get('_id'))==cid: c.update(upd); return jsonify({'success': True})
        return jsonify({'error': 'Not found'}), 404
    try: db.customers.update_one({'_id': ObjectId(cid)}, {'$set': upd})
    except: return jsonify({'error': 'Invalid ID'}), 400
    return jsonify({'success': True})

@app.route('/api/customers/<cid>/history', methods=['GET'])
def customer_history(cid):
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    limit = min(int(request.args.get('limit', 20)), 100)
    if USE_MEMORY_DB:
        sales = sorted([s for s in getattr(db,'sales',[]) if str(s.get('customer_id'))==cid],
                       key=lambda x: x.get('timestamp',''), reverse=True)[:limit]
        out = []
        for s in sales:
            sc = dict(s); sc['_id'] = str(sc.get('_id',''))
            if isinstance(sc.get('timestamp'), datetime): sc['timestamp'] = sc['timestamp'].isoformat()
            out.append(sc)
        return jsonify(out)
    sales = list(db.sales.find({'customer_id': cid}).sort('timestamp', -1).limit(limit))
    for s in sales:
        s['_id'] = str(s['_id'])
        if isinstance(s.get('timestamp'), datetime): s['timestamp'] = s['timestamp'].isoformat()
    return jsonify(sales)


@app.route('/api/customers/<cid>', methods=['DELETE'])
def delete_customer(cid):
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        db.customers = [c for c in getattr(db, 'customers', []) if str(c.get('_id', '')) != cid]
        return jsonify({'message': 'Customer deleted'})
    try:
        result = db.customers.delete_one({'_id': ObjectId(cid)})
        if result.deleted_count == 0:
            return jsonify({'message': 'Customer not found'}), 404
        return jsonify({'message': 'Customer deleted'})
    except Exception as e:
        return jsonify({'message': str(e)}), 500


# ── Promo Codes ──────────────────────────────────────────────────────────────

@app.route('/api/promo-codes', methods=['GET'])
def get_promo_codes():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        return jsonify([{**dict(p), '_id': str(p.get('_id',''))} for p in getattr(db,'promo_codes',[])])
    promos = list(db.promo_codes.find({}))
    for p in promos: p['_id'] = str(p['_id'])
    return jsonify(promos)

@app.route('/api/promo-codes', methods=['POST'])
def create_promo_code():
    if session.get('role') != 'admin': return jsonify({'error': 'Unauthorized'}), 403
    data = request.json or {}
    code = (data.get('code','') or '').strip().upper()
    if not code: return jsonify({'error': 'Promo code required'}), 400
    promo = {
        'code': code, 'type': data.get('type','percent'),
        'value': float(data.get('value', 0)), 'min_purchase': float(data.get('min_purchase', 0)),
        'max_uses': int(data.get('max_uses', 0)), 'uses': 0,
        'expiry': data.get('expiry',''), 'active': bool(data.get('active', True)),
        'description': data.get('description','').strip(), 'created_at': datetime.utcnow().isoformat(),
    }
    if USE_MEMORY_DB:
        promo['_id'] = get_next_id('promo_codes')
        if not hasattr(db,'promo_codes'): db.promo_codes = []
        db.promo_codes.append(promo); promo['_id'] = str(promo['_id'])
        return jsonify({'success': True, 'promo': promo}), 201
    result = db.promo_codes.insert_one(promo); promo['_id'] = str(result.inserted_id)
    return jsonify({'success': True, 'promo': promo}), 201

@app.route('/api/promo-codes/<pid>', methods=['PUT'])
def update_promo_code(pid):
    if session.get('role') != 'admin': return jsonify({'error': 'Unauthorized'}), 403
    data = request.json or {}
    upd = {k: data[k] for k in ['code','type','value','min_purchase','max_uses','expiry','active','description'] if k in data}
    if 'code' in upd: upd['code'] = upd['code'].strip().upper()
    if 'value' in upd: upd['value'] = float(upd['value'])
    if 'min_purchase' in upd: upd['min_purchase'] = float(upd['min_purchase'])
    if 'max_uses' in upd: upd['max_uses'] = int(upd['max_uses'])
    if USE_MEMORY_DB:
        for p in getattr(db,'promo_codes',[]):
            if str(p.get('_id'))==pid: p.update(upd); return jsonify({'success': True})
        return jsonify({'error': 'Not found'}), 404
    try: db.promo_codes.update_one({'_id': ObjectId(pid)}, {'$set': upd})
    except: return jsonify({'error': 'Invalid ID'}), 400
    return jsonify({'success': True})

@app.route('/api/promo-codes/<pid>', methods=['DELETE'])
def delete_promo_code(pid):
    if session.get('role') != 'admin': return jsonify({'error': 'Unauthorized'}), 403
    if USE_MEMORY_DB:
        if hasattr(db,'promo_codes'): db.promo_codes = [p for p in db.promo_codes if str(p.get('_id'))!=pid]
        return jsonify({'success': True})
    try: db.promo_codes.delete_one({'_id': ObjectId(pid)})
    except: return jsonify({'error': 'Invalid ID'}), 400
    return jsonify({'success': True})

@app.route('/api/promo-codes/validate', methods=['POST'])
def validate_promo_code():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    data = request.json or {}
    code = (data.get('code','') or '').strip().upper()
    cart_total = float(data.get('cart_total', 0))
    if not code: return jsonify({'error': 'No code provided'}), 400
    now_str = datetime.utcnow().date().isoformat()
    if USE_MEMORY_DB:
        promo = next((p for p in getattr(db,'promo_codes',[]) if p.get('code')==code), None)
    else:
        promo = db.promo_codes.find_one({'code': code})
    if not promo: return jsonify({'valid': False, 'error': 'Invalid promo code'}), 400
    if not promo.get('active', True): return jsonify({'valid': False, 'error': 'This promo code is inactive'}), 400
    if promo.get('expiry') and promo['expiry'] < now_str: return jsonify({'valid': False, 'error': 'Promo code has expired'}), 400
    max_uses = promo.get('max_uses', 0)
    if max_uses > 0 and promo.get('uses', 0) >= max_uses: return jsonify({'valid': False, 'error': 'Usage limit reached'}), 400
    min_pur = promo.get('min_purchase', 0)
    if cart_total < min_pur: return jsonify({'valid': False, 'error': f'Minimum purchase PKR {min_pur:.0f} required'}), 400
    discount = round(cart_total * promo['value'] / 100, 2) if promo['type']=='percent' else min(float(promo['value']), cart_total)
    return jsonify({'valid': True, 'discount': discount, 'type': promo['type'], 'value': promo['value'],
                    'code': code, 'description': promo.get('description','')})


# ── Shifts ───────────────────────────────────────────────────────────────────

@app.route('/api/shifts', methods=['GET'])
def get_shifts():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    limit  = min(int(request.args.get('limit', 20)), 100)
    status = request.args.get('status', '')
    if USE_MEMORY_DB:
        shifts = getattr(db, 'shifts', [])
        if status: shifts = [s for s in shifts if s.get('status')==status]
        shifts = sorted(shifts, key=lambda x: x.get('opened_at',''), reverse=True)[:limit]
        return jsonify([{**dict(s), '_id': str(s.get('_id',''))} for s in shifts])
    q = {}
    if status: q['status'] = status
    shifts = list(db.shifts.find(q).sort('opened_at', -1).limit(limit))
    for s in shifts: s['_id'] = str(s['_id'])
    return jsonify(shifts)

@app.route('/api/shifts', methods=['POST'])
def open_shift():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    if session.get('role') not in ['admin','sales_person','cashier']:
        return jsonify({'error': 'Unauthorized'}), 403
    cid = session['user_id']
    if USE_MEMORY_DB:
        existing = next((s for s in getattr(db,'shifts',[]) if s.get('cashier_id')==cid and s.get('status')=='open'), None)
    else:
        existing = db.shifts.find_one({'cashier_id': cid, 'status': 'open'})
    if existing:
        e = dict(existing); e['_id'] = str(e['_id'])
        return jsonify({'error': 'You already have an open shift', 'shift': e}), 400
    data = request.json or {}
    shift = {
        'shift_number': get_next_shift_number(), 'cashier_id': cid,
        'cashier_name': session.get('full_name',''), 'status': 'open',
        'opened_at': datetime.utcnow().isoformat(), 'closed_at': None,
        'opening_float': float(data.get('opening_float', 0)),
        'closing_float': None, 'expected_cash': None, 'cash_over_short': None,
        'notes': data.get('notes','').strip(), 'close_notes': '',
        'sales_count': 0, 'sales_total': 0, 'discount_total': 0,
        'payment_breakdown': {'cash': 0, 'card': 0, 'online': 0, 'credit': 0, 'split': 0},
    }
    if USE_MEMORY_DB:
        shift['_id'] = get_next_id('shifts')
        if not hasattr(db, 'shifts'): db.shifts = []
        db.shifts.append(shift); shift['_id'] = str(shift['_id'])
        return jsonify({'success': True, 'shift': shift}), 201
    result = db.shifts.insert_one(shift); shift['_id'] = str(result.inserted_id)
    return jsonify({'success': True, 'shift': shift}), 201

@app.route('/api/shifts/active', methods=['GET'])
def get_active_shift():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    cid = session['user_id']
    if USE_MEMORY_DB:
        s = next((x for x in getattr(db,'shifts',[]) if x.get('cashier_id')==cid and x.get('status')=='open'), None)
        if not s: return jsonify({'shift': None})
        return jsonify({'shift': {**dict(s), '_id': str(s['_id'])}})
    s = db.shifts.find_one({'cashier_id': cid, 'status': 'open'})
    if not s: return jsonify({'shift': None})
    s['_id'] = str(s['_id']); return jsonify({'shift': s})

@app.route('/api/shifts/<sid>/close', methods=['POST'])
def close_shift(sid):
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    data = request.json or {}
    closing_float = float(data.get('closing_float', 0))
    if USE_MEMORY_DB:
        shift = next((s for s in getattr(db,'shifts',[]) if str(s.get('_id'))==sid), None)
    else:
        try: shift = db.shifts.find_one({'_id': ObjectId(sid)})
        except: shift = None
    if not shift: return jsonify({'error': 'Shift not found'}), 404
    if shift.get('status') != 'open': return jsonify({'error': 'Shift is not open'}), 400
    cash_sales    = shift.get('payment_breakdown', {}).get('cash', 0)
    expected_cash = shift.get('opening_float', 0) + cash_sales
    over_short_val = round(closing_float - expected_cash, 2)
    upd = {'status': 'closed', 'closed_at': datetime.utcnow().isoformat(),
           'closing_float': closing_float, 'expected_cash': expected_cash,
           'over_short': over_short_val, 'cash_over_short': over_short_val,
           'close_notes': data.get('notes', '')}
    if USE_MEMORY_DB:
        shift.update(upd); return jsonify({'success': True, 'shift': {**dict(shift), '_id': str(shift['_id'])}})
    db.shifts.update_one({'_id': ObjectId(sid)}, {'$set': upd})
    shift.update(upd); shift['_id'] = str(shift['_id'])
    return jsonify({'success': True, 'shift': shift})

@app.route('/api/shifts/<sid>/summary', methods=['GET'])
def shift_summary(sid):
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        s = next((x for x in getattr(db,'shifts',[]) if str(x.get('_id'))==sid), None)
    else:
        try: s = db.shifts.find_one({'_id': ObjectId(sid)})
        except: s = None
    if not s: return jsonify({'error': 'Not found'}), 404
    sc = dict(s); sc['_id'] = str(sc['_id']); return jsonify(sc)


@app.route('/api/shifts/cashier-kpis', methods=['GET'])
def cashier_kpis():
    """Per-cashier KPIs for today — used by sales manager and admin dashboards."""
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    today_start = datetime.utcnow().replace(hour=0, minute=0, second=0, microsecond=0).isoformat()

    if USE_MEMORY_DB:
        shifts = getattr(db, 'shifts', [])
        today_shifts = [s for s in shifts if s.get('opened_at', '') >= today_start]
    else:
        today_shifts = list(db.shifts.find({'opened_at': {'$gte': today_start}}))
        for s in today_shifts:
            s['_id'] = str(s['_id'])

    # Group by cashier
    by_cashier = {}
    for s in today_shifts:
        cid = s.get('cashier_id') or s.get('cashier_name', 'Unknown')
        if cid not in by_cashier:
            by_cashier[cid] = {
                'cashier_id': cid,
                'cashier_name': s.get('cashier_name', 'Unknown'),
                'shifts': 0, 'open_shifts': 0,
                'transactions': 0, 'total_sales': 0.0,
                'cash_sales': 0.0, 'over_short': 0.0,
                'current_shift_id': None, 'current_shift_opened': None
            }
        c = by_cashier[cid]
        c['shifts'] += 1
        c['transactions'] += s.get('transaction_count', 0)
        c['total_sales'] += s.get('total_sales', 0)
        c['cash_sales'] += (s.get('payment_breakdown') or {}).get('cash', 0)
        if s.get('status') == 'open':
            c['open_shifts'] += 1
            c['current_shift_id'] = str(s.get('_id', ''))
            c['current_shift_opened'] = s.get('opened_at')
        if s.get('status') == 'closed' and s.get('cash_over_short') is not None:
            c['over_short'] += s.get('cash_over_short', 0)

    result = sorted(by_cashier.values(), key=lambda x: x['total_sales'], reverse=True)
    return jsonify(result)


# =============================================================================
# HR MODULE
# =============================================================================

def _next_emp_num():
    if USE_MEMORY_DB:
        nums = [int(e.get('emp_id','EMP-0000').split('-')[-1]) for e in getattr(db,'employees',[]) if e.get('emp_id','').startswith('EMP-')]
        return f'EMP-{(max(nums,default=0)+1):04d}'
    last = db.employees.find_one({'emp_id':{'$regex':'^EMP-'}}, sort=[('emp_id',-1)])
    return f'EMP-{(int(last["emp_id"].split("-")[-1])+1 if last else 1):04d}'


def _next_payroll_num(month):
    prefix = f'PAY-{month}'
    if USE_MEMORY_DB:
        count = sum(1 for p in getattr(db,'payroll',[]) if p.get('payroll_id','').startswith(prefix))
        return f'{prefix}-{(count+1):03d}'
    count = db.payroll.count_documents({'payroll_id':{'$regex':f'^{prefix}'}})
    return f'{prefix}-{(count+1):03d}'


def _next_pmt_num():
    if USE_MEMORY_DB:
        nums = [int(p.get('payment_id','PMT-0000').split('-')[-1]) for p in getattr(db,'payments',[]) if p.get('payment_id','').startswith('PMT-')]
        return f'PMT-{(max(nums,default=0)+1):04d}'
    last = db.payments.find_one({'payment_id':{'$regex':'^PMT-'}}, sort=[('payment_id',-1)])
    return f'PMT-{(int(last["payment_id"].split("-")[-1])+1 if last else 1):04d}'


def _emp_salary_totals(data):
    allowances = data.get('allowances', [])
    deductions  = data.get('deductions', [])
    allowances_total = round(sum(float(a.get('amount',0)) for a in allowances), 2)
    deductions_total = round(sum(float(d.get('amount',0)) for d in deductions), 2)
    basic  = round(float(data.get('basic_salary', 0)), 2)
    gross  = round(basic + allowances_total, 2)
    net    = round(gross - deductions_total, 2)
    return allowances, deductions, allowances_total, deductions_total, basic, gross, net


# ── Employees ─────────────────────────────────────────────────────────────────

@app.route('/api/hr/employees', methods=['GET'])
def get_employees():
    if 'user_id' not in session: return jsonify({'error':'Unauthorized'}), 401
    dept   = request.args.get('department','')
    status = request.args.get('status','')
    if USE_MEMORY_DB:
        emps = list(getattr(db,'employees',[]))
        if dept:   emps = [e for e in emps if e.get('department')==dept]
        if status: emps = [e for e in emps if e.get('status')==status]
        return jsonify([{**dict(e),'_id':str(e.get('_id',''))} for e in emps])
    q = {}
    if dept:   q['department'] = dept
    if status: q['status'] = status
    emps = list(db.employees.find(q).sort('name',1))
    for e in emps: e['_id'] = str(e['_id'])
    return jsonify(emps)


@app.route('/api/hr/employees', methods=['POST'])
def create_employee():
    if 'user_id' not in session: return jsonify({'error':'Unauthorized'}), 401
    if session.get('role') != 'admin': return jsonify({'error':'Forbidden'}), 403
    data = request.json or {}
    name = data.get('name','').strip()
    if not name: return jsonify({'message':'Employee name is required'}), 400
    al, ded, al_tot, ded_tot, basic, gross, net = _emp_salary_totals(data)
    doc = {
        'emp_id': _next_emp_num(),
        'name': name,
        'designation': data.get('designation',''),
        'department':  data.get('department',''),
        'phone':       data.get('phone',''),
        'email':       data.get('email',''),
        'cnic':        data.get('cnic',''),
        'joining_date':data.get('joining_date',''),
        'salary_type': data.get('salary_type','monthly'),
        'basic_salary':basic, 'allowances':al, 'deductions':ded,
        'allowances_total':al_tot, 'deductions_total':ded_tot,
        'gross_salary':gross, 'net_salary':net,
        'bank_name':   data.get('bank_name',''),
        'bank_account':data.get('bank_account',''),
        'status':      'active',
        'created_at':  datetime.utcnow().isoformat(),
        'created_by':  session.get('full_name','System')
    }
    if USE_MEMORY_DB:
        doc['_id'] = str(uuid.uuid4())
        if not hasattr(db,'employees'): db.employees = []
        db.employees.append(doc)
        return jsonify({'success':True,'employee':doc}), 201
    result = db.employees.insert_one(doc)
    doc['_id'] = str(result.inserted_id)
    return jsonify({'success':True,'employee':doc}), 201


@app.route('/api/hr/employees/<eid>', methods=['GET'])
def get_employee(eid):
    if 'user_id' not in session: return jsonify({'error':'Unauthorized'}), 401
    if USE_MEMORY_DB:
        e = next((x for x in getattr(db,'employees',[]) if str(x.get('_id'))==eid or x.get('emp_id')==eid), None)
    else:
        try: e = db.employees.find_one({'_id':ObjectId(eid)})
        except: e = db.employees.find_one({'emp_id':eid})
    if not e: return jsonify({'error':'Not found'}), 404
    ec = dict(e); ec['_id'] = str(ec.get('_id','')); return jsonify(ec)


@app.route('/api/hr/employees/<eid>', methods=['PUT'])
def update_employee(eid):
    if 'user_id' not in session: return jsonify({'error':'Unauthorized'}), 401
    if session.get('role') != 'admin': return jsonify({'error':'Forbidden'}), 403
    data = request.json or {}
    al, ded, al_tot, ded_tot, basic, gross, net = _emp_salary_totals(data)
    upd = {
        'name':        data.get('name','').strip(),
        'designation': data.get('designation',''),
        'department':  data.get('department',''),
        'phone':       data.get('phone',''),
        'email':       data.get('email',''),
        'cnic':        data.get('cnic',''),
        'joining_date':data.get('joining_date',''),
        'salary_type': data.get('salary_type','monthly'),
        'basic_salary':basic, 'allowances':al, 'deductions':ded,
        'allowances_total':al_tot, 'deductions_total':ded_tot,
        'gross_salary':gross, 'net_salary':net,
        'bank_name':   data.get('bank_name',''),
        'bank_account':data.get('bank_account',''),
        'status':      data.get('status','active'),
        'updated_at':  datetime.utcnow().isoformat()
    }
    if USE_MEMORY_DB:
        e = next((x for x in getattr(db,'employees',[]) if str(x.get('_id'))==eid), None)
        if not e: return jsonify({'error':'Not found'}), 404
        e.update(upd); return jsonify({'success':True})
    db.employees.update_one({'_id':ObjectId(eid)}, {'$set':upd})
    return jsonify({'success':True})


@app.route('/api/hr/employees/<eid>', methods=['DELETE'])
def delete_employee_hr(eid):
    if 'user_id' not in session: return jsonify({'error':'Unauthorized'}), 401
    if session.get('role') != 'admin': return jsonify({'error':'Forbidden'}), 403
    if USE_MEMORY_DB:
        db.employees = [e for e in getattr(db,'employees',[]) if str(e.get('_id'))!=eid]
        return jsonify({'success':True})
    db.employees.delete_one({'_id':ObjectId(eid)})
    return jsonify({'success':True})


@app.route('/api/hr/departments', methods=['GET'])
def get_departments_hr():
    if 'user_id' not in session: return jsonify({'error':'Unauthorized'}), 401
    default_depts = ['Management','Sales','Procurement','Finance','IT','Operations','HR','Warehouse']
    if USE_MEMORY_DB:
        from_emps = sorted(set(e.get('department','') for e in getattr(db,'employees',[]) if e.get('department')))
        return jsonify(sorted(set(default_depts + from_emps)))
    from_emps = db.employees.distinct('department')
    return jsonify(sorted(set(default_depts + [d for d in from_emps if d])))


@app.route('/api/hr/stats', methods=['GET'])
def hr_stats():
    if 'user_id' not in session: return jsonify({'error':'Unauthorized'}), 401
    if USE_MEMORY_DB:
        emps    = getattr(db,'employees',[])
        active  = [e for e in emps if e.get('status')=='active']
        payrolls= getattr(db,'payroll',[])
        pending_pay  = len([p for p in payrolls if p.get('status')=='pending'])
        return jsonify({
            'total_employees':  len(emps),
            'active_employees': len(active),
            'monthly_payroll_cost': sum(e.get('net_salary',0) for e in active),
            'pending_payroll':  pending_pay,
            'departments': list(set(e.get('department','') for e in active if e.get('department')))
        })
    total  = db.employees.count_documents({})
    active = db.employees.count_documents({'status':'active'})
    cost_r = list(db.employees.aggregate([{'$match':{'status':'active'}},{'$group':{'_id':None,'t':{'$sum':'$net_salary'}}}]))
    return jsonify({
        'total_employees':  total,
        'active_employees': active,
        'monthly_payroll_cost': cost_r[0]['t'] if cost_r else 0,
        'pending_payroll':  db.payroll.count_documents({'status':'pending'}),
        'departments': db.employees.distinct('department', {'status':'active'})
    })


# ── Payroll ───────────────────────────────────────────────────────────────────

@app.route('/api/hr/payroll', methods=['GET'])
def get_payroll():
    if 'user_id' not in session: return jsonify({'error':'Unauthorized'}), 401
    month  = request.args.get('month','')
    emp_id = request.args.get('employee_id','')
    status = request.args.get('status','')
    if USE_MEMORY_DB:
        recs = list(getattr(db,'payroll',[]))
        if month:  recs = [r for r in recs if r.get('month','').startswith(month)]
        if emp_id: recs = [r for r in recs if str(r.get('employee_id'))==emp_id]
        if status: recs = [r for r in recs if r.get('status')==status]
        return jsonify([{**dict(r),'_id':str(r.get('_id',''))} for r in sorted(recs, key=lambda x:x.get('month',''), reverse=True)])
    q = {}
    if month:  q['month'] = {'$regex':f'^{month}'}
    if emp_id: q['employee_id'] = emp_id
    if status: q['status'] = status
    recs = list(db.payroll.find(q).sort('month',-1))
    for r in recs: r['_id'] = str(r['_id'])
    return jsonify(recs)


@app.route('/api/hr/payroll/generate', methods=['POST'])
def generate_payroll():
    if 'user_id' not in session: return jsonify({'error':'Unauthorized'}), 401
    if session.get('role') != 'admin': return jsonify({'error':'Forbidden'}), 403
    data  = request.json or {}
    month = data.get('month','')
    if not month: return jsonify({'message':'Month is required (YYYY-MM)'}), 400
    if USE_MEMORY_DB:
        employees = [e for e in getattr(db,'employees',[]) if e.get('status')=='active']
        existing  = [p for p in getattr(db,'payroll',[]) if p.get('month')==month]
    else:
        employees = list(db.employees.find({'status':'active'}))
        existing  = list(db.payroll.find({'month':month}))
    if existing:
        return jsonify({'message':f'Payroll for {month} already generated ({len(existing)} records). Delete them first to regenerate.'}), 400
    if not employees:
        return jsonify({'message':'No active employees found'}), 400
    created = []
    for emp in employees:
        emp_oid = str(emp.get('_id',''))
        doc = {
            'payroll_id':   _next_payroll_num(month),
            'month':         month,
            'employee_id':   emp_oid,
            'employee_name': emp.get('name',''),
            'emp_number':    emp.get('emp_id',''),
            'designation':   emp.get('designation',''),
            'department':    emp.get('department',''),
            'basic_salary':  emp.get('basic_salary',0),
            'allowances':    emp.get('allowances',[]),
            'allowances_total': emp.get('allowances_total',0),
            'deductions':    emp.get('deductions',[]),
            'deductions_total': emp.get('deductions_total',0),
            'gross_salary':  emp.get('gross_salary',0),
            'net_salary':    emp.get('net_salary',0),
            'bank_name':     emp.get('bank_name',''),
            'bank_account':  emp.get('bank_account',''),
            'status':        'pending',
            'payment_id':    None,
            'created_at':    datetime.utcnow().isoformat(),
            'created_by':    session.get('full_name','System')
        }
        if USE_MEMORY_DB:
            doc['_id'] = str(uuid.uuid4())
            if not hasattr(db,'payroll'): db.payroll = []
            db.payroll.append(doc)
        else:
            r = db.payroll.insert_one(doc); doc['_id'] = str(r.inserted_id)
        created.append(doc)
    return jsonify({'success':True,'count':len(created),'total_net':sum(d['net_salary'] for d in created),'month':month}), 201


@app.route('/api/hr/payroll/<pid>/pay', methods=['POST'])
def pay_payroll_record(pid):
    if 'user_id' not in session: return jsonify({'error':'Unauthorized'}), 401
    if session.get('role') != 'admin': return jsonify({'error':'Forbidden'}), 403
    data = request.json or {}
    if USE_MEMORY_DB:
        pr = next((x for x in getattr(db,'payroll',[]) if str(x.get('_id'))==pid), None)
    else:
        try: pr = db.payroll.find_one({'_id':ObjectId(pid)})
        except: pr = None
    if not pr:  return jsonify({'error':'Payroll record not found'}), 404
    if pr.get('status')=='paid': return jsonify({'error':'Already paid'}), 400
    pmt_id = _next_pmt_num()
    pmt_doc = {
        'payment_id':   pmt_id,
        'payment_type': 'employee',
        'payee_id':     str(pr.get('employee_id','')),
        'payee_name':   pr.get('employee_name',''),
        'amount':       pr.get('net_salary',0),
        'payment_method': data.get('payment_method','bank_transfer'),
        'reference':    data.get('reference',''),
        'description':  f'Salary — {pr.get("employee_name","")} — {pr.get("month","")}',
        'category':     'salary',
        'related_payroll_id': pid,
        'related_invoice_id': '',
        'status':       'completed',
        'created_by':   session.get('full_name','System'),
        'created_at':   datetime.utcnow().isoformat()
    }
    if USE_MEMORY_DB:
        pmt_doc['_id'] = str(uuid.uuid4())
        if not hasattr(db,'payments'): db.payments = []
        db.payments.append(pmt_doc)
        pr.update({'status':'paid','payment_id':pmt_doc['_id'],'paid_at':datetime.utcnow().isoformat()})
    else:
        r = db.payments.insert_one(pmt_doc); pmt_doc['_id'] = str(r.inserted_id)
        db.payroll.update_one({'_id':pr['_id']}, {'$set':{'status':'paid','payment_id':str(r.inserted_id),'paid_at':datetime.utcnow().isoformat()}})
    return jsonify({'success':True,'payment':pmt_doc})


@app.route('/api/hr/payroll/bulk-pay', methods=['POST'])
def bulk_pay_payroll():
    """Pay all pending payroll for a given month."""
    if 'user_id' not in session: return jsonify({'error':'Unauthorized'}), 401
    if session.get('role') != 'admin': return jsonify({'error':'Forbidden'}), 403
    data  = request.json or {}
    month = data.get('month','')
    method= data.get('payment_method','bank_transfer')
    if not month: return jsonify({'message':'Month required'}), 400
    if USE_MEMORY_DB:
        pending = [p for p in getattr(db,'payroll',[]) if p.get('month')==month and p.get('status')=='pending']
    else:
        pending = list(db.payroll.find({'month':month,'status':'pending'}))
    if not pending: return jsonify({'message':'No pending payroll for this month'}), 400
    paid_count = 0
    total_paid = 0
    for pr in pending:
        pid = str(pr.get('_id',''))
        pmt_doc = {
            'payment_id':   _next_pmt_num(),
            'payment_type': 'employee',
            'payee_id':     str(pr.get('employee_id','')),
            'payee_name':   pr.get('employee_name',''),
            'amount':       pr.get('net_salary',0),
            'payment_method': method,
            'reference':    data.get('reference',''),
            'description':  f'Bulk salary — {pr.get("employee_name","")} — {month}',
            'category':     'salary',
            'related_payroll_id': pid,
            'related_invoice_id': '',
            'status':       'completed',
            'created_by':   session.get('full_name','System'),
            'created_at':   datetime.utcnow().isoformat()
        }
        if USE_MEMORY_DB:
            pmt_doc['_id'] = str(uuid.uuid4())
            if not hasattr(db,'payments'): db.payments = []
            db.payments.append(pmt_doc)
            pr.update({'status':'paid','payment_id':pmt_doc['_id'],'paid_at':datetime.utcnow().isoformat()})
        else:
            r = db.payments.insert_one(pmt_doc)
            db.payroll.update_one({'_id':pr['_id']}, {'$set':{'status':'paid','payment_id':str(r.inserted_id),'paid_at':datetime.utcnow().isoformat()}})
        paid_count += 1
        total_paid += pr.get('net_salary',0)
    return jsonify({'success':True,'paid_count':paid_count,'total_paid':total_paid})


# =============================================================================
# PAYMENTS MODULE
# =============================================================================

def _pmt_normalize(p):
    """Normalize payment doc to frontend field names."""
    d = dict(p)
    d['_id']    = str(d.get('_id',''))
    d['type']   = d.get('type') or d.pop('payment_type', 'other')
    d['payee']  = d.get('payee') or d.pop('payee_name', '')
    d['method'] = d.get('method') or d.pop('payment_method', 'cash')
    d['date']   = (d.get('date') or d.get('created_at',''))[:10]
    return d


@app.route('/api/payments', methods=['GET'])
def pmt_list():
    if 'user_id' not in session: return jsonify({'error':'Unauthorized'}), 401
    ptype   = request.args.get('type','')
    status  = request.args.get('status','')
    method  = request.args.get('method','')
    dfrom   = request.args.get('from','')
    dto     = request.args.get('to','')
    if USE_MEMORY_DB:
        pmts = list(getattr(db,'payments',[]))
        if ptype:  pmts = [p for p in pmts if (p.get('type') or p.get('payment_type'))==ptype]
        if status: pmts = [p for p in pmts if p.get('status')==status]
        if method: pmts = [p for p in pmts if (p.get('method') or p.get('payment_method'))==method]
        if dfrom:  pmts = [p for p in pmts if p.get('created_at','')>=dfrom]
        if dto:    pmts = [p for p in pmts if p.get('created_at','')<=dto+'T23:59:59']
        pmts = sorted(pmts, key=lambda x:x.get('created_at',''), reverse=True)
        return jsonify({'payments':[_pmt_normalize(p) for p in pmts]})
    q = {}
    if ptype:  q['$or'] = [{'type':ptype},{'payment_type':ptype}]
    if status: q['status'] = status
    if method: q['$or'] = q.get('$or',[]) + [{'method':method},{'payment_method':method}]
    if dfrom or dto:
        q['created_at'] = {}
        if dfrom: q['created_at']['$gte'] = dfrom
        if dto:   q['created_at']['$lte'] = dto+'T23:59:59'
    pmts = list(db.payments.find(q).sort('created_at',-1))
    return jsonify({'payments':[_pmt_normalize(p) for p in pmts]})


@app.route('/api/payments', methods=['POST'])
def pmt_create():
    if 'user_id' not in session: return jsonify({'error':'Unauthorized'}), 401
    if session.get('role') != 'admin': return jsonify({'error':'Forbidden'}), 403
    data   = request.json or {}
    ptype  = data.get('type', data.get('payment_type','other'))
    amount = round(float(data.get('amount',0)), 2)
    if not amount: return jsonify({'error':'Amount is required'}), 400
    payee  = (data.get('payee') or data.get('payee_name','')).strip()
    doc = {
        'payment_id':  _next_pmt_num(),
        'type':        ptype,
        'payee':       payee,
        'amount':      amount,
        'method':      data.get('method', data.get('payment_method','cash')),
        'date':        data.get('date', datetime.utcnow().strftime('%Y-%m-%d')),
        'reference':   data.get('reference','').strip(),
        'notes':       data.get('notes','').strip(),
        'category':    data.get('category','other'),
        'description': data.get('description','').strip(),
        'employee_id': data.get('employee_id',''),
        'payroll_id':  data.get('payroll_id',''),
        'vendor_id':   data.get('vendor_id',''),
        'invoice_id':  data.get('invoice_id',''),
        'status':      data.get('status','completed'),
        'created_by':  session.get('full_name','System'),
        'created_at':  datetime.utcnow().isoformat()
    }
    # Mark linked records paid
    if doc['invoice_id'] and ptype == 'vendor':
        _mark_invoice_paid(doc['invoice_id'], doc['payment_id'])
    if doc['payroll_id'] and ptype == 'employee':
        _mark_payroll_paid(doc['payroll_id'], doc['payment_id'])
    if USE_MEMORY_DB:
        doc['_id'] = str(uuid.uuid4())
        if not hasattr(db,'payments'): db.payments = []
        db.payments.append(doc)
        return jsonify({'success':True,'payment':_pmt_normalize(doc)}), 201
    result = db.payments.insert_one(doc)
    doc['_id'] = str(result.inserted_id)
    return jsonify({'success':True,'payment':_pmt_normalize(doc)}), 201


def _mark_invoice_paid(inv_id, pmt_id):
    upd = {'status':'paid','payment_id':pmt_id,'paid_at':datetime.utcnow().isoformat()}
    if USE_MEMORY_DB:
        inv = next((x for x in getattr(db,'vendor_invoices',[]) if str(x.get('_id'))==inv_id), None)
        if inv: inv.update(upd)
        inv2 = next((x for x in getattr(db,'invoices',[]) if str(x.get('_id'))==inv_id), None)
        if inv2: inv2.update(upd)
    else:
        try: db.vendor_invoices.update_one({'_id':ObjectId(inv_id)}, {'$set':upd})
        except: pass
        try: db.invoices.update_one({'_id':ObjectId(inv_id)}, {'$set':upd})
        except: pass


def _mark_payroll_paid(pr_id, pmt_id):
    upd = {'status':'paid','payment_id':pmt_id,'paid_at':datetime.utcnow().isoformat()}
    if USE_MEMORY_DB:
        pr = next((x for x in getattr(db,'payroll',[]) if str(x.get('_id'))==pr_id), None)
        if pr: pr.update(upd)
    else:
        try: db.payroll.update_one({'_id':ObjectId(pr_id)}, {'$set':upd})
        except: pass


@app.route('/api/payments/<pid>', methods=['GET'])
def get_payment_detail(pid):
    if 'user_id' not in session: return jsonify({'error':'Unauthorized'}), 401
    if USE_MEMORY_DB:
        p = next((x for x in getattr(db,'payments',[]) if str(x.get('_id'))==pid or x.get('payment_id')==pid), None)
    else:
        try: p = db.payments.find_one({'_id':ObjectId(pid)})
        except: p = db.payments.find_one({'payment_id':pid})
    if not p: return jsonify({'error':'Not found'}), 404
    pc = dict(p); pc['_id'] = str(pc.get('_id','')); return jsonify(pc)


@app.route('/api/payments/stats', methods=['GET'])
def pmt_stats():
    if 'user_id' not in session: return jsonify({'error':'Unauthorized'}), 401
    now         = datetime.utcnow()
    month_start = now.replace(day=1,hour=0,minute=0,second=0,microsecond=0).isoformat()
    year_start  = now.replace(month=1,day=1,hour=0,minute=0,second=0,microsecond=0).isoformat()
    def _ptype(p): return p.get('type') or p.get('payment_type','other')
    if USE_MEMORY_DB:
        all_pmts   = list(getattr(db,'payments',[]))
        this_month = [p for p in all_pmts if p.get('created_at','')>=month_start]
        this_year  = [p for p in all_pmts if p.get('created_at','')>=year_start]
        emp_mo  = [p for p in this_month if _ptype(p)=='employee']
        vnd_mo  = [p for p in this_month if _ptype(p)=='vendor']
        oth_mo  = [p for p in this_month if _ptype(p)=='other']
        return jsonify({
            'this_month': {
                'total': sum(p.get('amount',0) for p in this_month),
                'count': len(this_month),
                'employee': sum(p.get('amount',0) for p in emp_mo),
                'employee_count': len(emp_mo),
                'vendor': sum(p.get('amount',0) for p in vnd_mo),
                'vendor_count': len(vnd_mo),
                'other': sum(p.get('amount',0) for p in oth_mo),
                'other_count': len(oth_mo),
            },
            'ytd': {'total': sum(p.get('amount',0) for p in this_year)}
        })
    pipe_mo = [
        {'$match':{'created_at':{'$gte':month_start}}},
        {'$group':{'_id':{'$ifNull':['$type','$payment_type']},'total':{'$sum':'$amount'},'count':{'$sum':1}}}
    ]
    mo_data  = {r['_id']:(r['total'],r['count']) for r in db.payments.aggregate(pipe_mo)}
    ytd_pipe = [{'$match':{'created_at':{'$gte':year_start}}},{'$group':{'_id':None,'total':{'$sum':'$amount'}}}]
    ytd_res  = list(db.payments.aggregate(ytd_pipe))
    def _mt(k): return mo_data.get(k,(0,0))[0]
    def _mc(k): return mo_data.get(k,(0,0))[1]
    return jsonify({
        'this_month': {
            'total': sum(v[0] for v in mo_data.values()),
            'count': sum(v[1] for v in mo_data.values()),
            'employee': _mt('employee'), 'employee_count': _mc('employee'),
            'vendor':   _mt('vendor'),   'vendor_count':   _mc('vendor'),
            'other':    _mt('other'),    'other_count':    _mc('other'),
        },
        'ytd': {'total': ytd_res[0]['total'] if ytd_res else 0}
    })


@app.route('/api/payments/report', methods=['GET'])
def pmt_report():
    if 'user_id' not in session: return jsonify({'error':'Unauthorized'}), 401
    def _ptype(p): return p.get('type') or p.get('payment_type','other')
    now        = datetime.utcnow()
    year_start = now.replace(month=1,day=1,hour=0,minute=0,second=0,microsecond=0).isoformat()
    if USE_MEMORY_DB:
        all_pmts = sorted(getattr(db,'payments',[]), key=lambda x:x.get('created_at',''))
        year_pmts = [p for p in all_pmts if p.get('created_at','')>=year_start]
        by_mo = {}
        for p in year_pmts:
            mo = p.get('created_at','')[:7]; pt = _ptype(p)
            if not mo: continue
            if mo not in by_mo: by_mo[mo] = {'month':mo,'employee':0,'vendor':0,'other':0}
            by_mo[mo][pt] = by_mo[mo].get(pt,0) + p.get('amount',0)
        emp_pmts = [p for p in all_pmts if _ptype(p)=='employee']
        vnd_pmts = [p for p in all_pmts if _ptype(p)=='vendor']
        oth_pmts = [p for p in all_pmts if _ptype(p)=='other']
        return jsonify({
            'monthly': sorted(by_mo.values(), key=lambda x:x['month']),
            'summary': {
                'employee_total': sum(p.get('amount',0) for p in emp_pmts),
                'employee_count': len(emp_pmts),
                'vendor_total':   sum(p.get('amount',0) for p in vnd_pmts),
                'vendor_count':   len(vnd_pmts),
                'other_total':    sum(p.get('amount',0) for p in oth_pmts),
                'other_count':    len(oth_pmts),
            }
        })
    pipe = [
        {'$match':{'created_at':{'$gte':year_start}}},
        {'$group':{'_id':{'mo':{'$substr':['$created_at',0,7]},'t':{'$ifNull':['$type','$payment_type']}},
                   'total':{'$sum':'$amount'}}}
    ]
    by_mo = {}
    for r in db.payments.aggregate(pipe):
        mo=r['_id']['mo']; pt=r['_id']['t']
        if mo not in by_mo: by_mo[mo]={'month':mo,'employee':0,'vendor':0,'other':0}
        by_mo[mo][pt]=r['total']
    sum_pipe = [{'$group':{'_id':{'$ifNull':['$type','$payment_type']},'total':{'$sum':'$amount'},'count':{'$sum':1}}}]
    sum_data = {r['_id']:(r['total'],r['count']) for r in db.payments.aggregate(sum_pipe)}
    def _st(k): return sum_data.get(k,(0,0))[0]
    def _sc(k): return sum_data.get(k,(0,0))[1]
    return jsonify({
        'monthly': sorted(by_mo.values(), key=lambda x:x['month']),
        'summary': {
            'employee_total': _st('employee'), 'employee_count': _sc('employee'),
            'vendor_total':   _st('vendor'),   'vendor_count':   _sc('vendor'),
            'other_total':    _st('other'),    'other_count':    _sc('other'),
        }
    })


# ── HR & Payments page routes ─────────────────────────────────────────────────

@app.route('/hr')
def hr_page():
    if 'user_id' not in session: return redirect(url_for('login_page'))
    if session.get('role') != 'admin': return redirect_to_dashboard(session.get('role'))
    return render_template('hr_dashboard.html')


@app.route('/payments')
def payments_page():
    if 'user_id' not in session: return redirect(url_for('login_page'))
    if session.get('role') != 'admin': return redirect_to_dashboard(session.get('role'))
    return render_template('payments.html')


# ── POS Customer & Shift pages ───────────────────────────────────────────────

@app.route('/pos/customers')
def pos_customers_page():
    if 'user_id' not in session: return redirect(url_for('login_page'))
    if session.get('role') not in ['admin', 'sales_manager', 'sales_person']:
        return redirect_to_dashboard(session.get('role'))
    return render_template('pos_customers.html', user=session.get('full_name'), role=session.get('role'))

@app.route('/pos/shifts')
def pos_shifts_page():
    if 'user_id' not in session: return redirect(url_for('login_page'))
    if session.get('role') not in ['admin', 'sales_manager', 'sales_person']:
        return redirect_to_dashboard(session.get('role'))
    return render_template('pos_shifts.html', user=session.get('full_name'), role=session.get('role'))


# =============================================================================
# PHASE 3 — ANALYTICS API ENDPOINTS
# =============================================================================

@app.route('/api/analytics/sales-summary', methods=['GET'])
def analytics_sales_summary():
    """Multi-period revenue summary: today, yesterday, this week, this month, last month."""
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    now   = datetime.utcnow()
    today = now.date()
    def period_stats(start, end):
        if USE_MEMORY_DB:
            sales = [s for s in getattr(db, 'sales', [])
                     if isinstance(s.get('timestamp'), datetime) and start <= s['timestamp'] <= end]
            return {'revenue': sum(s.get('total',0) for s in sales), 'transactions': len(sales)}
        r = list(db.sales.aggregate([
            {'$match': {'timestamp': {'$gte': start, '$lte': end}}},
            {'$group': {'_id': None, 'revenue': {'$sum': '$total'}, 'transactions': {'$sum': 1}}}
        ]))
        return {'revenue': r[0]['revenue'] if r else 0, 'transactions': r[0]['transactions'] if r else 0}

    from datetime import timedelta
    week_start  = datetime.combine(today - timedelta(days=today.weekday()), datetime.min.time())
    month_start = datetime.combine(today.replace(day=1), datetime.min.time())
    lm_start    = datetime.combine((today.replace(day=1) - timedelta(days=1)).replace(day=1), datetime.min.time())
    lm_end      = datetime.combine(today.replace(day=1) - timedelta(days=1), datetime.max.time())

    td  = period_stats(datetime.combine(today, datetime.min.time()), datetime.combine(today, datetime.max.time()))
    yd  = period_stats(datetime.combine(today-timedelta(days=1), datetime.min.time()), datetime.combine(today-timedelta(days=1), datetime.max.time()))
    wk  = period_stats(week_start, datetime.combine(today, datetime.max.time()))
    mo  = period_stats(month_start, datetime.combine(today, datetime.max.time()))
    lmo = period_stats(lm_start, lm_end)

    def change(curr, prev):
        if prev == 0: return None
        return round((curr - prev) / prev * 100, 1)

    return jsonify({
        'today':      {**td,  'avg_basket': round(td['revenue']/td['transactions'],2) if td['transactions'] else 0},
        'yesterday':  yd,
        'this_week':  wk,
        'this_month': mo,
        'last_month': lmo,
        'vs_yesterday': change(td['revenue'], yd['revenue']),
        'vs_last_month': change(mo['revenue'], lmo['revenue']),
    })


@app.route('/api/analytics/hourly-sales', methods=['GET'])
def analytics_hourly_sales():
    """Today's sales broken down by hour (0-23)."""
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    today = datetime.utcnow().date()
    day_start = datetime.combine(today, datetime.min.time())
    day_end   = datetime.combine(today, datetime.max.time())
    hours = [{'hour': h, 'label': f'{h:02d}:00', 'revenue': 0, 'transactions': 0} for h in range(24)]
    if USE_MEMORY_DB:
        for s in getattr(db, 'sales', []):
            ts = s.get('timestamp')
            if isinstance(ts, datetime) and day_start <= ts <= day_end:
                h = ts.hour
                hours[h]['revenue']      += s.get('total', 0)
                hours[h]['transactions'] += 1
    else:
        r = list(db.sales.aggregate([
            {'$match': {'timestamp': {'$gte': day_start, '$lte': day_end}}},
            {'$group': {'_id': {'$hour': '$timestamp'}, 'revenue': {'$sum': '$total'}, 'transactions': {'$sum': 1}}}
        ]))
        for row in r:
            h = row['_id']
            if 0 <= h < 24:
                hours[h]['revenue']      = row['revenue']
                hours[h]['transactions'] = row['transactions']
    return jsonify(hours)


@app.route('/api/analytics/payment-methods', methods=['GET'])
def analytics_payment_methods():
    """Payment method breakdown for a given period (default: this month)."""
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    period = request.args.get('period', 'month')
    now   = datetime.utcnow()
    today = now.date()
    from datetime import timedelta
    if period == 'today':
        start = datetime.combine(today, datetime.min.time())
    elif period == 'week':
        start = datetime.combine(today - timedelta(days=today.weekday()), datetime.min.time())
    else:
        start = datetime.combine(today.replace(day=1), datetime.min.time())
    end = datetime.combine(today, datetime.max.time())
    if USE_MEMORY_DB:
        tally = {}
        for s in getattr(db, 'sales', []):
            ts = s.get('timestamp')
            if isinstance(ts, datetime) and start <= ts <= end:
                pm = s.get('payment_method', 'cash').title()
                tally.setdefault(pm, {'revenue': 0, 'count': 0})
                tally[pm]['revenue'] += s.get('total', 0)
                tally[pm]['count']   += 1
        return jsonify([{'method': k, **v} for k, v in tally.items()])
    r = list(db.sales.aggregate([
        {'$match': {'timestamp': {'$gte': start, '$lte': end}}},
        {'$group': {'_id': '$payment_method', 'revenue': {'$sum': '$total'}, 'count': {'$sum': 1}}}
    ]))
    return jsonify([{'method': (x['_id'] or 'cash').title(), 'revenue': x['revenue'], 'count': x['count']} for x in r])


@app.route('/api/analytics/sales-report', methods=['GET'])
def analytics_sales_report():
    """Detailed sales report with date range filter. Returns per-day summary + transaction list."""
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    date_from = request.args.get('from')
    date_to   = request.args.get('to')
    limit     = min(int(request.args.get('limit', 200)), 500)
    try:
        now   = datetime.utcnow()
        today = now.date()
        from datetime import timedelta
        if date_from:
            start = datetime.combine(datetime.strptime(date_from, '%Y-%m-%d').date(), datetime.min.time())
        else:
            start = datetime.combine(today, datetime.min.time())
        if date_to:
            end = datetime.combine(datetime.strptime(date_to, '%Y-%m-%d').date(), datetime.max.time())
        else:
            end = datetime.combine(today, datetime.max.time())
    except Exception:
        return jsonify({'error': 'Invalid date format'}), 400

    if USE_MEMORY_DB:
        sales = [s for s in getattr(db, 'sales', [])
                 if isinstance(s.get('timestamp'), datetime) and start <= s['timestamp'] <= end]
        sales.sort(key=lambda x: x.get('timestamp', datetime.min), reverse=True)
        total_rev  = sum(s.get('total', 0) for s in sales)
        total_txn  = len(sales)
        out = []
        for s in sales[:limit]:
            sc = dict(s)
            sc['_id'] = str(sc.get('_id', ''))
            if isinstance(sc.get('timestamp'), datetime):
                sc['timestamp'] = sc['timestamp'].isoformat()
            out.append(sc)
        return jsonify({'transactions': out, 'total_revenue': total_rev, 'total_transactions': total_txn,
                        'avg_basket': round(total_rev/total_txn, 2) if total_txn else 0})

    match = {'timestamp': {'$gte': start, '$lte': end}}
    agg = list(db.sales.aggregate([{'$match': match}, {'$group': {'_id': None, 'rev': {'$sum': '$total'}, 'cnt': {'$sum': 1}}}]))
    total_rev = agg[0]['rev'] if agg else 0
    total_txn = agg[0]['cnt'] if agg else 0
    sales = list(db.sales.find(match).sort('timestamp', -1).limit(limit))
    for s in sales:
        s['_id'] = str(s['_id'])
        if isinstance(s.get('timestamp'), datetime):
            s['timestamp'] = s['timestamp'].isoformat()
    return jsonify({'transactions': sales, 'total_revenue': total_rev, 'total_transactions': total_txn,
                    'avg_basket': round(total_rev/total_txn, 2) if total_txn else 0})


@app.route('/api/analytics/sales-trend', methods=['GET'])
def analytics_sales_trend():
    """Last 7 days daily sales: revenue + transaction count."""
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    days = []
    for i in range(6, -1, -1):
        day = (datetime.utcnow() - timedelta(days=i)).date()
        day_start = datetime.combine(day, datetime.min.time())
        day_end   = datetime.combine(day, datetime.max.time())
        if USE_MEMORY_DB:
            sales = getattr(db, 'sales', [])
            day_s = [s for s in sales if isinstance(s.get('timestamp'), datetime) and day_start <= s['timestamp'] <= day_end]
            days.append({'date': day.strftime('%d %b'), 'revenue': sum(s.get('total', 0) for s in day_s), 'transactions': len(day_s)})
        else:
            r = list(db.sales.aggregate([
                {'$match': {'timestamp': {'$gte': day_start, '$lte': day_end}}},
                {'$group': {'_id': None, 'total': {'$sum': '$total'}, 'count': {'$sum': 1}}}
            ]))
            days.append({'date': day.strftime('%d %b'), 'revenue': r[0]['total'] if r else 0, 'transactions': r[0]['count'] if r else 0})
    return jsonify(days)


@app.route('/api/analytics/top-products', methods=['GET'])
def analytics_top_products():
    """Top N products by revenue from all sales."""
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    limit = min(int(request.args.get('limit', 10)), 20)
    if USE_MEMORY_DB:
        tally = {}
        for sale in getattr(db, 'sales', []):
            for it in sale.get('items', []):
                name = it.get('product_name', 'Unknown')
                tally.setdefault(name, {'quantity': 0, 'revenue': 0})
                tally[name]['quantity'] += it.get('quantity', 0)
                tally[name]['revenue']  += it.get('subtotal', 0)
        return jsonify(sorted([{'name': k, **v} for k, v in tally.items()], key=lambda x: x['revenue'], reverse=True)[:limit])
    r = list(db.sales.aggregate([
        {'$unwind': '$items'},
        {'$group': {'_id': '$items.product_name', 'quantity': {'$sum': '$items.quantity'}, 'revenue': {'$sum': '$items.subtotal'}}},
        {'$sort': {'revenue': -1}},
        {'$limit': limit}
    ]))
    return jsonify([{'name': x['_id'] or 'Unknown', 'quantity': x['quantity'], 'revenue': x['revenue']} for x in r])


@app.route('/api/analytics/category-breakdown', methods=['GET'])
def analytics_category_breakdown():
    """Sales revenue grouped by product category."""
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        tally = {}
        for sale in getattr(db, 'sales', []):
            for it in sale.get('items', []):
                pid  = it.get('product_id')
                prod = db.find_one('products', {'_id': pid}) if pid else None
                cat  = (prod.get('category') or 'Uncategorized') if prod else 'Uncategorized'
                tally.setdefault(cat, {'revenue': 0, 'quantity': 0})
                tally[cat]['revenue']  += it.get('subtotal', 0)
                tally[cat]['quantity'] += it.get('quantity', 0)
        return jsonify(sorted([{'category': k, **v} for k, v in tally.items()], key=lambda x: x['revenue'], reverse=True)[:10])
    r = list(db.sales.aggregate([
        {'$unwind': '$items'},
        {'$lookup': {'from': 'products', 'localField': 'items.product_id', 'foreignField': '_id', 'as': 'prod'}},
        {'$unwind': {'path': '$prod', 'preserveNullAndEmptyArrays': True}},
        {'$group': {
            '_id': {'$ifNull': ['$prod.category', 'Uncategorized']},
            'revenue': {'$sum': '$items.subtotal'},
            'quantity': {'$sum': '$items.quantity'}
        }},
        {'$sort': {'revenue': -1}},
        {'$limit': 10}
    ]))
    return jsonify([{'category': x['_id'] or 'Uncategorized', 'revenue': x['revenue'], 'quantity': x['quantity']} for x in r])


@app.route('/api/analytics/procurement-spend', methods=['GET'])
def analytics_procurement_spend():
    """PO total spend grouped by vendor (top 10)."""
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        tally = {}
        for po in getattr(db, 'purchase_orders_v2', []):
            v = po.get('vendor_name') or 'Unknown'
            tally.setdefault(v, {'value': 0, 'count': 0})
            tally[v]['value'] += po.get('total_amount', 0)
            tally[v]['count'] += 1
        return jsonify(sorted([{'vendor': k, **v} for k, v in tally.items()], key=lambda x: x['value'], reverse=True)[:10])
    r = list(db.purchase_orders_v2.aggregate([
        {'$group': {'_id': '$vendor_name', 'value': {'$sum': '$total_amount'}, 'count': {'$sum': 1}}},
        {'$sort': {'value': -1}},
        {'$limit': 10}
    ]))
    return jsonify([{'vendor': x['_id'] or 'Unknown', 'value': x['value'], 'count': x['count']} for x in r])


@app.route('/api/analytics/procurement-kpis', methods=['GET'])
def analytics_procurement_kpis():
    """Procurement KPIs: pending pipeline stages, overdue invoices, AP aging buckets."""
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    zero = {'pr_pending': 0, 'po_pending_delivery': 0, 'gr_pending': 0,
            'overdue_invoices': 0, 'overdue_amount': 0, 'total_ap': 0,
            'aging': {'current': 0, '1_30': 0, '31_60': 0, '61_90': 0, '90_plus': 0},
            'aging_amounts': {'current': 0, '1_30': 0, '31_60': 0, '61_90': 0, '90_plus': 0}}
    if USE_MEMORY_DB:
        return jsonify(zero)
    now = datetime.utcnow()
    pr_pending = db.purchase_requisitions.count_documents({'status': 'pending_approval'})
    po_pending = db.purchase_orders_v2.count_documents({'status': 'sent'})
    gr_pending = db.goods_receipts.count_documents({'status': 'pending'})
    unpaid     = list(db.vendor_invoices.find({'status': {'$in': ['pending', 'partial']}}))
    total_ap   = sum(i.get('amount_due', 0) for i in unpaid)
    aging = {'current': 0, '1_30': 0, '31_60': 0, '61_90': 0, '90_plus': 0}
    aging_amounts = {'current': 0, '1_30': 0, '31_60': 0, '61_90': 0, '90_plus': 0}
    overdue_count, overdue_amount = 0, 0
    for inv in unpaid:
        due = inv.get('due_date')
        amt = inv.get('amount_due', 0)
        if not due:
            aging['current'] += 1; aging_amounts['current'] += amt; continue
        due_dt = due if isinstance(due, datetime) else datetime.fromisoformat(str(due))
        days_past = (now - due_dt).days
        if days_past <= 0:
            b = 'current'
        elif days_past <= 30:
            b = '1_30'; overdue_count += 1; overdue_amount += amt
        elif days_past <= 60:
            b = '31_60'; overdue_count += 1; overdue_amount += amt
        elif days_past <= 90:
            b = '61_90'; overdue_count += 1; overdue_amount += amt
        else:
            b = '90_plus'; overdue_count += 1; overdue_amount += amt
        aging[b] += 1; aging_amounts[b] += amt
    return jsonify({
        'pr_pending': pr_pending, 'po_pending_delivery': po_pending, 'gr_pending': gr_pending,
        'overdue_invoices': overdue_count, 'overdue_amount': overdue_amount, 'total_ap': total_ap,
        'aging': aging, 'aging_amounts': aging_amounts,
    })


# =============================================================================
# AUDIT TRAIL
# =============================================================================

def log_audit(action, module, record_id='', details='', user=None):
    """Write an audit entry. Call this from any write operation."""
    try:
        entry = {
            'action':     action,
            'module':     module,
            'record_id':  str(record_id),
            'details':    details,
            'user':       user or session.get('full_name', 'System'),
            'user_id':    session.get('user_id', ''),
            'role':       session.get('role', ''),
            'ip':         request.remote_addr,
            'created_at': datetime.utcnow().isoformat()
        }
        if USE_MEMORY_DB:
            if not hasattr(db, 'audit_logs'): db.audit_logs = []
            db.audit_logs.insert(0, entry)
            if len(db.audit_logs) > 2000: db.audit_logs = db.audit_logs[:2000]
        else:
            db.audit_logs.insert_one(entry)
    except Exception:
        pass  # never let audit failure break the main flow


@app.route('/api/audit-log', methods=['GET'])
def get_audit_log():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    if session.get('role') != 'admin': return jsonify({'error': 'Forbidden'}), 403
    module = request.args.get('module', '')
    action = request.args.get('action', '')
    user   = request.args.get('user', '')
    dfrom  = request.args.get('from', '')
    dto    = request.args.get('to', '')
    page   = int(request.args.get('page', 1))
    per    = 50
    if USE_MEMORY_DB:
        logs = list(getattr(db, 'audit_logs', []))
        if module: logs = [l for l in logs if l.get('module') == module]
        if action: logs = [l for l in logs if action.lower() in l.get('action','').lower()]
        if user:   logs = [l for l in logs if user.lower() in l.get('user','').lower()]
        if dfrom:  logs = [l for l in logs if l.get('created_at','') >= dfrom]
        if dto:    logs = [l for l in logs if l.get('created_at','') <= dto + 'T23:59:59']
        total = len(logs)
        logs  = logs[(page-1)*per : page*per]
        for l in logs: l['_id'] = str(l.get('_id',''))
        return jsonify({'logs': logs, 'total': total, 'page': page, 'pages': -(-total//per)})
    q = {}
    if module: q['module'] = module
    if action: q['action'] = {'$regex': action, '$options': 'i'}
    if user:   q['user']   = {'$regex': user,   '$options': 'i'}
    if dfrom or dto:
        q['created_at'] = {}
        if dfrom: q['created_at']['$gte'] = dfrom
        if dto:   q['created_at']['$lte'] = dto + 'T23:59:59'
    total = db.audit_logs.count_documents(q)
    logs  = list(db.audit_logs.find(q).sort('created_at', -1).skip((page-1)*per).limit(per))
    for l in logs: l['_id'] = str(l['_id'])
    return jsonify({'logs': logs, 'total': total, 'page': page, 'pages': -(-total//per)})


@app.route('/audit-log')
def audit_log_page():
    if 'user_id' not in session: return redirect(url_for('login_page'))
    if session.get('role') != 'admin': return redirect_to_dashboard(session.get('role'))
    return render_template('audit_log.html')


# =============================================================================
# SALES RETURNS & REFUNDS
# =============================================================================

def _next_return_num():
    prefix = f'RET-{datetime.utcnow().strftime("%Y%m")}-'
    if USE_MEMORY_DB:
        nums = [int(r.get('return_id','RET-000000-000').split('-')[-1])
                for r in getattr(db,'sales_returns',[]) if r.get('return_id','').startswith(prefix)]
        return prefix + f'{(max(nums, default=0)+1):03d}'
    last = db.sales_returns.find_one({'return_id': {'$regex': f'^{prefix}'}}, sort=[('return_id', -1)])
    n = int(last['return_id'].split('-')[-1]) + 1 if last else 1
    return prefix + f'{n:03d}'


@app.route('/api/sales/returns', methods=['GET'])
def get_sales_returns():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    status = request.args.get('status', '')
    dfrom  = request.args.get('from', '')
    dto    = request.args.get('to', '')
    if USE_MEMORY_DB:
        rets = list(getattr(db, 'sales_returns', []))
        if status: rets = [r for r in rets if r.get('status') == status]
        if dfrom:  rets = [r for r in rets if r.get('created_at','') >= dfrom]
        if dto:    rets = [r for r in rets if r.get('created_at','') <= dto + 'T23:59:59']
        rets = sorted(rets, key=lambda x: x.get('created_at',''), reverse=True)
        return jsonify({'returns': [{**dict(r), '_id': str(r.get('_id',''))} for r in rets]})
    q = {}
    if status: q['status'] = status
    if dfrom or dto:
        q['created_at'] = {}
        if dfrom: q['created_at']['$gte'] = dfrom
        if dto:   q['created_at']['$lte'] = dto + 'T23:59:59'
    rets = list(db.sales_returns.find(q).sort('created_at', -1))
    for r in rets: r['_id'] = str(r['_id'])
    return jsonify({'returns': rets})


@app.route('/api/sales/returns', methods=['POST'])
def create_sales_return():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    data    = request.json or {}
    items   = data.get('items', [])
    if not items: return jsonify({'error': 'No items provided'}), 400
    total   = round(sum(float(i.get('subtotal', 0)) for i in items), 2)
    doc = {
        'return_id':      _next_return_num(),
        'original_sale_id': data.get('original_sale_id', ''),
        'customer_name':  data.get('customer_name', 'Walk-in'),
        'customer_phone': data.get('customer_phone', ''),
        'items':          items,
        'total_refund':   total,
        'refund_method':  data.get('refund_method', 'cash'),
        'reason':         data.get('reason', ''),
        'notes':          data.get('notes', ''),
        'status':         'completed',
        'processed_by':   session.get('full_name', 'Cashier'),
        'created_at':     datetime.utcnow().isoformat()
    }
    # Restock items
    for item in items:
        pid = item.get('product_id', '')
        qty = int(item.get('qty', 0))
        if pid and qty > 0:
            if USE_MEMORY_DB:
                prod = next((p for p in getattr(db,'products',[]) if str(p.get('_id')) == pid), None)
                if prod: prod['quantity'] = int(prod.get('quantity', 0)) + qty
            else:
                try: db.products.update_one({'_id': ObjectId(pid)}, {'$inc': {'quantity': qty}})
                except: pass
    if USE_MEMORY_DB:
        doc['_id'] = str(uuid.uuid4())
        if not hasattr(db, 'sales_returns'): db.sales_returns = []
        db.sales_returns.append(doc)
        log_audit('RETURN_CREATED', 'Sales Returns', doc['_id'], f"Return {doc['return_id']} — PKR {total}")
        return jsonify({'success': True, 'return': doc}), 201
    result = db.sales_returns.insert_one(doc)
    doc['_id'] = str(result.inserted_id)
    log_audit('RETURN_CREATED', 'Sales Returns', doc['_id'], f"Return {doc['return_id']} — PKR {total}")
    return jsonify({'success': True, 'return': doc}), 201


@app.route('/api/sales/returns/<rid>', methods=['GET'])
def get_sales_return(rid):
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        r = next((x for x in getattr(db,'sales_returns',[]) if str(x.get('_id'))==rid or x.get('return_id')==rid), None)
    else:
        try: r = db.sales_returns.find_one({'_id': ObjectId(rid)})
        except: r = db.sales_returns.find_one({'return_id': rid})
    if not r: return jsonify({'error': 'Not found'}), 404
    rc = dict(r); rc['_id'] = str(rc.get('_id','')); return jsonify(rc)


@app.route('/api/sales/returns/stats', methods=['GET'])
def returns_stats():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    today = datetime.utcnow().strftime('%Y-%m-%d')
    month = datetime.utcnow().strftime('%Y-%m')
    if USE_MEMORY_DB:
        rets = list(getattr(db, 'sales_returns', []))
        today_rets = [r for r in rets if r.get('created_at','').startswith(today)]
        month_rets = [r for r in rets if r.get('created_at','').startswith(month)]
        return jsonify({
            'total_returns':      len(rets),
            'today_returns':      len(today_rets),
            'today_refund':       sum(r.get('total_refund',0) for r in today_rets),
            'month_returns':      len(month_rets),
            'month_refund':       sum(r.get('total_refund',0) for r in month_rets),
            'total_refund_alltime': sum(r.get('total_refund',0) for r in rets),
        })
    today_rets = list(db.sales_returns.find({'created_at': {'$gte': today}}))
    month_rets = list(db.sales_returns.find({'created_at': {'$gte': month}}))
    all_rets   = list(db.sales_returns.find())
    return jsonify({
        'total_returns':      len(all_rets),
        'today_returns':      len(today_rets),
        'today_refund':       sum(r.get('total_refund',0) for r in today_rets),
        'month_returns':      len(month_rets),
        'month_refund':       sum(r.get('total_refund',0) for r in month_rets),
        'total_refund_alltime': sum(r.get('total_refund',0) for r in all_rets),
    })


@app.route('/pos/returns')
def pos_returns_page():
    if 'user_id' not in session: return redirect(url_for('login_page'))
    if session.get('role') not in ['admin', 'sales_manager', 'sales_person']:
        return redirect_to_dashboard(session.get('role'))
    return render_template('pos_returns.html')


# =============================================================================
# END-OF-DAY Z-REPORT
# =============================================================================

@app.route('/api/reports/z-report', methods=['GET'])
def get_z_report():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    report_date = request.args.get('date', datetime.utcnow().strftime('%Y-%m-%d'))
    date_start  = report_date + 'T00:00:00'
    date_end    = report_date + 'T23:59:59'

    if USE_MEMORY_DB:
        sales   = [s for s in getattr(db,'sales',[]) if date_start <= s.get('created_at','') <= date_end]
        returns = [r for r in getattr(db,'sales_returns',[]) if date_start <= r.get('created_at','') <= date_end]
        shifts  = [sh for sh in getattr(db,'shifts',[]) if sh.get('date','') == report_date or
                   sh.get('opened_at','').startswith(report_date)]
    else:
        sales   = list(db.sales.find({'created_at': {'$gte': date_start, '$lte': date_end}}))
        returns = list(db.sales_returns.find({'created_at': {'$gte': date_start, '$lte': date_end}}))
        shifts  = list(db.shifts.find({'$or': [
            {'date': report_date},
            {'opened_at': {'$gte': date_start, '$lte': date_end}}
        ]}))

    gross_sales   = sum(s.get('total', 0) for s in sales)
    discounts     = sum(s.get('discount_amount', 0) for s in sales)
    tax_collected = sum(s.get('tax_amount', 0) for s in sales)
    net_sales     = gross_sales - discounts
    total_refunds = sum(r.get('total_refund', 0) for r in returns)
    net_revenue   = net_sales - total_refunds

    cash_sales  = sum(s.get('total', 0) for s in sales if s.get('payment_method') == 'cash')
    card_sales  = sum(s.get('total', 0) for s in sales if s.get('payment_method') in ['card','credit_card','debit_card'])
    other_sales = gross_sales - cash_sales - card_sales

    opening_float = sum(sh.get('opening_float', 0) for sh in shifts)
    cash_refunds  = sum(r.get('total_refund', 0) for r in returns if r.get('refund_method') == 'cash')
    expected_cash = opening_float + cash_sales - cash_refunds

    promo_used    = len([s for s in sales if s.get('promo_code')])
    items_sold    = sum(sum(int(i.get('qty', 1)) for i in s.get('items', [])) for s in sales)

    by_category = {}
    for s in sales:
        for item in s.get('items', []):
            cat = item.get('category', 'Uncategorized')
            by_category[cat] = by_category.get(cat, 0) + float(item.get('subtotal', 0))

    top_products = {}
    for s in sales:
        for item in s.get('items', []):
            name = item.get('name', '')
            if name:
                if name not in top_products:
                    top_products[name] = {'qty': 0, 'revenue': 0}
                top_products[name]['qty']     += int(item.get('qty', 1))
                top_products[name]['revenue'] += float(item.get('subtotal', 0))
    top5 = sorted(top_products.items(), key=lambda x: x[1]['revenue'], reverse=True)[:5]

    return jsonify({
        'report_date':     report_date,
        'transactions':    len(sales),
        'items_sold':      items_sold,
        'gross_sales':     round(gross_sales, 2),
        'discounts':       round(discounts, 2),
        'tax_collected':   round(tax_collected, 2),
        'net_sales':       round(net_sales, 2),
        'total_refunds':   round(total_refunds, 2),
        'net_revenue':     round(net_revenue, 2),
        'cash_sales':      round(cash_sales, 2),
        'card_sales':      round(card_sales, 2),
        'other_sales':     round(other_sales, 2),
        'cash_refunds':    round(cash_refunds, 2),
        'opening_float':   round(opening_float, 2),
        'expected_cash':   round(expected_cash, 2),
        'promo_used':      promo_used,
        'shifts':          len(shifts),
        'by_category':     by_category,
        'top_products':    [{'name': k, **v} for k, v in top5],
    })


@app.route('/reports/z-report')
def z_report_page():
    if 'user_id' not in session: return redirect(url_for('login_page'))
    if session.get('role') not in ['admin', 'sales_manager']:
        return redirect_to_dashboard(session.get('role'))
    return render_template('z_report.html')


# =============================================================================
# STOCK ADJUSTMENTS & PHYSICAL COUNT
# =============================================================================

def _next_adj_num():
    prefix = f'ADJ-{datetime.utcnow().strftime("%Y%m")}-'
    if USE_MEMORY_DB:
        nums = [int(a.get('adj_id','ADJ-000000-000').split('-')[-1])
                for a in getattr(db,'stock_adjustments',[]) if a.get('adj_id','').startswith(prefix)]
        return prefix + f'{(max(nums, default=0)+1):03d}'
    last = db.stock_adjustments.find_one({'adj_id': {'$regex': f'^{prefix}'}}, sort=[('adj_id', -1)])
    n = int(last['adj_id'].split('-')[-1]) + 1 if last else 1
    return prefix + f'{n:03d}'


@app.route('/api/inventory/adjustments', methods=['GET'])
def get_adjustments():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    adj_type = request.args.get('type', '')
    dfrom    = request.args.get('from', '')
    dto      = request.args.get('to', '')
    if USE_MEMORY_DB:
        adjs = list(getattr(db, 'stock_adjustments', []))
        if adj_type: adjs = [a for a in adjs if a.get('adj_type') == adj_type]
        if dfrom:    adjs = [a for a in adjs if a.get('created_at','') >= dfrom]
        if dto:      adjs = [a for a in adjs if a.get('created_at','') <= dto + 'T23:59:59']
        adjs = sorted(adjs, key=lambda x: x.get('created_at',''), reverse=True)
        return jsonify({'adjustments': [{**dict(a), '_id': str(a.get('_id',''))} for a in adjs]})
    q = {}
    if adj_type: q['adj_type'] = adj_type
    if dfrom or dto:
        q['created_at'] = {}
        if dfrom: q['created_at']['$gte'] = dfrom
        if dto:   q['created_at']['$lte'] = dto + 'T23:59:59'
    adjs = list(db.stock_adjustments.find(q).sort('created_at', -1))
    for a in adjs: a['_id'] = str(a['_id'])
    return jsonify({'adjustments': adjs})


@app.route('/api/inventory/adjustments', methods=['POST'])
def create_adjustment():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    if session.get('role') not in ['admin', 'procurement_manager']:
        return jsonify({'error': 'Forbidden'}), 403
    data     = request.json or {}
    items    = data.get('items', [])
    adj_type = data.get('adj_type', 'manual')  # manual, damage, expiry, physical_count, transfer
    if not items: return jsonify({'error': 'No items provided'}), 400
    doc = {
        'adj_id':    _next_adj_num(),
        'adj_type':  adj_type,
        'reason':    data.get('reason', ''),
        'notes':     data.get('notes', ''),
        'items':     items,
        'status':    'completed',
        'created_by': session.get('full_name', 'System'),
        'created_at': datetime.utcnow().isoformat()
    }
    # Apply qty changes
    for item in items:
        pid   = item.get('product_id', '')
        delta = int(item.get('delta', 0))  # +ve = add, -ve = remove
        if pid and delta != 0:
            if USE_MEMORY_DB:
                prod = next((p for p in getattr(db,'products',[]) if str(p.get('_id')) == pid), None)
                if prod:
                    prod['quantity'] = max(0, int(prod.get('quantity', 0)) + delta)
            else:
                try: db.products.update_one({'_id': ObjectId(pid)}, {'$inc': {'quantity': delta}})
                except: pass
    if USE_MEMORY_DB:
        doc['_id'] = str(uuid.uuid4())
        if not hasattr(db, 'stock_adjustments'): db.stock_adjustments = []
        db.stock_adjustments.append(doc)
    else:
        result = db.stock_adjustments.insert_one(doc)
        doc['_id'] = str(result.inserted_id)
    log_audit('ADJUSTMENT_CREATED', 'Inventory', doc['_id'],
              f"{doc['adj_id']} — {adj_type} — {len(items)} items")
    return jsonify({'success': True, 'adjustment': doc}), 201


@app.route('/api/inventory/adjustments/stats', methods=['GET'])
def adjustment_stats():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    month = datetime.utcnow().strftime('%Y-%m')
    if USE_MEMORY_DB:
        adjs = list(getattr(db, 'stock_adjustments', []))
        month_adjs = [a for a in adjs if a.get('created_at','').startswith(month)]
        by_type = {}
        for a in adjs:
            t = a.get('adj_type','manual')
            by_type[t] = by_type.get(t, 0) + 1
        return jsonify({'total': len(adjs), 'this_month': len(month_adjs), 'by_type': by_type})
    total = db.stock_adjustments.count_documents({})
    mo    = db.stock_adjustments.count_documents({'created_at': {'$gte': month}})
    pipe  = [{'$group': {'_id': '$adj_type', 'count': {'$sum': 1}}}]
    by_type = {r['_id']: r['count'] for r in db.stock_adjustments.aggregate(pipe)}
    return jsonify({'total': total, 'this_month': mo, 'by_type': by_type})


@app.route('/inventory/adjustments')
def stock_adjustments_page():
    if 'user_id' not in session: return redirect(url_for('login_page'))
    if session.get('role') not in ['admin', 'procurement_manager']:
        return redirect_to_dashboard(session.get('role'))
    return render_template('stock_adjustments.html')


# =============================================================================
# EXPIRY DATE TRACKING
# =============================================================================

@app.route('/api/inventory/expiring', methods=['GET'])
def get_expiring_products():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    days    = int(request.args.get('days', 30))
    today   = datetime.utcnow().date()
    cutoff  = (today + timedelta(days=days)).isoformat()
    today_s = today.isoformat()
    if USE_MEMORY_DB:
        prods = list(getattr(db, 'products', []))
        expiring = []
        for p in prods:
            exp = p.get('expiry_date', '')
            if exp and exp <= cutoff:
                pc = dict(p); pc['_id'] = str(pc.get('_id',''))
                pc['days_to_expiry'] = (datetime.fromisoformat(exp).date() - today).days
                pc['is_expired'] = exp < today_s
                expiring.append(pc)
        expiring.sort(key=lambda x: x.get('expiry_date',''))
        return jsonify({'products': expiring, 'count': len(expiring)})
    q = {'expiry_date': {'$exists': True, '$ne': '', '$lte': cutoff}}
    prods = list(db.products.find(q).sort('expiry_date', 1))
    expiring = []
    for p in prods:
        p['_id'] = str(p['_id'])
        exp = p.get('expiry_date','')
        try: p['days_to_expiry'] = (datetime.fromisoformat(exp).date() - today).days
        except: p['days_to_expiry'] = 0
        p['is_expired'] = exp < today_s
        expiring.append(p)
    return jsonify({'products': expiring, 'count': len(expiring)})


# =============================================================================
# CREDIT SALES / ACCOUNTS RECEIVABLE
# =============================================================================

def _next_credit_num():
    prefix = f'CR-{datetime.utcnow().strftime("%Y%m")}-'
    if USE_MEMORY_DB:
        nums = [int(c.get('credit_id','CR-000000-000').split('-')[-1])
                for c in getattr(db,'credit_sales',[]) if c.get('credit_id','').startswith(prefix)]
        return prefix + f'{(max(nums, default=0)+1):03d}'
    last = db.credit_sales.find_one({'credit_id': {'$regex': f'^{prefix}'}}, sort=[('credit_id', -1)])
    n = int(last['credit_id'].split('-')[-1]) + 1 if last else 1
    return prefix + f'{n:03d}'


@app.route('/api/credit-sales', methods=['GET'])
def get_credit_sales():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    status = request.args.get('status', '')
    search = request.args.get('search', '')
    if USE_MEMORY_DB:
        cs = list(getattr(db, 'credit_sales', []))
        if status: cs = [c for c in cs if c.get('status') == status]
        if search: cs = [c for c in cs if search.lower() in (c.get('customer_name','') + c.get('customer_phone','')).lower()]
        cs = sorted(cs, key=lambda x: x.get('created_at',''), reverse=True)
        return jsonify({'credit_sales': [{**dict(c), '_id': str(c.get('_id',''))} for c in cs]})
    q = {}
    if status: q['status'] = status
    if search: q['$or'] = [{'customer_name': {'$regex': search, '$options':'i'}},
                            {'customer_phone': {'$regex': search}}]
    cs = list(db.credit_sales.find(q).sort('created_at', -1))
    for c in cs: c['_id'] = str(c['_id'])
    return jsonify({'credit_sales': cs})


@app.route('/api/credit-sales', methods=['POST'])
def create_credit_sale():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    data   = request.json or {}
    amount = round(float(data.get('amount', 0)), 2)
    if not amount: return jsonify({'error': 'Amount required'}), 400
    if not data.get('customer_name','').strip(): return jsonify({'error': 'Customer name required'}), 400
    due_date = data.get('due_date', (datetime.utcnow() + timedelta(days=30)).strftime('%Y-%m-%d'))
    doc = {
        'credit_id':       _next_credit_num(),
        'customer_name':   data.get('customer_name','').strip(),
        'customer_phone':  data.get('customer_phone','').strip(),
        'customer_id':     data.get('customer_id',''),
        'original_sale_id': data.get('original_sale_id',''),
        'items':           data.get('items', []),
        'amount':          amount,
        'amount_paid':     0.0,
        'balance':         amount,
        'due_date':        due_date,
        'notes':           data.get('notes',''),
        'status':          'outstanding',
        'payments':        [],
        'created_by':      session.get('full_name','Cashier'),
        'created_at':      datetime.utcnow().isoformat()
    }
    if USE_MEMORY_DB:
        doc['_id'] = str(uuid.uuid4())
        if not hasattr(db, 'credit_sales'): db.credit_sales = []
        db.credit_sales.append(doc)
    else:
        result = db.credit_sales.insert_one(doc)
        doc['_id'] = str(result.inserted_id)
    log_audit('CREDIT_CREATED', 'Credit Sales', doc['_id'],
              f"{doc['credit_id']} — {doc['customer_name']} — PKR {amount}")
    return jsonify({'success': True, 'credit_sale': doc}), 201


@app.route('/api/credit-sales/<cid>/pay', methods=['POST'])
def pay_credit_sale(cid):
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    data   = request.json or {}
    amount = round(float(data.get('amount', 0)), 2)
    if not amount: return jsonify({'error': 'Amount required'}), 400
    pmt = {
        'amount': amount,
        'method': data.get('method', 'cash'),
        'reference': data.get('reference',''),
        'paid_at': datetime.utcnow().isoformat(),
        'received_by': session.get('full_name','Cashier')
    }
    if USE_MEMORY_DB:
        cr = next((c for c in getattr(db,'credit_sales',[]) if str(c.get('_id'))==cid or c.get('credit_id')==cid), None)
        if not cr: return jsonify({'error': 'Not found'}), 404
        if 'payments' not in cr: cr['payments'] = []
        cr['payments'].append(pmt)
        cr['amount_paid'] = round(cr.get('amount_paid',0) + amount, 2)
        cr['balance']     = round(cr.get('amount',0) - cr['amount_paid'], 2)
        cr['status']      = 'settled' if cr['balance'] <= 0 else 'partial'
        log_audit('CREDIT_PAYMENT', 'Credit Sales', cid, f"PKR {amount} received")
        return jsonify({'success': True, 'balance': cr['balance'], 'status': cr['status']})
    try:
        cr = db.credit_sales.find_one({'_id': ObjectId(cid)})
    except:
        cr = db.credit_sales.find_one({'credit_id': cid})
    if not cr: return jsonify({'error': 'Not found'}), 404
    new_paid    = round(float(cr.get('amount_paid',0)) + amount, 2)
    new_balance = round(float(cr.get('amount',0)) - new_paid, 2)
    new_status  = 'settled' if new_balance <= 0 else 'partial'
    try:
        db.credit_sales.update_one({'_id': ObjectId(cid)}, {
            '$push': {'payments': pmt},
            '$set':  {'amount_paid': new_paid, 'balance': new_balance, 'status': new_status}
        })
    except: pass
    log_audit('CREDIT_PAYMENT', 'Credit Sales', cid, f"PKR {amount} received")
    return jsonify({'success': True, 'balance': new_balance, 'status': new_status})


@app.route('/api/credit-sales/stats', methods=['GET'])
def credit_sales_stats():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        cs = list(getattr(db, 'credit_sales', []))
        outstanding = [c for c in cs if c.get('status') == 'outstanding']
        partial     = [c for c in cs if c.get('status') == 'partial']
        overdue     = [c for c in cs if c.get('status') not in ['settled'] and
                       c.get('due_date','') < datetime.utcnow().strftime('%Y-%m-%d')]
        return jsonify({
            'total_accounts':    len(cs),
            'outstanding_count': len(outstanding) + len(partial),
            'total_outstanding': round(sum(c.get('balance',0) for c in cs if c.get('status')!='settled'), 2),
            'overdue_count':     len(overdue),
            'overdue_amount':    round(sum(c.get('balance',0) for c in overdue), 2),
            'settled_count':     len([c for c in cs if c.get('status')=='settled']),
        })
    total   = db.credit_sales.count_documents({})
    pipe    = [{'$match': {'status': {'$ne': 'settled'}}},
               {'$group': {'_id': None, 'count': {'$sum':1}, 'bal': {'$sum':'$balance'}}}]
    res     = list(db.credit_sales.aggregate(pipe))
    today   = datetime.utcnow().strftime('%Y-%m-%d')
    overdue = list(db.credit_sales.find({'status': {'$ne':'settled'}, 'due_date': {'$lt': today}}))
    return jsonify({
        'total_accounts':    total,
        'outstanding_count': res[0]['count'] if res else 0,
        'total_outstanding': res[0]['bal']   if res else 0,
        'overdue_count':     len(overdue),
        'overdue_amount':    round(sum(c.get('balance',0) for c in overdue), 2),
        'settled_count':     db.credit_sales.count_documents({'status':'settled'}),
    })


@app.route('/credit-sales')
def credit_sales_page():
    if 'user_id' not in session: return redirect(url_for('login_page'))
    if session.get('role') not in ['admin', 'sales_manager', 'finance_manager']:
        return redirect_to_dashboard(session.get('role'))
    return render_template('credit_sales.html')


# =============================================================================
# NOTIFICATIONS & ALERTS
# =============================================================================

@app.route('/api/notifications', methods=['GET'])
def get_notifications():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    role  = session.get('role', '')
    notes = []

    # Low stock alerts
    try:
        if USE_MEMORY_DB:
            prods = [p for p in getattr(db,'products',[]) if int(p.get('quantity',0)) <= int(p.get('reorder_level',10))]
        else:
            prods = list(db.products.find({'$expr': {'$lte': ['$quantity', '$reorder_level']}}))
        for p in prods[:10]:
            notes.append({'type':'low_stock','level':'warning',
                          'title':f"Low Stock: {p.get('name','')}",
                          'message':f"Only {p.get('quantity',0)} units left (reorder at {p.get('reorder_level',10)})",
                          'link':'/procurement/stock','created_at':datetime.utcnow().isoformat()})
    except: pass

    # Expiring products (within 7 days)
    try:
        cutoff = (datetime.utcnow().date() + timedelta(days=7)).isoformat()
        today  = datetime.utcnow().strftime('%Y-%m-%d')
        if USE_MEMORY_DB:
            exp_prods = [p for p in getattr(db,'products',[])
                         if p.get('expiry_date','') and p.get('expiry_date','') <= cutoff]
        else:
            exp_prods = list(db.products.find({'expiry_date': {'$exists':True,'$ne':'','$lte':cutoff}}))
        for p in exp_prods[:5]:
            exp = p.get('expiry_date','')
            days = (datetime.fromisoformat(exp).date() - datetime.utcnow().date()).days if exp else 0
            level = 'danger' if days <= 0 else 'warning'
            notes.append({'type':'expiry','level':level,
                          'title':f"{'EXPIRED' if days<=0 else 'Expiring Soon'}: {p.get('name','')}",
                          'message':f"{'Expired' if days<=0 else f'Expires in {days} day(s)'} — {exp}",
                          'link':'/inventory/adjustments','created_at':datetime.utcnow().isoformat()})
    except: pass

    # Overdue vendor invoices
    if role in ['admin', 'procurement_manager', 'finance_manager']:
        try:
            today = datetime.utcnow().strftime('%Y-%m-%d')
            if USE_MEMORY_DB:
                inv = [i for i in getattr(db,'vendor_invoices',[])
                       if i.get('status') not in ['paid'] and i.get('due_date','') and i.get('due_date','') < today]
            else:
                inv = list(db.vendor_invoices.find({'status':{'$ne':'paid'},
                                                     'due_date':{'$lt':today,'$exists':True,'$ne':''}}))
            if inv:
                notes.append({'type':'overdue_invoice','level':'danger',
                              'title':f"{len(inv)} Overdue Vendor Invoice(s)",
                              'message':f"PKR {sum(float(i.get('total',0)) for i in inv):,.2f} overdue",
                              'link':'/procurement/invoices','created_at':datetime.utcnow().isoformat()})
        except: pass

    # Pending payroll
    if role in ['admin', 'finance_manager']:
        try:
            if USE_MEMORY_DB:
                pending_pr = [p for p in getattr(db,'payroll',[]) if p.get('status') == 'pending']
            else:
                pending_pr = list(db.payroll.find({'status':'pending'}))
            if pending_pr:
                notes.append({'type':'payroll','level':'info',
                              'title':f"{len(pending_pr)} Payroll Record(s) Pending",
                              'message':f"PKR {sum(float(p.get('net_salary',0)) for p in pending_pr):,.2f} to be processed",
                              'link':'/hr','created_at':datetime.utcnow().isoformat()})
        except: pass

    # Overdue credit sales
    if role in ['admin', 'sales_manager', 'finance_manager']:
        try:
            today = datetime.utcnow().strftime('%Y-%m-%d')
            if USE_MEMORY_DB:
                over_cr = [c for c in getattr(db,'credit_sales',[])
                           if c.get('status') not in ['settled'] and c.get('due_date','') < today]
            else:
                over_cr = list(db.credit_sales.find({'status':{'$ne':'settled'},
                                                      'due_date':{'$lt':today,'$exists':True,'$ne':''}}))
            if over_cr:
                notes.append({'type':'credit','level':'warning',
                              'title':f"{len(over_cr)} Overdue Credit Account(s)",
                              'message':f"PKR {sum(float(c.get('balance',0)) for c in over_cr):,.2f} outstanding",
                              'link':'/credit-sales','created_at':datetime.utcnow().isoformat()})
        except: pass

    return jsonify({'notifications': notes, 'count': len(notes),
                    'unread': len([n for n in notes if n['level'] in ['danger','warning']])})


# =============================================================================
# TAX / GST REPORT
# =============================================================================

@app.route('/api/reports/tax', methods=['GET'])
def get_tax_report():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    if session.get('role') not in ['admin', 'finance_manager']:
        return jsonify({'error': 'Forbidden'}), 403
    period_type = request.args.get('period', 'month')  # month or quarter
    year  = int(request.args.get('year',  datetime.utcnow().year))
    month = int(request.args.get('month', datetime.utcnow().month))

    if period_type == 'quarter':
        q_start_month = ((month - 1) // 3) * 3 + 1
        date_start = f'{year}-{q_start_month:02d}-01'
        end_month  = q_start_month + 2
        end_year   = year + (1 if end_month > 12 else 0)
        end_month  = end_month if end_month <= 12 else end_month - 12
        import calendar
        date_end = f'{end_year}-{end_month:02d}-{calendar.monthrange(end_year, end_month)[1]}'
    else:
        import calendar
        date_start = f'{year}-{month:02d}-01'
        date_end   = f'{year}-{month:02d}-{calendar.monthrange(year, month)[1]}'

    date_start_iso = date_start + 'T00:00:00'
    date_end_iso   = date_end   + 'T23:59:59'

    if USE_MEMORY_DB:
        sales = [s for s in getattr(db,'sales',[])
                 if date_start_iso <= s.get('created_at','') <= date_end_iso]
        invoices = [i for i in getattr(db,'vendor_invoices',[])
                    if date_start_iso <= i.get('created_at','') <= date_end_iso]
    else:
        sales    = list(db.sales.find({'created_at': {'$gte': date_start_iso, '$lte': date_end_iso}}))
        invoices = list(db.vendor_invoices.find({'created_at': {'$gte': date_start_iso, '$lte': date_end_iso}}))

    output_tax  = round(sum(float(s.get('tax_amount', 0)) for s in sales), 2)
    gross_sales = round(sum(float(s.get('total', 0)) for s in sales), 2)
    input_tax   = round(sum(float(i.get('tax_amount', 0)) for i in invoices), 2)
    net_tax     = round(output_tax - input_tax, 2)

    # Monthly breakdown within period
    monthly = {}
    for s in sales:
        mo = s.get('created_at','')[:7]
        if mo not in monthly: monthly[mo] = {'output_tax':0,'gross_sales':0,'transactions':0}
        monthly[mo]['output_tax']   += float(s.get('tax_amount',0))
        monthly[mo]['gross_sales']  += float(s.get('total',0))
        monthly[mo]['transactions'] += 1

    return jsonify({
        'period': {'type': period_type, 'from': date_start, 'to': date_end},
        'output_tax':   output_tax,
        'input_tax':    input_tax,
        'net_tax':      net_tax,
        'gross_sales':  gross_sales,
        'transactions': len(sales),
        'monthly':      [{'month': k, **v} for k, v in sorted(monthly.items())],
    })


# =============================================================================
# FBR / POS INTEGRATION (UI SHELL — pending FBR registration)
# =============================================================================

@app.route('/fbr')
def fbr_page():
    if 'user_id' not in session: return redirect(url_for('login_page'))
    if session.get('role') != 'admin': return redirect_to_dashboard(session.get('role'))
    return render_template('fbr_dashboard.html')


@app.route('/api/fbr/status', methods=['GET'])
def fbr_status():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    # Returns placeholder status until FBR registration is complete
    return jsonify({
        'registered':     False,
        'registration_no': '',
        'pos_id':         '',
        'status':         'pending_registration',
        'message':        'FBR registration pending. Module will activate after IRIS registration.',
        'invoices_sent':  0,
        'last_sync':      None,
    })


# =============================================================================
# PRODUCT BUNDLES / COMBO DEALS
# =============================================================================

def _next_bundle_num():
    prefix = 'BND-'
    if USE_MEMORY_DB:
        nums = [int(b.get('bundle_id','BND-0000').split('-')[-1])
                for b in getattr(db,'bundles',[]) if b.get('bundle_id','').startswith(prefix)]
        return prefix + f'{(max(nums, default=0)+1):04d}'
    last = db.bundles.find_one({'bundle_id': {'$regex': f'^{prefix}'}}, sort=[('bundle_id', -1)])
    n = int(last['bundle_id'].split('-')[-1]) + 1 if last else 1
    return prefix + f'{n:04d}'


@app.route('/api/bundles', methods=['GET'])
def get_bundles():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        bundles = list(getattr(db, 'bundles', []))
        return jsonify({'bundles': [{**dict(b), '_id': str(b.get('_id',''))} for b in bundles]})
    bundles = list(db.bundles.find().sort('created_at', -1))
    for b in bundles: b['_id'] = str(b['_id'])
    return jsonify({'bundles': bundles})


@app.route('/api/bundles', methods=['POST'])
def create_bundle():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    if session.get('role') not in ['admin', 'sales_manager']:
        return jsonify({'error': 'Forbidden'}), 403
    data  = request.json or {}
    items = data.get('items', [])
    if not data.get('name','').strip(): return jsonify({'error': 'Bundle name required'}), 400
    if len(items) < 2: return jsonify({'error': 'Bundle must have at least 2 items'}), 400
    regular_price = round(sum(float(i.get('unit_price',0)) * int(i.get('qty',1)) for i in items), 2)
    bundle_price  = round(float(data.get('bundle_price', regular_price)), 2)
    doc = {
        'bundle_id':     _next_bundle_num(),
        'name':          data.get('name','').strip(),
        'description':   data.get('description',''),
        'items':         items,
        'regular_price': regular_price,
        'bundle_price':  bundle_price,
        'discount_pct':  round((1 - bundle_price/regular_price)*100, 1) if regular_price else 0,
        'is_active':     data.get('is_active', True),
        'valid_from':    data.get('valid_from',''),
        'valid_to':      data.get('valid_to',''),
        'created_by':    session.get('full_name','Admin'),
        'created_at':    datetime.utcnow().isoformat()
    }
    if USE_MEMORY_DB:
        doc['_id'] = str(uuid.uuid4())
        if not hasattr(db, 'bundles'): db.bundles = []
        db.bundles.append(doc)
        return jsonify({'success': True, 'bundle': doc}), 201
    result = db.bundles.insert_one(doc)
    doc['_id'] = str(result.inserted_id)
    return jsonify({'success': True, 'bundle': doc}), 201


@app.route('/api/bundles/<bid>', methods=['PUT'])
def update_bundle(bid):
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    data  = request.json or {}
    items = data.get('items', [])
    regular_price = round(sum(float(i.get('unit_price',0)) * int(i.get('qty',1)) for i in items), 2) if items else 0
    bundle_price  = round(float(data.get('bundle_price', regular_price)), 2)
    upd = {
        'name': data.get('name',''), 'description': data.get('description',''),
        'items': items, 'regular_price': regular_price, 'bundle_price': bundle_price,
        'discount_pct': round((1 - bundle_price/regular_price)*100, 1) if regular_price else 0,
        'is_active': data.get('is_active', True),
        'valid_from': data.get('valid_from',''), 'valid_to': data.get('valid_to',''),
    }
    if USE_MEMORY_DB:
        b = next((x for x in getattr(db,'bundles',[]) if str(x.get('_id'))==bid or x.get('bundle_id')==bid), None)
        if not b: return jsonify({'error':'Not found'}), 404
        b.update(upd); return jsonify({'success':True,'bundle':b})
    try: db.bundles.update_one({'_id': ObjectId(bid)}, {'$set': upd})
    except: db.bundles.update_one({'bundle_id': bid}, {'$set': upd})
    return jsonify({'success': True})


@app.route('/api/bundles/<bid>', methods=['DELETE'])
def delete_bundle(bid):
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    if USE_MEMORY_DB:
        bundles = getattr(db,'bundles',[])
        db.bundles = [b for b in bundles if str(b.get('_id'))!=bid and b.get('bundle_id')!=bid]
        return jsonify({'success': True})
    try: db.bundles.delete_one({'_id': ObjectId(bid)})
    except: db.bundles.delete_one({'bundle_id': bid})
    return jsonify({'success': True})


@app.route('/bundles')
def bundles_page():
    if 'user_id' not in session: return redirect(url_for('login_page'))
    if session.get('role') not in ['admin', 'sales_manager']:
        return redirect_to_dashboard(session.get('role'))
    return render_template('bundles.html')


# =============================================================================
# SUPPLIER PRICE HISTORY
# =============================================================================

@app.route('/api/vendors/<vid>/price-history', methods=['GET'])
def vendor_price_history(vid):
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    product = request.args.get('product', '')
    if USE_MEMORY_DB:
        pos = list(getattr(db,'purchase_orders',[]))
        history = []
        for po in pos:
            if str(po.get('vendor_id','')) != vid and po.get('vendor_id','') != vid: continue
            for item in po.get('items',[]):
                if product and product.lower() not in item.get('product_name','').lower(): continue
                history.append({
                    'po_number':    po.get('po_number',''),
                    'date':         po.get('created_at','')[:10],
                    'product_name': item.get('product_name',''),
                    'product_id':   item.get('product_id',''),
                    'unit_price':   item.get('unit_price', item.get('unit_cost',0)),
                    'qty':          item.get('qty', item.get('quantity',0)),
                    'status':       po.get('status',''),
                })
        history.sort(key=lambda x: x['date'], reverse=True)
        return jsonify({'history': history})
    q = {'vendor_id': vid}
    pos = list(db.purchase_orders.find(q).sort('created_at',-1).limit(50))
    history = []
    for po in pos:
        for item in po.get('items',[]):
            if product and product.lower() not in item.get('product_name','').lower(): continue
            history.append({
                'po_number':    po.get('po_number',''),
                'date':         po.get('created_at','')[:10],
                'product_name': item.get('product_name',''),
                'product_id':   str(item.get('product_id','')),
                'unit_price':   item.get('unit_price', item.get('unit_cost',0)),
                'qty':          item.get('qty', item.get('quantity',0)),
                'status':       po.get('status',''),
            })
    return jsonify({'history': history})


# =============================================================================
# DATA EXPORT (CSV)
# =============================================================================

@app.route('/api/export/<module>', methods=['GET'])
def export_csv(module):
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    allowed = {
        'sales':       ('sales',       ['created_at','receipt_number','customer_name','total','payment_method','cashier_name']),
        'products':    ('products',    ['name','sku','category','quantity','sale_price','cost_price','reorder_level']),
        'employees':   ('employees',   ['emp_id','name','department','designation','basic_salary','gross_salary','net_salary','status']),
        'payments':    ('payments',    ['payment_id','type','payee','amount','method','date','status']),
        'returns':     ('sales_returns',['return_id','customer_name','total_refund','refund_method','reason','created_at']),
        'credit_sales':('credit_sales',['credit_id','customer_name','customer_phone','amount','amount_paid','balance','due_date','status']),
        'adjustments': ('stock_adjustments',['adj_id','adj_type','reason','created_by','created_at']),
        'vendors':     ('vendors',     ['name','contact_person','email','phone','city','status']),
        'payroll':     ('payroll',     ['payroll_number','employee_name','month','basic_salary','gross_salary','net_salary','status']),
    }
    if module not in allowed: return jsonify({'error': 'Unknown module'}), 400
    collection_name, fields = allowed[module]
    if USE_MEMORY_DB:
        rows = list(getattr(db, collection_name, []))
    else:
        rows = list(getattr(db, collection_name).find({}, {f: 1 for f in fields}))
    output = BytesIO()
    writer_text = []
    writer_text.append(','.join(fields))
    for row in rows:
        line = []
        for f in fields:
            val = str(row.get(f,'')).replace(',','').replace('\n',' ')
            line.append(val)
        writer_text.append(','.join(line))
    csv_content = '\n'.join(writer_text)
    output = BytesIO(csv_content.encode('utf-8'))
    output.seek(0)
    filename = f'{module}_{datetime.utcnow().strftime("%Y%m%d_%H%M%S")}.csv'
    return send_file(output, mimetype='text/csv',
                     as_attachment=True, download_name=filename)


# =============================================================================
# PREDICTIVE ANALYSIS — ADMIN KPI DASHBOARD
# =============================================================================

def _month_add(year, month, delta):
    month += delta
    while month > 12: month -= 12; year += 1
    while month < 1:  month += 12; year -= 1
    return year, month

def _prev_months(n=12):
    now = datetime.utcnow()
    out = []
    for i in range(n - 1, -1, -1):
        y, m = _month_add(now.year, now.month, -i)
        out.append((f'{y}-{m:02d}', datetime(y, m, 1).strftime('%b %Y')))
    return out

def _next_months(n=6):
    now = datetime.utcnow()
    out = []
    for i in range(1, n + 1):
        y, m = _month_add(now.year, now.month, i)
        out.append((f'{y}-{m:02d}', datetime(y, m, 1).strftime('%b %Y')))
    return out

def _lr(values):
    n = len(values)
    if n < 2: return 0, (values[0] if values else 0)
    x_mean = (n - 1) / 2
    y_mean = sum(values) / n
    denom = sum((i - x_mean)**2 for i in range(n))
    if denom == 0: return 0, y_mean
    slope = sum((i - x_mean) * (v - y_mean) for i, v in enumerate(values)) / denom
    return slope, y_mean - slope * x_mean

def _forecast(values, periods=6):
    slope, intercept = _lr(values)
    n = len(values)
    return [max(0, round(intercept + slope * (n + i), 2)) for i in range(periods)]

def _growth_pct(values):
    non_zero = [v for v in values if v > 0]
    if len(non_zero) < 2: return 0.0
    return round((non_zero[-1] - non_zero[-2]) / non_zero[-2] * 100, 1)

def _trend_label(slope):
    if slope > 0.05: return 'growing'
    if slope < -0.05: return 'declining'
    return 'stable'


@app.route('/api/predictive/sales')
def api_predictive_sales():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    if session.get('role') != 'admin': return jsonify({'error': 'Forbidden'}), 403
    hist  = _prev_months(12)
    fcast = _next_months(6)
    rev   = {k: 0.0 for k, _ in hist}
    tx    = {k: 0   for k, _ in hist}
    disc  = {k: 0.0 for k, _ in hist}
    product_rev = {}
    pm_counts   = {}
    cutoff = hist[0][0] + '-01'
    if USE_MEMORY_DB:
        sales_list = [s for s in getattr(db,'sales',[]) if s.get('created_at','') >= cutoff]
    else:
        sales_list = list(db.sales.find({'created_at': {'$gte': cutoff}}))
    for s in sales_list:
        k = s.get('created_at','')[:7]
        if k in rev:
            rev[k]  += float(s.get('total', 0))
            tx[k]   += 1
            disc[k] += float(s.get('discount_amount', 0))
        for it in s.get('items', []):
            pname = it.get('name', 'Unknown')
            product_rev[pname] = product_rev.get(pname, 0) + float(it.get('subtotal', 0))
        pm = s.get('payment_method', 'other')
        pm_counts[pm] = pm_counts.get(pm, 0) + 1
    rev_v  = [round(rev[k], 2) for k, _ in hist]
    tx_v   = [tx[k] for k, _ in hist]
    avg_v  = [round(rev[k]/tx[k], 2) if tx[k] else 0 for k, _ in hist]
    rev_fc = _forecast(rev_v, 6)
    tx_fc  = _forecast(tx_v, 6)
    slope, _ = _lr(rev_v)
    growth   = _growth_pct(rev_v)
    total_rev = sum(rev_v)
    top_products = sorted(product_rev.items(), key=lambda x: x[1], reverse=True)[:10]
    six_proj = round(sum(rev_fc), 2)
    total_disc = sum(disc[k] for k, _ in hist)
    total_pm   = sum(pm_counts.values()) or 1
    recs = []
    if growth < -5:
        recs.append({'level':'danger','icon':'fa-arrow-trend-down','title':'Revenue Declining',
                     'message':f'Revenue dropped {abs(growth):.1f}% vs last month. Review pricing, run promotions, and check peak hour coverage.'})
    elif growth >= 10:
        recs.append({'level':'success','icon':'fa-arrow-trend-up','title':'Strong Revenue Growth',
                     'message':f'Revenue grew {growth:.1f}% vs last month. Scale inventory of top products and replicate what\'s working.'})
    else:
        recs.append({'level':'info','icon':'fa-chart-line','title':'Revenue Stable',
                     'message':f'Revenue changed {growth:+.1f}%. Introduce new product categories or bundles to accelerate growth.'})
    if top_products and total_rev:
        top_share = top_products[0][1] / total_rev * 100
        if top_share > 40:
            recs.append({'level':'warning','icon':'fa-exclamation-triangle','title':'Revenue Concentration Risk',
                         'message':f'"{top_products[0][0]}" drives {top_share:.0f}% of revenue. Diversify your product mix to reduce dependency.'})
    if total_rev and total_disc / total_rev > 0.15:
        recs.append({'level':'warning','icon':'fa-tags','title':'High Discount Rate',
                     'message':f'Discounts are {total_disc/total_rev*100:.1f}% of gross revenue. Tighten discount rules to protect margins.'})
    if pm_counts.get('cash', 0) / total_pm > 0.75:
        recs.append({'level':'info','icon':'fa-credit-card','title':'Low Digital Payment Adoption',
                     'message':'75%+ transactions are cash. Enable card/digital payments to improve checkout speed and reduce cash risk.'})
    recs.append({'level':'info','icon':'fa-robot','title':'6-Month AI Projection',
                 'message':f'Linear trend projects PKR {six_proj:,.0f} revenue over next 6 months. {"Invest in growth." if slope > 0 else "Take corrective action now."}'})
    return jsonify({
        'hist_labels': [l for _, l in hist], 'fcast_labels': [l for _, l in fcast],
        'revenue': rev_v, 'rev_forecast': rev_fc,
        'transactions': tx_v, 'tx_forecast': tx_fc, 'avg_transaction': avg_v,
        'top_products': [{'name': k, 'revenue': round(v, 2)} for k, v in top_products],
        'payment_methods': pm_counts,
        'kpis': {'total_revenue': round(total_rev, 2), 'avg_monthly': round(total_rev/12, 2),
                 'growth_rate': growth, 'transactions_total': sum(tx_v),
                 'six_month_proj': six_proj, 'trend': _trend_label(slope)},
        'recommendations': recs,
    })


@app.route('/api/predictive/procurement')
def api_predictive_procurement():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    if session.get('role') != 'admin': return jsonify({'error': 'Forbidden'}), 403
    hist  = _prev_months(12)
    fcast = _next_months(6)
    po_val = {k: 0.0 for k, _ in hist}
    po_cnt = {k: 0   for k, _ in hist}
    gr_cnt = {k: 0   for k, _ in hist}
    pr_cnt = {k: 0   for k, _ in hist}
    vendor_spend = {}
    cutoff = hist[0][0] + '-01'
    if USE_MEMORY_DB:
        pos  = [p for p in getattr(db,'purchase_orders_v2',[]) if p.get('created_at','') >= cutoff]
        grs  = [g for g in getattr(db,'goods_receipts',[]) if g.get('created_at','') >= cutoff and g.get('status') == 'confirmed']
        prs  = [p for p in getattr(db,'purchase_requisitions',[]) if p.get('created_at','') >= cutoff]
    else:
        pos  = list(db.purchase_orders_v2.find({'created_at': {'$gte': cutoff}}))
        grs  = list(db.goods_receipts.find({'created_at': {'$gte': cutoff}, 'status': 'confirmed'}))
        prs  = list(db.purchase_requisitions.find({'created_at': {'$gte': cutoff}}))
    for po in pos:
        k = po.get('created_at','')[:7]
        if k in po_val:
            po_val[k] += float(po.get('total', 0))
            po_cnt[k] += 1
        vn = po.get('vendor_name', 'Unknown')
        vendor_spend[vn] = vendor_spend.get(vn, 0) + float(po.get('total', 0))
    for gr in grs:
        k = gr.get('created_at','')[:7]
        if k in gr_cnt: gr_cnt[k] += 1
    for pr in prs:
        k = pr.get('created_at','')[:7]
        if k in pr_cnt: pr_cnt[k] += 1
    po_v = [round(po_val[k], 2) for k, _ in hist]
    po_c = [po_cnt[k] for k, _ in hist]
    gr_v = [gr_cnt[k] for k, _ in hist]
    pr_v = [pr_cnt[k] for k, _ in hist]
    po_fc = _forecast(po_v, 6)
    slope, _ = _lr(po_v)
    growth = _growth_pct(po_v)
    total_spend = sum(po_v)
    top_vendors = sorted(vendor_spend.items(), key=lambda x: x[1], reverse=True)[:8]
    total_pr = sum(pr_v); total_po = sum(po_c)
    conversion = round(total_po / total_pr * 100, 1) if total_pr else 0
    six_proj = round(sum(po_fc), 2)
    recs = []
    if growth > 15:
        recs.append({'level':'warning','icon':'fa-arrow-up','title':'Procurement Spend Increasing Fast',
                     'message':f'PO value grew {growth:.1f}% last month. Check if driven by price inflation or volume — negotiate better vendor rates.'})
    elif growth < -10:
        recs.append({'level':'info','icon':'fa-arrow-down','title':'Procurement Spend Decreasing',
                     'message':f'PO value dropped {abs(growth):.1f}%. Verify stock levels to avoid stockouts from reduced ordering.'})
    else:
        recs.append({'level':'success','icon':'fa-check-circle','title':'Procurement Spend Stable',
                     'message':f'Procurement spend changed {growth:+.1f}%. Good control. Monitor vendor price trends proactively.'})
    if conversion < 60:
        recs.append({'level':'warning','icon':'fa-exchange-alt','title':'Low PR-to-PO Conversion',
                     'message':f'Only {conversion:.0f}% of PRs convert to POs. Review approval bottlenecks and pending requisitions.'})
    if top_vendors and total_spend:
        top_share = top_vendors[0][1] / total_spend * 100
        if top_share > 50:
            recs.append({'level':'warning','icon':'fa-building','title':'Vendor Concentration Risk',
                         'message':f'"{top_vendors[0][0]}" accounts for {top_share:.0f}% of spend. Diversify suppliers to reduce risk.'})
    recs.append({'level':'info','icon':'fa-robot','title':'6-Month Spend Projection',
                 'message':f'Projected procurement spend: PKR {six_proj:,.0f}. {"Plan budget for growth." if slope > 0 else "Declining trend — review stock adequacy."}'})
    return jsonify({
        'hist_labels': [l for _, l in hist], 'fcast_labels': [l for _, l in fcast],
        'po_values': po_v, 'po_forecast': po_fc, 'po_counts': po_c, 'gr_counts': gr_v, 'pr_counts': pr_v,
        'top_vendors': [{'name': k, 'spend': round(v, 2)} for k, v in top_vendors],
        'kpis': {'total_spend': round(total_spend, 2), 'avg_monthly': round(total_spend/12, 2),
                 'growth_rate': growth, 'total_pos': total_po, 'total_grs': sum(gr_v),
                 'conversion_rate': conversion, 'six_month_proj': six_proj, 'trend': _trend_label(slope)},
        'recommendations': recs,
    })


@app.route('/api/predictive/finance')
def api_predictive_finance():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    if session.get('role') != 'admin': return jsonify({'error': 'Forbidden'}), 403
    hist  = _prev_months(12)
    fcast = _next_months(6)
    revenue  = {k: 0.0 for k, _ in hist}
    expenses = {k: 0.0 for k, _ in hist}
    cutoff = hist[0][0] + '-01'
    if USE_MEMORY_DB:
        sales_list = [s for s in getattr(db,'sales',[]) if s.get('created_at','') >= cutoff]
        inv_list   = [i for i in getattr(db,'vendor_invoices',[]) if i.get('created_at','') >= cutoff]
        pay_list   = [p for p in getattr(db,'payroll',[]) if p.get('created_at','') >= cutoff]
    else:
        sales_list = list(db.sales.find({'created_at': {'$gte': cutoff}}))
        inv_list   = list(db.vendor_invoices.find({'created_at': {'$gte': cutoff}}))
        pay_list   = list(db.payroll.find({'created_at': {'$gte': cutoff}}))
    for s in sales_list:
        k = s.get('created_at','')[:7]
        if k in revenue: revenue[k] += float(s.get('total', 0))
    for inv in inv_list:
        k = inv.get('created_at','')[:7]
        if k in expenses: expenses[k] += float(inv.get('total', 0))
    for p in pay_list:
        k = p.get('created_at','')[:7]
        if k in expenses: expenses[k] += float(p.get('net_salary', 0))
    rev_v    = [round(revenue[k], 2) for k, _ in hist]
    exp_v    = [round(expenses[k], 2) for k, _ in hist]
    profit_v = [round(rev_v[i] - exp_v[i], 2) for i in range(12)]
    rev_fc     = _forecast(rev_v, 6)
    exp_fc     = _forecast(exp_v, 6)
    profit_fc  = [round(rev_fc[i] - exp_fc[i], 2) for i in range(6)]
    slope_r, _ = _lr(rev_v)
    slope_e, _ = _lr(exp_v)
    total_rev    = sum(rev_v)
    total_exp    = sum(exp_v)
    total_profit = sum(profit_v)
    margin       = round(total_profit / total_rev * 100, 1) if total_rev else 0
    if USE_MEMORY_DB:
        pending_inv = [i for i in getattr(db,'vendor_invoices',[]) if i.get('status') not in ['paid']]
    else:
        pending_inv = list(db.vendor_invoices.find({'status': {'$ne': 'paid'}}))
    outstanding = sum(float(i.get('amount_due', i.get('total', 0))) for i in pending_inv)
    recs = []
    if margin < 10:
        recs.append({'level':'danger','icon':'fa-exclamation-circle','title':'Critical: Low Profit Margin',
                     'message':f'Margin is {margin:.1f}%. Target 20%+. Reduce procurement costs and eliminate low-margin products immediately.'})
    elif margin < 20:
        recs.append({'level':'warning','icon':'fa-chart-pie','title':'Margin Below Target',
                     'message':f'Margin is {margin:.1f}%. Aim for 20%+. Focus on higher-margin products and negotiate better vendor pricing.'})
    else:
        recs.append({'level':'success','icon':'fa-chart-pie','title':'Healthy Profit Margin',
                     'message':f'Margin is {margin:.1f}%. Strong performance. Reinvest profits in inventory expansion and new categories.'})
    if slope_e > slope_r and slope_e > 0:
        recs.append({'level':'warning','icon':'fa-balance-scale','title':'Expenses Growing Faster Than Revenue',
                     'message':'Cost growth is outpacing revenue. Audit vendor contracts and payroll costs immediately.'})
    if outstanding > total_rev * 0.1:
        recs.append({'level':'warning','icon':'fa-file-invoice','title':'High Outstanding Payables',
                     'message':f'PKR {outstanding:,.0f} in unpaid vendor invoices. Prioritize cash flow management and overdue payments.'})
    six_proj_profit = round(sum(profit_fc), 2)
    recs.append({'level':'info','icon':'fa-robot','title':'6-Month Profit Projection',
                 'message':f'Projected profit: PKR {six_proj_profit:,.0f}. {"Maintain current trajectory." if six_proj_profit > 0 else "Losses projected — take immediate cost reduction action."}'})
    return jsonify({
        'hist_labels': [l for _, l in hist], 'fcast_labels': [l for _, l in fcast],
        'revenue': rev_v, 'rev_forecast': rev_fc,
        'expenses': exp_v, 'exp_forecast': exp_fc,
        'profit': profit_v, 'profit_forecast': profit_fc,
        'kpis': {'total_revenue': round(total_rev, 2), 'total_expenses': round(total_exp, 2),
                 'total_profit': round(total_profit, 2), 'profit_margin': margin,
                 'outstanding_payables': round(outstanding, 2),
                 'six_month_proj_profit': six_proj_profit,
                 'trend_revenue': _trend_label(slope_r), 'trend_expenses': _trend_label(slope_e)},
        'recommendations': recs,
    })


@app.route('/api/predictive/vendors')
def api_predictive_vendors():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    if session.get('role') != 'admin': return jsonify({'error': 'Forbidden'}), 403
    hist   = _prev_months(12)
    cutoff = hist[0][0] + '-01'
    today  = datetime.utcnow().strftime('%Y-%m-%d')
    if USE_MEMORY_DB:
        pos      = [p for p in getattr(db,'purchase_orders_v2',[]) if p.get('created_at','') >= cutoff]
        invoices = [i for i in getattr(db,'vendor_invoices',[]) if i.get('created_at','') >= cutoff]
        payments = [p for p in getattr(db,'vendor_payments_v2',[]) if p.get('created_at','') >= cutoff]
    else:
        pos      = list(db.purchase_orders_v2.find({'created_at': {'$gte': cutoff}}))
        invoices = list(db.vendor_invoices.find({'created_at': {'$gte': cutoff}}))
        payments = list(db.vendor_payments_v2.find({'created_at': {'$gte': cutoff}}))
    vendor_data = {}
    for po in pos:
        vn = po.get('vendor_name', 'Unknown')
        if vn not in vendor_data:
            vendor_data[vn] = {'spend': 0.0, 'invoices': 0, 'paid': 0.0, 'overdue': 0, 'pos': 0}
        vendor_data[vn]['spend'] += float(po.get('total', 0))
        vendor_data[vn]['pos']   += 1
    for inv in invoices:
        vn = inv.get('vendor_name', 'Unknown')
        if vn not in vendor_data:
            vendor_data[vn] = {'spend': 0.0, 'invoices': 0, 'paid': 0.0, 'overdue': 0, 'pos': 0}
        vendor_data[vn]['invoices'] += 1
        due = inv.get('due_date', '')
        if inv.get('status') != 'paid' and due and due < today:
            vendor_data[vn]['overdue'] += 1
    for pmt in payments:
        vn = pmt.get('vendor_name', 'Unknown')
        if vn not in vendor_data:
            vendor_data[vn] = {'spend': 0.0, 'invoices': 0, 'paid': 0.0, 'overdue': 0, 'pos': 0}
        vendor_data[vn]['paid'] += float(pmt.get('amount', 0))
    def _vscore(vd):
        base = 70
        if vd['invoices']:
            base -= min(30, vd['overdue'] / vd['invoices'] * 100)
        if vd['spend'] > 0 and vd['paid'] / vd['spend'] > 0.8:
            base = min(100, base + 10)
        return round(base)
    top_vendors = sorted(vendor_data.items(), key=lambda x: x[1]['spend'], reverse=True)[:10]
    total_spend = sum(v['spend'] for v in vendor_data.values())
    overdue_count = sum(1 for v in vendor_data.values() if v['overdue'] > 0)
    avg_score = round(sum(_vscore(v) for v in vendor_data.values()) / len(vendor_data), 1) if vendor_data else 0
    vendors_out = [{'name': n, 'spend': round(vd['spend'], 2), 'invoices': vd['invoices'],
                    'paid': round(vd['paid'], 2), 'overdue': vd['overdue'],
                    'pos': vd['pos'], 'score': _vscore(vd)} for n, vd in top_vendors]
    recs = []
    if top_vendors and total_spend:
        top_share = top_vendors[0][1]['spend'] / total_spend * 100
        if top_share > 50:
            recs.append({'level':'danger','icon':'fa-building','title':'High Vendor Dependency',
                         'message':f'"{top_vendors[0][0]}" drives {top_share:.0f}% of spend. Onboard alternative vendors to reduce supply risk.'})
    if overdue_count:
        recs.append({'level':'warning','icon':'fa-clock','title':f'{overdue_count} Vendors Have Overdue Invoices',
                     'message':'Settle overdue invoices to protect vendor relationships and maintain credit terms.'})
    low_score = [n for n, vd in top_vendors if _vscore(vd) < 70]
    if low_score:
        recs.append({'level':'warning','icon':'fa-star-half-alt','title':'Low-Reliability Vendors Detected',
                     'message':f'{", ".join(low_score[:3])} scored below 70. Renegotiate terms or find alternatives.'})
    if len(vendor_data) < 3:
        recs.append({'level':'danger','icon':'fa-exclamation-triangle','title':'Too Few Active Vendors',
                     'message':'Fewer than 3 active vendors detected. Onboard additional suppliers immediately to ensure supply continuity.'})
    recs.append({'level':'success','icon':'fa-handshake','title':'Vendor Expansion Opportunity',
                 'message':'Add 2-3 new vendors per product category to improve price competition and reduce single-source risk.'})
    return jsonify({
        'vendors': vendors_out, 'hist_labels': [l for _, l in hist],
        'kpis': {'total_vendors': len(vendor_data), 'total_spend': round(total_spend, 2),
                 'avg_score': avg_score, 'overdue_count': overdue_count},
        'recommendations': recs,
    })


@app.route('/api/predictive/overview')
def api_predictive_overview():
    if 'user_id' not in session: return jsonify({'error': 'Unauthorized'}), 401
    if session.get('role') != 'admin': return jsonify({'error': 'Forbidden'}), 403
    hist  = _prev_months(12)
    fcast = _next_months(12)
    cutoff = hist[0][0] + '-01'
    revenue  = {k: 0.0 for k, _ in hist}
    expenses = {k: 0.0 for k, _ in hist}
    if USE_MEMORY_DB:
        sales_list = [s for s in getattr(db,'sales',[]) if s.get('created_at','') >= cutoff]
        inv_list   = [i for i in getattr(db,'vendor_invoices',[]) if i.get('created_at','') >= cutoff]
    else:
        sales_list = list(db.sales.find({'created_at': {'$gte': cutoff}}))
        inv_list   = list(db.vendor_invoices.find({'created_at': {'$gte': cutoff}}))
    for s in sales_list:
        k = s.get('created_at','')[:7]
        if k in revenue: revenue[k] += float(s.get('total', 0))
    for inv in inv_list:
        k = inv.get('created_at','')[:7]
        if k in expenses: expenses[k] += float(inv.get('total', 0))
    rev_v    = [round(revenue[k], 2) for k, _ in hist]
    exp_v    = [round(expenses[k], 2) for k, _ in hist]
    profit_v = [round(rev_v[i] - exp_v[i], 2) for i in range(12)]
    rev_fc   = _forecast(rev_v, 12)
    exp_fc   = _forecast(exp_v, 12)
    profit_fc = [round(rev_fc[i] - exp_fc[i], 2) for i in range(12)]
    slope_r, _ = _lr(rev_v)
    total_rev    = sum(rev_v)
    total_profit = sum(profit_v)
    margin       = round(total_profit / total_rev * 100, 1) if total_rev else 0
    # Health score
    slope_norm = max(-1.0, min(1.0, slope_r / (max(rev_v) + 1) * 12))
    sales_score = max(5, min(25, int(15 + slope_norm * 10)))
    if USE_MEMORY_DB:
        total_prs = len(getattr(db,'purchase_requisitions',[]))
        total_pos_count = len(getattr(db,'purchase_orders_v2',[]))
        vcount = len(set(p.get('vendor_name','') for p in getattr(db,'purchase_orders_v2',[])))
    else:
        total_prs = db.purchase_requisitions.count_documents({})
        total_pos_count = db.purchase_orders_v2.count_documents({})
        vcount = len(db.purchase_orders_v2.distinct('vendor_name'))
    conv = total_pos_count / total_prs if total_prs else 0
    proc_score   = int(min(25, conv * 25))
    fin_score    = 25 if margin >= 25 else (18 if margin >= 15 else (10 if margin >= 5 else 3))
    vendor_score = min(25, vcount * 5)
    health_score = sales_score + proc_score + fin_score + vendor_score
    year_rev_proj = round(sum(rev_fc), 2)
    recs = []
    if health_score >= 80:
        recs.append({'level':'success','icon':'fa-trophy','title':'Business in Excellent Health',
                     'message':f'Score {health_score}/100. Outstanding performance across all departments. Focus on scaling operations.'})
    elif health_score >= 60:
        recs.append({'level':'info','icon':'fa-chart-bar','title':'Business Performing Well',
                     'message':f'Score {health_score}/100. Good foundation. Address weaker areas to reach the next performance level.'})
    elif health_score >= 40:
        recs.append({'level':'warning','icon':'fa-exclamation-triangle','title':'Business Needs Attention',
                     'message':f'Score {health_score}/100. Significant gaps detected. Prioritize procurement efficiency and margin improvement.'})
    else:
        recs.append({'level':'danger','icon':'fa-exclamation-circle','title':'Business Requires Immediate Action',
                     'message':f'Score {health_score}/100. Critical issues across multiple areas. Take immediate corrective steps in all departments.'})
    recs.append({'level':'info','icon':'fa-robot','title':'1-Year Revenue AI Projection',
                 'message':f'Based on current trend, projected revenue for next 12 months: PKR {year_rev_proj:,.0f}. {"Strong trajectory." if slope_r > 0 else "Declining trend — intervention needed."}'})
    if margin < 15:
        recs.append({'level':'warning','icon':'fa-coins','title':'Improve Profit Margins',
                     'message':f'Current margin {margin:.1f}%. Review your top expense categories and eliminate inefficiencies to reach 20%+.'})
    return jsonify({
        'hist_labels': [l for _, l in hist], 'fcast_labels': [l for _, l in fcast],
        'revenue': rev_v, 'rev_forecast': rev_fc,
        'expenses': exp_v, 'exp_forecast': exp_fc,
        'profit': profit_v, 'profit_forecast': profit_fc,
        'health_score': health_score,
        'score_breakdown': {'sales': sales_score, 'procurement': proc_score,
                            'finance': fin_score, 'vendors': vendor_score},
        'kpis': {'total_revenue': round(total_rev, 2), 'total_profit': round(total_profit, 2),
                 'profit_margin': margin, 'year_rev_proj': year_rev_proj, 'trend': _trend_label(slope_r)},
        'recommendations': recs,
    })


@app.route('/admin/predictive')
def predictive_overview_page():
    if 'user_id' not in session: return redirect(url_for('login_page'))
    if session.get('role') != 'admin': return redirect_to_dashboard(session.get('role'))
    return render_template('predictive_overview.html')

@app.route('/admin/predictive/sales')
def predictive_sales_page():
    if 'user_id' not in session: return redirect(url_for('login_page'))
    if session.get('role') != 'admin': return redirect_to_dashboard(session.get('role'))
    return render_template('predictive_sales.html')

@app.route('/admin/predictive/procurement')
def predictive_procurement_page():
    if 'user_id' not in session: return redirect(url_for('login_page'))
    if session.get('role') != 'admin': return redirect_to_dashboard(session.get('role'))
    return render_template('predictive_procurement.html')

@app.route('/admin/predictive/finance')
def predictive_finance_page():
    if 'user_id' not in session: return redirect(url_for('login_page'))
    if session.get('role') != 'admin': return redirect_to_dashboard(session.get('role'))
    return render_template('predictive_finance.html')

@app.route('/admin/predictive/vendors')
def predictive_vendors_page():
    if 'user_id' not in session: return redirect(url_for('login_page'))
    if session.get('role') != 'admin': return redirect_to_dashboard(session.get('role'))
    return render_template('predictive_vendors.html')


# =============================================================================
# DOCUMENTATION ROUTES
# =============================================================================

@app.route('/docs/user-guide')
def doc_user_guide():
    if session.get('role') not in ['admin', 'manual_viewer']:
        return redirect(url_for('login_page'))
    docs_dir = os.path.join(os.path.abspath(os.path.dirname(__file__)), 'docs')
    return send_from_directory(docs_dir, 'user_guide.html')


@app.route('/docs/use-cases')
def doc_use_cases():
    if session.get('role') not in ['admin', 'manual_viewer']:
        return redirect(url_for('login_page'))
    docs_dir = os.path.join(os.path.abspath(os.path.dirname(__file__)), 'docs')
    return send_from_directory(docs_dir, 'use_cases.html')


@app.route('/docs')
def docs_home():
    """Landing page for manual_viewer role — shows both doc links."""
    if session.get('role') not in ['admin', 'manual_viewer']:
        return redirect(url_for('login_page'))
    name = session.get('full_name', 'User')
    return f'''<!DOCTYPE html>
<html><head><meta charset="UTF-8"><title>POS Documentation</title>
<link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/font-awesome/6.4.0/css/all.min.css">
<style>
  *{{margin:0;padding:0;box-sizing:border-box}}
  body{{font-family:'Segoe UI',sans-serif;background:linear-gradient(135deg,#0d1f3c,#1e3a5f,#2d6a9f);min-height:100vh;display:flex;align-items:center;justify-content:center;padding:24px}}
  .box{{background:#fff;border-radius:20px;padding:48px 44px;max-width:560px;width:100%;box-shadow:0 20px 60px rgba(0,0,0,.35);text-align:center}}
  .logo{{font-size:52px;margin-bottom:16px}}
  h1{{font-size:24px;font-weight:900;color:#1e3a5f;margin-bottom:6px}}
  .sub{{font-size:14px;color:#6b7a99;margin-bottom:36px}}
  .card{{display:flex;align-items:center;gap:18px;background:#f4f7fb;border:1.5px solid #dde3ee;border-radius:14px;padding:20px 24px;margin-bottom:16px;text-decoration:none;transition:transform .2s,box-shadow .2s;text-align:left}}
  .card:hover{{transform:translateY(-3px);box-shadow:0 8px 24px rgba(0,0,0,.1)}}
  .card-icon{{width:52px;height:52px;border-radius:12px;display:flex;align-items:center;justify-content:center;font-size:24px;flex-shrink:0;color:#fff}}
  .card-title{{font-size:15px;font-weight:800;color:#1e3a5f;margin-bottom:3px}}
  .card-desc{{font-size:12.5px;color:#6b7a99}}
  .logout{{margin-top:28px;font-size:13px;color:#aab}}
  .logout a{{color:#e74c3c;text-decoration:none;font-weight:700}}
</style></head>
<body>
<div class="box">
  <div class="logo">📚</div>
  <h1>POS System Documentation</h1>
  <div class="sub">Welcome, {name}. Download or read the manuals below.</div>
  <a href="/docs/user-guide" class="card">
    <div class="card-icon" style="background:linear-gradient(135deg,#1e3a5f,#2d6a9f)"><i class="fas fa-book"></i></div>
    <div>
      <div class="card-title">User Guide</div>
      <div class="card-desc">Complete system manual — every screen, every field explained</div>
    </div>
  </a>
  <a href="/docs/use-cases" class="card">
    <div class="card-icon" style="background:linear-gradient(135deg,#6c3483,#8e44ad)"><i class="fas fa-clipboard-list"></i></div>
    <div>
      <div class="card-title">Use Cases</div>
      <div class="card-desc">30 real-world business scenarios with step-by-step instructions</div>
    </div>
  </a>
  <div class="logout"><a href="/logout"><i class="fas fa-sign-out-alt"></i> Logout</a></div>
</div>
</body></html>'''


# =============================================================================
# MAIN ENTRY
# =============================================================================

if __name__ == '__main__':
    print("Initializing POS System...")
    init_db()
    init_gl_accounts()
    print(f"Database: {'MongoDB' if not USE_MEMORY_DB else 'In-Memory (demo mode)'}")
    port = int(os.environ.get('PORT', 5000))
    print(f"Starting POS server on http://0.0.0.0:{port}")
    app.run(debug=False, port=port, host='0.0.0.0')